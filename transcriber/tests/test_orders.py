"""Перечень поручений: разбор ответа «Искры» (`app/orders.py`) и Word по эталону
(`app/orders_docx.py` — копия общего модуля `portal/shared/orders_docx.py`).

Эталон — `moderator/docs/примеры/Перечень поручений 25.05.2026 (пример).docx` и решения
пользователя 02.10 (TRACK.md → «Модератор» → «Поручения»).
"""
from __future__ import annotations

import io
import json
import sys
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm

HERE = Path(__file__).resolve().parent
SERVICE = HERE.parent
sys.path.insert(0, str(SERVICE))

from app import orders, orders_docx, summarizer  # noqa: E402

WORKSPACE = SERVICE.parent
REFERENCE = WORKSPACE / "moderator" / "docs" / "примеры" / "Перечень поручений 25.05.2026 (пример).docx"
SHARED = WORKSPACE / "portal" / "shared" / "orders_docx.py"

ROLE = "Заместитель Председателя Правительства Российской Федерации"
ORDERS = [
    {"text": "Минпросвещения России (С.С. Кравцову) подготовить предложения — по итогам совещания.",
     "assignee": "С.С. Кравцов", "deadline": "25.06.2026"},
    {"text": "Рособрнадзору (А.А. Музаеву) обеспечить доступ к системе.", "assignee": "", "deadline": ""},
]


# ── разбор ответа «Искры» ─────────────────────────────────────────────────────

def scripted_run(*answers, spent=(30, 10)):
    calls = []

    def fake(message, instructions, cancelled=lambda: False, error=None, cancel=None, what=""):
        calls.append(message)
        answer = answers[min(len(calls) - 1, len(answers) - 1)]
        return answer, {"in": spent[0], "out": spent[1], "total": spent[0] + spent[1]}

    return fake, calls


def test_json_po_skheme_pustye_otbrasyvayutsya_srok_ne_vydumyvaetsya(monkeypatch):
    answer = ('```json\n{"orders":['
              '{"text":"Минтруду России (А.О. Котякову) представить доклад — о зарплате.",'
              '"assignee":"А.О. Котяков","deadline":"2026-07-01"},'
              '{"text":"   ","assignee":"кто-то","deadline":"01.08.2026"},'
              '{"text":"Минфину России проработать вопрос.","assignee":"Минфин России","deadline":"не указан"},'
              '{"text":"Росстату подготовить данные."},'
              '"Рособрнадзору провести проверку.",'
              '42]}\n```')
    fake, calls = scripted_run(answer)
    monkeypatch.setattr(summarizer, "run", fake)
    found, spent = orders.extract("[00:00] Говорящий 1: поручаю …")
    assert len(calls) == 1
    assert [o["text"] for o in found] == [
        "Минтруду России (А.О. Котякову) представить доклад – о зарплате.",   # тире короткое
        "Минфину России проработать вопрос.",
        "Росстату подготовить данные.",
        "Рособрнадзору провести проверку.",
    ]
    assert [o["deadline"] for o in found] == ["01.07.2026", "", "", ""]   # срока нет — пусто, не выдуман
    assert found[0]["assignee"] == "А.О. Котяков"
    assert spent == {"in": 30, "out": 10, "total": 40}
    assert "дословная расшифровка" in calls[0] and "Говорящий 1" in calls[0]


def test_proza_vmesto_json_povtor_strogo_tokeny_skladyvayutsya(monkeypatch):
    fake, calls = scripted_run("Вот перечень поручений: …", '{"orders": []}')
    monkeypatch.setattr(summarizer, "run", fake)
    found, spent = orders.extract("текст")
    assert found == []
    assert len(calls) == 2 and calls[1].startswith(orders.STRICT)
    assert spent["total"] == 80


def test_dvazhdy_proza_oshibka_s_raskhodom(monkeypatch):
    fake, _calls = scripted_run("не json", "снова не json")
    monkeypatch.setattr(summarizer, "run", fake)
    with pytest.raises(orders.OrdersError) as err:
        orders.extract("текст")
    assert err.value.spent["total"] == 80


def test_check_schema_ne_obekt():
    with pytest.raises(ValueError):
        orders.check_schema("строка")
    assert orders.check_schema({"orders": {"text": "Одно поручение."}}) == [
        {"text": "Одно поручение.", "assignee": "", "deadline": ""}]


def test_otmena_i_oshibka_iskry_nesut_raskhod(monkeypatch):
    def failing(message, instructions, cancelled=lambda: False, error=None, cancel=None, what=""):
        raise error("«Искра» прервала перечень поручений: сбой", {"in": 7, "out": 1, "total": 8})

    monkeypatch.setattr(summarizer, "run", failing)
    with pytest.raises(orders.OrdersError) as err:
        orders.extract("текст")
    assert err.value.spent["total"] == 8


# ── Word по эталону ───────────────────────────────────────────────────────────

def build(orders_=ORDERS, chair="Чернышенко Дмитрий Николаевич", role=ROLE):
    return Document(io.BytesIO(orders_docx.build(orders_, chair, role, date(2026, 5, 25))))


def texts(doc):
    return [p.text for p in doc.paragraphs]


def test_polya_shrift_otstup_kak_v_etalone():
    doc = build()
    s = doc.sections[0]
    assert round(s.page_width.mm) == 210 and round(s.page_height.mm) == 297
    assert s.page_width < s.page_height                                   # книжный
    assert [round(x.mm, 1) for x in (s.top_margin, s.right_margin, s.bottom_margin, s.left_margin)] == \
        [20.0, 15.0, 20.0, 20.0]
    body = [p for p in doc.paragraphs if p.text.startswith(("1.", "2.", "Срок"))]
    assert body
    for p in body:
        assert p.style.name == "No Spacing"
        assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert p.paragraph_format.line_spacing == pytest.approx(1.15)
        assert abs(p.paragraph_format.first_line_indent - Mm(12.5)) < Mm(0.05)
        for run in p.runs:
            assert run.font.name == "Times New Roman" and run.font.size.pt == 16


@pytest.mark.skipif(not REFERENCE.exists(), reason="эталон «Совещания» не рядом (копия одного сервиса)")
def test_svoystva_sovpadayut_s_etalonom():
    ref, doc = Document(str(REFERENCE)), build()
    rs, ds = ref.sections[0], doc.sections[0]
    for attr in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        assert abs(getattr(rs, attr) - getattr(ds, attr)) < Mm(0.05), attr
    ref_item = next(p for p in ref.paragraphs if p.text.startswith("3."))
    our_item = next(p for p in doc.paragraphs if p.text.startswith("1."))
    for attr in ("alignment", "line_spacing", "first_line_indent"):
        assert getattr(ref_item.paragraph_format, attr) == getattr(our_item.paragraph_format, attr), attr
    assert ref_item.style.name == our_item.style.name == "No Spacing"
    assert ref_item.runs[0].font.size == our_item.runs[0].font.size
    # шапка: по центру, полужирно, две строки заглавными — как в эталоне
    ref_head, our_head = ref.paragraphs[0], doc.paragraphs[0]
    assert ref_head.alignment == our_head.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert [l.strip() for l in our_head.text.split("\n")] == [l.strip() for l in ref_head.text.split("\n")]
    assert all(r.bold for r in our_head.runs)
    # подпись справа последним абзацем
    assert ref.paragraphs[-1].alignment == doc.paragraphs[-1].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert ref.paragraphs[-1].text == doc.paragraphs[-1].text == "Д.Н. Чернышенко"
    # строка срока датой — в том же виде, что в эталоне
    assert "Срок – до 25 июня 2026 г." in texts(ref) and "Срок – до 25 июня 2026 г." in texts(doc)


def test_poryadok_abzacev_i_net_sroka_bez_daty():
    lines = texts(build())
    assert lines == [
        "ЗАМЕСТИТЕЛЬ ПРЕДСЕДАТЕЛЯ ПРАВИТЕЛЬСТВА \nРОССИЙСКОЙ ФЕДЕРАЦИИ",
        "", "",                                           # две пустые строки под шапкой
        "1. Минпросвещения России (С.С. Кравцову) подготовить предложения – по итогам совещания.",
        "Срок – до 25 июня 2026 г.",
        "",                                               # пустая строка между поручениями
        "2. Рособрнадзору (А.А. Музаеву) обеспечить доступ к системе.",   # даты нет — «Срок» нет
        "", "", "",                                       # три пустые строки перед подписью, как в эталоне
        "Д.Н. Чернышенко",
    ]
    assert sum(1 for l in lines if l.startswith("Срок")) == 1
    joined = "\n".join(lines)
    assert "—" not in joined                               # только короткое тире
    assert "25.05.2026" not in joined and "Перечень" not in joined   # названия и даты в теле нет


def test_bez_dolzhnosti_i_predsedatelya_net_shapki_i_podpisi():
    lines = texts(build(chair="", role=""))
    assert lines[0].startswith("1. ")
    assert lines[-1].startswith("2. ")


def test_srok_slovami_i_formaty_daty():
    assert orders_docx.deadline_line("") == ""
    assert orders_docx.deadline_line("2026-06-25") == "Срок – до 25 июня 2026 г."
    assert orders_docx.deadline_line("1.9.2026") == "Срок – до 1 сентября 2026 г."
    assert orders_docx.deadline_line("начиная с 1 сентября 2026 г.") == "Срок – начиная с 1 сентября 2026 г."
    assert orders_docx.signature("Чернышенко Д.Н.") == "Д.Н. Чернышенко"
    assert orders_docx.signature("Д. Н. Чернышенко") == "Д.Н. Чернышенко"


def test_imya_fayla():
    assert orders_docx.file_name(date(2026, 5, 25)) == "Перечень поручений 25.05.2026.docx"
    assert orders_docx.file_name(None) == "Перечень поручений.docx"


def test_dlinnaya_transkripciya_po_chastyam_bez_poter(monkeypatch):
    """Длиннее предела — не вырезаем середину (ревью RC-17, Н-4): саммари по частям и итог,
    поручения по частям с объединением без дублей; расход складывается."""
    monkeypatch.setattr(summarizer, "MAX_CHARS", 200)
    paragraphs = [f"[{i:02d}:00] Говорящий 1: абзац номер {i} " + "слово " * 10 for i in range(12)]
    text = "\n\n".join(paragraphs)
    parts = summarizer.split_text(text)
    assert len(parts) > 2 and all(len(p) <= 200 for p in parts)
    assert "\n\n".join(parts) == text                      # ни знака не потеряно, порядок тот же

    calls = []

    def fake_run(message, instructions, cancelled=lambda: False, error=None, cancel=None, what=""):
        calls.append(message)
        if message.startswith(summarizer.MERGE_PROMPT):
            return "итог", {"in": 1, "out": 1, "total": 2}
        if "Схема ответа" in message:
            n = len([c for c in calls if "Схема ответа" in c])
            orders_ = [{"text": "Минфину России проработать вопрос.", "deadline": "01.09.2026" if n == 3 else ""},
                       {"text": f"Поручение из части {n}."}]
            return json.dumps({"orders": orders_}), {"in": 3, "out": 1, "total": 4}
        return f"саммари части {len(calls)}", {"in": 5, "out": 1, "total": 6}

    monkeypatch.setattr(summarizer, "run", fake_run)
    summary, spent = summarizer.summarize(text)
    assert summary == "итог"
    assert len(calls) == len(parts) + 1
    assert "абзац номер 6" in "".join(calls)                  # середина ушла в «Искру»
    assert spent["total"] == 6 * len(parts) + 2

    calls.clear()
    found, spent = orders.extract(text)
    assert len(calls) == len(parts)
    texts_ = [o["text"] for o in found]
    assert texts_.count("Минфину России проработать вопрос.") == 1       # дубль частей — один раз
    assert found[0]["deadline"] == "01.09.2026"                          # срок взят у дубля, где он назван
    assert [t for t in texts_ if t.startswith("Поручение из части")] == \
        [f"Поручение из части {n}." for n in range(1, len(parts) + 1)]
    assert spent["total"] == 4 * len(parts)


def test_sboy_chasti_neset_raskhod_predydushchih(monkeypatch):
    monkeypatch.setattr(summarizer, "MAX_CHARS", 50)
    calls = []

    def fake_run(message, instructions, cancelled=lambda: False, error=None, cancel=None, what=""):
        calls.append(1)
        if len(calls) == 3:
            raise (error or summarizer.SummaryError)("сбой", {"in": 1, "out": 0, "total": 1})
        return "ok", {"in": 10, "out": 0, "total": 10}

    monkeypatch.setattr(summarizer, "run", fake_run)
    with pytest.raises(summarizer.SummaryError) as err:
        summarizer.summarize("\n\n".join("абзац " * 8 for _ in range(6)))
    assert err.value.spent["total"] == 21


@pytest.mark.skipif(not SHARED.exists(), reason="portal/shared не рядом (копия одного сервиса)")
def test_kopiya_sovpadaet_s_obshchim_modulem():
    """Копию раскладывает `portal/scripts/sync_shared.sh`; правка копии руками — ошибка."""
    assert (SERVICE / "app" / "orders_docx.py").read_bytes() == SHARED.read_bytes()
