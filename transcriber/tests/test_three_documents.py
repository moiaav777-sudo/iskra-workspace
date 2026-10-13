"""Поток «Разобрать с Искрой» → три документа (решение пользователя 02.10).

Запись после разбора содержит транскрипцию, саммари и перечень поручений; каждый правится с
версией и скачивается в Word; три документа уходят в «Совещание» (тела запросов готовит сервер).
Движок «Искры» здесь настоящий — подменена только сеть (`iskra_engine._urlopen`), саммари и
поручения — управляемыми подделками, как в `test_contract.py`.
"""
from __future__ import annotations

import io
import json
import sqlite3
import sys
import threading
import time
from pathlib import Path

import pytest
from unittest.mock import ANY
from docx import Document
from fastapi.testclient import TestClient

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

from app import audit, config, db, jobs, main, summarizer  # noqa: E402
from app.engine import iskra_engine  # noqa: E402
from test_iskra_engine import FakeIskra, fast  # noqa: E402

from app import orders as _orders_module  # noqa: E402

# настоящие функции — до подмены фикстурой (для тестов, где нужен настоящий прогон с подменной сетью)
REAL_SUMMARIZE = summarizer.summarize
REAL_EXTRACT = _orders_module.extract

SECRET = "stend-secret-3"
EMP = {"X-Portal-Secret": SECRET, "X-Portal-User": "emp", "X-Portal-Role": "employee",
       "X-Portal-Name": "%D0%98%D0%B2%D0%B0%D0%BD%D0%BE%D0%B2"}
ADMIN = {**EMP, "X-Portal-User": "admin", "X-Portal-Role": "admin"}

# основной путь (правила API «Искры» 05.10): расшифровка приходит TXT-файлом в результате прогона
TRANSCRIPT = (
    "[00:00] Говорящий 1: Добрый день, коллеги. Обсуждаем систему показателей.\n"
    "[00:45] Говорящий 2: Готовы доложить по первому вопросу.\n"
    "[01:30] Говорящий 1: Поручаю Минпросвещения подготовить доклад до двадцать пятого июня.\n"
)
ORDER = {"text": "Минпросвещения России (С.С. Кравцову) подготовить доклад о системе показателей",
         "assignee": "С.С. Кравцов", "deadline": "25.06.2026"}


def fake_summary(text, cancelled=None):
    return "Обсудили систему показателей.\n\nРешили готовить доклад.", {"in": 100, "out": 20, "total": 120}


def fake_orders(text, cancelled=None):
    return [dict(ORDER), {"text": "Рособрнадзору проработать доступ", "assignee": "", "deadline": ""}], \
        {"in": 40, "out": 10, "total": 50}


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "transcriber.db")
    monkeypatch.setattr(config, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(config, "PORTAL_SECRET", SECRET)
    monkeypatch.setattr(config, "DEV", False)
    monkeypatch.setattr(config, "ENGINE", "iskra")
    monkeypatch.setenv("ISKRA_API_KEY", "test-key")
    monkeypatch.setenv("ISKRA_URL", "https://iskra.test")
    fake = FakeIskra(file_text=TRANSCRIPT)
    fast(monkeypatch)
    monkeypatch.setattr(iskra_engine, "_urlopen", fake)
    monkeypatch.setattr(main.summarizer, "summarize", fake_summary)
    monkeypatch.setattr(main.orders_mod, "extract", fake_orders)
    audit.sent.clear()
    jobs.reset()
    db.init()
    with TestClient(main.app) as c:
        c.iskra = fake
        yield c
    jobs.reset()


def upload(client, headers=EMP, name="Совещание 25.05.mp3", content=b"ID3" + b"\x00" * 5000):
    return client.post("/api/records", headers=headers,
                       files={"file": (name, io.BytesIO(content), "audio/mpeg")})


def finish(client, response, timeout=10.0):
    assert response.status_code == 202, response.text
    job = response.json()
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        job = client.get(f"/api/jobs/{job['id']}", headers=EMP).json()
        if job["status"] in ("done", "error", "cancelled"):
            return job
        time.sleep(0.02)
    raise AssertionError("обработка не закончилась")


def record_after_upload(client):
    job = finish(client, upload(client))
    assert job["status"] == "done", job
    return client.get(f"/api/records/{job['record']['id']}", headers=EMP).json()


# ── три документа ─────────────────────────────────────────────────────────────

def test_posle_razbora_tri_dokumenta(client):
    record = record_after_upload(client)
    assert record["state"] == "done"
    assert record["engine"] == "iskra"
    assert record["draft_text"].startswith("[00:00] Говорящий 1: Добрый день, коллеги.")
    assert "[01:30] Говорящий 1: Поручаю" in record["draft_text"]
    assert record["summary"].startswith("Обсудили")
    # html поручения для блока карточки выводится из текста (решение пользователя 13.10)
    # и постоянный id (ревью RC-2026-10-13-2, К-1)
    assert record["orders"][0] == {**ORDER, "html": f"<p>{ORDER['text']}</p>", "id": ANY} and len(record["orders"]) == 2
    assert len({o["id"] for o in record["orders"]}) == 2
    # расход: расшифровка (один прогон, 15) + саммари 120 + поручения 50
    assert record["tokens"] == 15 + 120 + 50
    # файл записи удалён после разбора — приватность
    assert list(config.UPLOAD_DIR.iterdir()) == []


def test_zhurnal_tri_sobytiya_iskry_bez_soderzhaniya(client):
    record_after_upload(client)
    sent = [r for r in audit.sent if r["service"] == "transcriber"]
    by_action = {r["action"]: r for r in sent}
    assert (by_action["запись расшифрована"]["tokens_in"], by_action["запись расшифрована"]["tokens_out"]) == (10, 5)
    assert by_action["саммари записи готово"]["tokens_in"] == 100
    assert by_action["перечень поручений записи готов"]["tokens_out"] == 10
    assert sum(1 for r in sent if r["action"] == "запись расшифрована") == 1
    for r in sent:
        for secret in ("Добрый день", "Минпросвещения", "систему показателей", "Рособрнадзор"):
            assert secret not in r["details"], (r["action"], r["details"])


def test_bez_klyucha_chestnaya_oshibka_servis_zhiv(client, monkeypatch):
    monkeypatch.delenv("ISKRA_API_KEY", raising=False)
    job = finish(client, upload(client))
    assert job["status"] == "error"
    assert "Ключ «Искры» не настроен" in job["message"]
    record = client.get(f"/api/records/{job['record_id']}", headers=EMP).json()
    assert record["state"] == "error" and "Ключ" in record["state_note"]
    assert list(config.UPLOAD_DIR.iterdir()) == []         # файл не остался на диске
    assert client.get("/healthz").status_code == 200
    assert client.get("/api/engine", headers=EMP).json()["ready"] is False


def test_413_iskry_zapis_s_oshibkoy(client):
    client.iskra.upload_status = 413
    job = finish(client, upload(client))
    assert job["status"] == "error" and "500 МБ" in job["message"]


def test_fayl_bolshe_predela_otklonyaetsya_srazu(client, monkeypatch):
    monkeypatch.setattr(config, "MAX_FILE_BYTES", 1024)
    response = upload(client, content=b"0" * 4096)
    assert response.status_code == 413
    assert "500 МБ" in response.json()["message"]
    assert client.iskra.uploads == []


def test_otmena_vo_vremya_rasshifrovki_schitaet_tokeny(client):
    started = {"go": threading.Event(), "stop": threading.Event()}

    def hold(n):
        started["go"].set()
        started["stop"].wait(5)

    client.iskra.on_run = hold
    response = upload(client)
    job_id = response.json()["id"]
    assert started["go"].wait(5)
    client.post(f"/api/jobs/{job_id}/cancel", headers=EMP)
    started["stop"].set()
    job = finish(client, response)
    assert job["status"] == "cancelled"
    record = client.get(f"/api/records/{job['record_id']}", headers=EMP).json()
    assert record["tokens"] == 15                           # снятый прогон оплачен — учтён
    cancelled = [r for r in audit.sent if r["action"] == "расшифровка записи отменена"]
    assert len(cancelled) == 1 and cancelled[0]["tokens_in"] == 10


# ── скачивание трёх документов ───────────────────────────────────────────────

def docx_text(response):
    return "\n".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)


def test_skachivanie_treh_word(client):
    record = record_after_upload(client)
    rid = record["id"]
    tr = client.get(f"/api/records/{rid}/export/transcript.docx", headers=EMP)
    su = client.get(f"/api/records/{rid}/export/summary.docx", headers=EMP)
    orr = client.get(f"/api/records/{rid}/export/orders.docx", headers=EMP)
    for r in (tr, su, orr):
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert "attachment" in r.headers["content-disposition"]
    text = docx_text(tr)
    assert "[00:45] Говорящий 2: Готовы доложить" in text          # таймкоды и говорящие печатаются
    bold = [run.text for p in Document(io.BytesIO(tr.content)).paragraphs for run in p.runs if run.bold]
    assert "Говорящий 2: " in bold
    assert "Решили готовить доклад." in docx_text(su)
    orders_text = docx_text(orr)
    assert "1. Минпросвещения России (С.С. Кравцову)" in orders_text
    assert "Срок – до 25 июня 2026 г." in orders_text
    assert orders_text.count("Срок") == 1                          # у второго срока нет — строки нет
    today = time.strftime("%d.%m.%Y")                               # даты совещания нет — день загрузки
    from urllib.parse import unquote
    assert unquote(orr.headers["content-disposition"].split("UTF-8''")[1]) == f"Перечень поручений {today}.docx"
    assert client.get(f"/api/records/{rid}/export/nothing.docx", headers=EMP).status_code == 404


def test_skachivanie_bez_portala_otkaz_a_ne_padenie(client):
    rid = record_after_upload(client)["id"]
    for kind in ("transcript", "summary", "orders"):
        r = client.get(f"/api/records/{rid}/export/{kind}.docx")
        assert r.status_code == 403
        assert r.json()["message"]


def test_shapka_i_podpis_perechnya_iz_kartochki(client):
    rid = record_after_upload(client)["id"]
    saved = client.put(f"/api/records/{rid}/meta", headers=EMP, json={
        "chair": "Чернышенко Дмитрий Николаевич",
        "chair_role": "Заместитель Председателя Правительства Российской Федерации",
        "meeting_date": "2026-05-25"})
    assert saved.status_code == 200
    assert saved.json()["meeting_date"] == "25.05.2026"
    r = client.get(f"/api/records/{rid}/export/orders.docx", headers=EMP)
    paragraphs = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
    assert paragraphs[0] == "ЗАМЕСТИТЕЛЬ ПРЕДСЕДАТЕЛЯ ПРАВИТЕЛЬСТВА \nРОССИЙСКОЙ ФЕДЕРАЦИИ"
    assert paragraphs[-1] == "Д.Н. Чернышенко"
    from urllib.parse import unquote
    assert unquote(r.headers["content-disposition"].split("UTF-8''")[1]) == "Перечень поручений 25.05.2026.docx"
    assert client.put(f"/api/records/{rid}/meta", headers=EMP,
                      json={"meeting_date": "вчера"}).status_code == 400
    assert any(a["action"] == "изменены сведения для документов записи" for a in audit.sent)


# ── правка с версией ─────────────────────────────────────────────────────────

def test_pravka_porucheniy_s_versiey(client):
    record = record_after_upload(client)
    rid, version = record["id"], record["orders_version"]
    assert client.put(f"/api/records/{rid}/orders", headers=EMP,
                      json={"orders": []}).status_code == 400              # без версии — нельзя
    new = [{"text": "Минфину России (А.Г. Силуанову) проработать финансирование",
            "assignee": "А.Г. Силуанов", "deadline": "2026-09-01"},
           {"text": "   ", "assignee": "пусто", "deadline": ""}]
    ok = client.put(f"/api/records/{rid}/orders", headers=EMP, json={"orders": new, "version": version})
    assert ok.status_code == 200
    body = ok.json()
    assert body["orders_version"] == version + 1
    assert body["orders"] == [{"text": "Минфину России (А.Г. Силуанову) проработать финансирование",
                               "assignee": "А.Г. Силуанов", "deadline": "01.09.2026",
                               "html": "<p>Минфину России (А.Г. Силуанову) проработать финансирование</p>",
                               "id": ANY}]
    late = client.put(f"/api/records/{rid}/orders", headers=EMP, json={"orders": [], "version": version})
    assert late.status_code == 409
    assert client.get(f"/api/records/{rid}", headers=EMP).json()["orders"] == body["orders"]
    doc = docx_text(client.get(f"/api/records/{rid}/export/orders.docx", headers=EMP))
    assert "Срок – до 1 сентября 2026 г." in doc


def test_pravka_summari_s_versiey(client):
    record = record_after_upload(client)
    rid, version = record["id"], record["summary_version"]
    ok = client.put(f"/api/records/{rid}/summary", headers=EMP,
                    json={"text": "Новое саммари", "version": version})
    assert ok.status_code == 200 and ok.json()["summary"] == "Новое саммари"
    assert client.put(f"/api/records/{rid}/summary", headers=EMP,
                      json={"text": "x", "version": version}).status_code == 409
    assert client.put(f"/api/records/{rid}/summary", headers=EMP, json={"text": "x"}).status_code == 400


def test_pravka_cheloveka_vo_vremya_iskry_ne_zatiraetsya(client, monkeypatch):
    """Перечень поправили, пока «Искра» его составляла, — ответ «Искры» правку не затирает."""
    rid = record_after_upload(client)["id"]
    record = client.get(f"/api/records/{rid}", headers=EMP).json()

    def orders_while_human_edits(text, cancelled=None):
        db.save_orders(rid, [{"text": "Правка человека", "assignee": "", "deadline": ""}],
                       {"login": "emp2", "name": "Петров"}, record["orders_version"])
        return [dict(ORDER)], {"in": 1, "out": 1, "total": 2}

    monkeypatch.setattr(main.orders_mod, "extract", orders_while_human_edits)
    job = finish(client, client.post(f"/api/records/{rid}/rebuild", headers=EMP))
    assert job["status"] == "done"
    after = client.get(f"/api/records/{rid}", headers=EMP).json()
    assert after["orders"] == [{"text": "Правка человека", "assignee": "", "deadline": "",
                                "html": "<p>Правка человека</p>", "id": ANY}]
    assert "не записан" in after["state_note"]


def test_peresborka_po_ispravlennoy_transkripcii(client, monkeypatch):
    record = record_after_upload(client)
    rid = record["id"]
    client.put(f"/api/records/{rid}/text", headers=EMP,
               json={"text": "Исправленный текст про кварц", "version": record["draft_version"]})
    seen = []
    monkeypatch.setattr(main.summarizer, "summarize",
                        lambda text, cancelled=None: (seen.append(text) or "Саммари про кварц",
                                                      {"in": 1, "out": 1, "total": 2}))
    first = client.post(f"/api/records/{rid}/rebuild", headers=EMP)
    job = finish(client, first)
    assert job["status"] == "done"
    assert seen == ["Исправленный текст про кварц"]
    after = client.get(f"/api/records/{rid}", headers=EMP).json()
    assert after["summary"] == "Саммари про кварц"
    assert client.iskra.uploads.__len__() == 1                 # расшифровку не повторяли


def test_peresborka_dvoynoe_nazhatie_odno_zadanie(client, monkeypatch):
    rid = record_after_upload(client)["id"]
    gate = threading.Event()

    def slow_summary(text, cancelled=None):
        gate.wait(5)
        return "ok", {"in": 0, "out": 0, "total": 0}

    monkeypatch.setattr(main.summarizer, "summarize", slow_summary)
    a = client.post(f"/api/records/{rid}/rebuild", headers=EMP)
    b = client.post(f"/api/records/{rid}/rebuild", headers=EMP)
    assert a.status_code == b.status_code == 202
    assert a.json()["id"] == b.json()["id"]
    gate.set()
    assert finish(client, a)["status"] == "done"


# ── в «Совещание» ─────────────────────────────────────────────────────────────

def test_dokumenty_dlya_soveshchaniya(client):
    record = record_after_upload(client)
    rid = record["id"]
    client.put(f"/api/records/{rid}/meta", headers=EMP, json={"meeting_date": "25.05.2026"})
    res = client.get(f"/api/records/{rid}/results", headers=EMP)
    assert res.status_code == 200
    docs = res.json()["documents"]
    assert [d["kind"] for d in docs] == ["transcript", "summary", "decisions"]
    for d in docs:
        assert set(d) == {"kind", "title", "html", "text", "docx_url", "source"}
        assert d["source"] == {"service": "transcriber", "id": str(rid), "url": f"/transcriber/#record={rid}"}
        assert d["docx_url"].startswith(f"/transcriber/api/records/{rid}/export/")
        # адрес «Скачать Word» существует за порталом (nginx снимает /transcriber)
        r = client.get(d["docx_url"][len("/transcriber"):], headers=EMP)
        assert r.status_code == 200 and r.content[:2] == b"PK"
    assert docs[0]["title"] == "Транскрипция совещания 25.05.2026"
    assert "<p>[00:00] <b>Говорящий 1:</b> Добрый день" in docs[0]["html"]
    assert docs[1]["html"] == "<p>Обсудили систему показателей.</p><p>Решили готовить доклад.</p>"
    assert docs[2]["html"].startswith("<ol><li>Минпросвещения России")
    assert "<br>Срок – до 25 июня 2026 г.</li>" in docs[2]["html"]
    assert docs[2]["text"].startswith("1. Минпросвещения России")
    # повтор — те же source.id: «Совещание» заменит, а не добавит
    again = client.get(f"/api/records/{rid}/results", headers=EMP).json()["documents"]
    assert [d["source"]["id"] for d in again] == [str(rid)] * 3


def test_html_ekraniruetsya(client, monkeypatch):
    monkeypatch.setattr(main.summarizer, "summarize",
                        lambda text, cancelled=None: ("<script>alert(1)</script>", {"in": 0, "out": 0, "total": 0}))
    rid = record_after_upload(client)["id"]
    docs = client.get(f"/api/records/{rid}/results", headers=EMP).json()["documents"]
    assert "<script>" not in docs[1]["html"] and "&lt;script&gt;" in docs[1]["html"]


def test_otmetka_soveshchaniya_v_kartochke(client):
    rid = record_after_upload(client)["id"]
    r = client.post(f"/api/records/{rid}/meeting", headers=EMP, json={
        "id": "m-42", "title": "Совещание о показателях", "date": "25.05.2026",
        "url": "https://evil.example/", "kinds": ["transcript", "summary", "decisions", "мусор"]})
    assert r.status_code == 200
    meeting = r.json()["meeting"]
    assert meeting["id"] == "m-42" and meeting["url"] == "/moderator/#meeting=m-42"   # чужой адрес не берём
    assert meeting["kinds"] == ["transcript", "summary", "decisions"]
    assert client.get("/api/records", headers=EMP).json()["records"][0]["meeting"]["id"] == "m-42"
    assert any(a["action"] == "документы записи приложены к совещанию" for a in audit.sent)
    assert client.post(f"/api/records/{rid}/meeting", headers=EMP, json={}).status_code == 400


def test_stranitsa_otdaet_adresa_soseda_tolko_za_portalom(client, monkeypatch):
    html = client.get("/", headers=EMP).text
    assert "/moderator/api/meetings/{id}/results" in html.replace("&quot;", '"')
    assert 'src="static/meeting.js' in html and "Разобрать с Искрой" in html
    assert 'class="dropzone"' in html and "аудио/видео до 500 МБ" in html
    monkeypatch.setattr(config, "DEV", True)
    dev = client.get("/").text
    assert 'data-links="{}"' in dev


# ── база первой версии ───────────────────────────────────────────────────────

def test_migraciya_bazy_pervoy_versii(tmp_path, monkeypatch):
    path = tmp_path / "old.db"
    monkeypatch.setattr(config, "DB_PATH", path)
    conn = sqlite3.connect(path)
    conn.executescript("""
        CREATE TABLE records (id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL,
          state TEXT NOT NULL DEFAULT '', state_note TEXT NOT NULL DEFAULT '',
          duration REAL NOT NULL DEFAULT 0, draft TEXT NOT NULL DEFAULT '',
          draft_version INTEGER NOT NULL DEFAULT 1, summary TEXT NOT NULL DEFAULT '',
          tokens INTEGER NOT NULL DEFAULT 0, engine TEXT NOT NULL DEFAULT '',
          archived INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL, created_by TEXT NOT NULL,
          created_name TEXT NOT NULL DEFAULT '', updated_at TEXT NOT NULL, updated_by TEXT NOT NULL,
          updated_name TEXT NOT NULL DEFAULT '');
    """)
    conn.execute("INSERT INTO records (title, state, draft, summary, created_at, created_by, updated_at,"
                 " updated_by) VALUES ('старая', 'done', ?, 'саммари', '2026-09-27 10:00:00', 'emp',"
                 " '2026-09-27 10:00:00', 'emp')",
                 (json.dumps([{"start": 0, "end": 5, "text": "текст"}, {"start": 0, "end": 0, "text": "правка"}]),))
    conn.commit()
    conn.close()
    db.init()
    db.init()                                                  # повторный запуск ничего не ломает
    record = db.get_record(1)
    assert record["orders"] == [] and record["has_orders"] is False
    assert record["summary_version"] == 1 and record["chair"] == ""
    assert record["draft_text"] == "[00:00] текст\n\nправка"
    db.set_state(1, "ordering")
    assert db.interrupt_stale() == 1                           # новый «рабочий» шаг тоже снимается


# ── ревью RC-2026-10-02-17 ───────────────────────────────────────────────────

def test_otkaz_iskry_bez_teksta_v_zhurnale_i_zametke(client, capsys):
    """Н-2: при отказе «Искры» в журнал, заметку записи и лог идёт только код ситуации."""
    refusal = "К сожалению, я не могу прослушать этот аудиофайл, секретное слово кварц."
    client.iskra.parts = [refusal]
    job = finish(client, upload(client))
    assert job["status"] == "error"
    record = client.get(f"/api/records/{job['record_id']}", headers=EMP).json()
    assert "«Искра» не вернула расшифровку" in record["state_note"]
    printed = capsys.readouterr()
    for place in [job["message"], record["state_note"], printed.out, printed.err] + \
            [r["details"] for r in audit.sent]:
        assert "кварц" not in place and "прослушать" not in place


def test_korotkaya_rech_s_ne_mogu_razbiraetsya(client):
    client.iskra.parts = ["Говорящий 1: Я не могу прийти завтра, не удалось перенести поездку.\n[КОНЕЦ]"]
    record = record_after_upload(client)
    assert "не могу прийти завтра" in record["draft_text"]
    assert record["summary"] and record["state"] == "done"


def test_net_metki_konca_vidno_na_kartochke(client):
    client.iskra.parts = ["[00:00] Говорящий 1: Добрый день."]
    record = record_after_upload(client)
    assert "не отметила конец" in record["state_note"]


# ── повторное ревью RC-2026-10-02-19 ─────────────────────────────────────────

def test_otkaz_s_metkoy_konec_ne_rasshifrovka(client):
    """Н-12: послушный отказ кончается «[КОНЕЦ]» — всё равно отказ, саммари и поручения не тратятся."""
    calls = []
    client.iskra.parts = ["К сожалению, я не могу прослушать аудиофайл.\n[КОНЕЦ]"]
    main.summarizer.summarize, old = (lambda *a, **k: calls.append(1)), main.summarizer.summarize
    try:
        job = finish(client, upload(client))
    finally:
        main.summarizer.summarize = old
    assert job["status"] == "error" and job["message"] == "«Искра» не вернула расшифровку"
    assert calls == []
    record = client.get(f"/api/records/{job['record_id']}", headers=EMP).json()
    assert record["draft_text"] == ""


class _NotJson:
    status = 200

    def read(self):
        return "<html>шлюз: секретная речь кварц</html>".encode()

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


@pytest.mark.parametrize("step", ["summary", "orders"])
def test_ne_json_ot_runs_ne_ostavlyaet_zapis_v_obrabotke(client, monkeypatch, step):
    """Н-14: 200 не в JSON на /api/v1/runs — запись в ошибке с понятным текстом, событие в журнале,
    «Пересобрать» доступно; текста ответа нигде нет."""
    monkeypatch.setattr(main.summarizer, "summarize", REAL_SUMMARIZE if step == "summary" else fake_summary)
    monkeypatch.setattr(main.orders_mod, "extract", REAL_EXTRACT)
    monkeypatch.setattr(summarizer.urllib.request, "urlopen", lambda *a, **k: _NotJson())
    job = finish(client, upload(client))
    assert job["status"] == "error"
    rid = job["record_id"]
    record = client.get(f"/api/records/{rid}", headers=EMP).json()
    assert record["state"] == "error"
    assert "не JSON" in record["state_note"]
    action = "«Искра» не сделала саммари записи" if step == "summary" \
        else "«Искра» не составила перечень поручений записи"
    events = [r for r in audit.sent if r["action"] == action]
    assert len(events) == 1 and "tokens_in" in events[0]
    for place in [job["message"], record["state_note"]] + [r["details"] for r in audit.sent]:
        assert "кварц" not in place and "<html" not in place
    monkeypatch.setattr(main.summarizer, "summarize", fake_summary)
    monkeypatch.setattr(main.orders_mod, "extract", fake_orders)
    again = client.post(f"/api/records/{rid}/rebuild", headers=EMP)
    assert again.status_code == 202
    assert finish(client, again)["status"] == "done"


def test_nepredvidennyy_sboy_ne_ostavlyaet_zapis_v_obrabotke(client, monkeypatch):
    def boom(text, cancelled=None):
        raise KeyError("неожиданно")

    monkeypatch.setattr(main.orders_mod, "extract", boom)
    job = finish(client, upload(client))
    assert job["status"] == "error" and "Пересобрать" in job["message"]
    record = client.get(f"/api/records/{job['record_id']}", headers=EMP).json()
    assert record["state"] == "error"
    assert any(r["action"] == "сбой при подготовке документов записи" for r in audit.sent)
