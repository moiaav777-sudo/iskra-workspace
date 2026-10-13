"""Единый формат блоков (решение пользователя 13.10: «каждый раздел — отдельным блоком для
форматирования, как в модераторе; единый формат»).

Карточка показывает транскрипцию, саммари и поручения блоками общего компонента
(`portal/static/blocks.js`); правка с оформлением приходит HTML. Здесь — серверная сторона:
очистка HTML на сервере всегда, текст выводится из HTML, версия (409), сброс оформления, когда
текст перезаписывает «Искра», Word из HTML, «Совещанию» — сохранённый HTML, копия общего
модуля `richtext_common.py` совпадает с `portal/shared/richtext.py`.
"""
from __future__ import annotations

import io
import json
import sqlite3
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX

import sys

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))   # app/ — и при сборе этого файла первым (rc.sh)
from app import config, db, main, orders, richtext_common  # noqa: E402
from test_three_documents import EMP, ORDER, client, finish, record_after_upload  # noqa: E402,F401

HERE = Path(__file__).resolve().parent
SERVICE = HERE.parent
SHARED = SERVICE.parent / "portal" / "shared" / "richtext.py"

XSS = ('<p onclick="steal()">Текст <b>жирный</b><img src=x onerror=alert(1)>'
       '<a href="javascript:alert(2)">ссылка</a></p><script>alert(3)</script>'
       '<iframe src="https://evil.example"></iframe><style>p{}</style>')


def raw_row(rid: int) -> sqlite3.Row:
    conn = sqlite3.connect(config.DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        return conn.execute("SELECT * FROM records WHERE id = ?", (rid,)).fetchone()
    finally:
        conn.close()


def no_xss(value: str) -> None:
    low = value.lower()
    for bad in ("<script", "onclick", "onerror", "javascript:", "<iframe", "<img", "<style", "alert("):
        assert bad not in low, (bad, value)


# ── копия общего модуля ──────────────────────────────────────────────────────

@pytest.mark.skipif(not SHARED.exists(), reason="portal/shared не рядом (копия одного сервиса)")
def test_kopiya_richtext_sovpadaet_s_obshchim_modulem():
    """Копию раскладывает общий трек; правка копии руками — ошибка (как у orders_docx.py)."""
    assert (SERVICE / "app" / "richtext_common.py").read_bytes() == SHARED.read_bytes()


# ── карточка: HTML для блоков ────────────────────────────────────────────────

def test_kartochka_otdaet_html_iz_teksta(client):
    record = record_after_upload(client)
    assert record["draft_html"].startswith(
        '<p><span class="tc">[00:00]</span> <b>Говорящий 1:</b> Добрый день, коллеги.')
    assert record["summary_html"] == "<p>Обсудили систему показателей.</p><p>Решили готовить доклад.</p>"
    assert record["orders"][1]["html"] == "<p>Рособрнадзору проработать доступ</p>"
    assert record["draft_styled"] is False and record["summary_styled"] is False
    # в базе html нет — он выведен из текста при отдаче
    row = raw_row(record["id"])
    assert row["draft_html"] == "" and row["summary_html"] == ""
    assert "html" not in json.loads(row["orders"])[0]          # ответ «Искры» html не несёт


# ── транскрипция ─────────────────────────────────────────────────────────────

def test_transkripciya_html_ochishchaetsya_tekst_vyvoditsya(client):
    record = record_after_upload(client)
    rid, version = record["id"], record["draft_version"]
    html = ('<p>[00:00] <strong>Говорящий 1:</strong> Добрый <em>день</em>.</p>'
            '<p>[00:45] Говорящий 2: Готовы.</p>' + XSS)
    r = client.put(f"/api/records/{rid}/text", headers=EMP, json={"html": html, "version": version})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["draft_version"] == version + 1
    assert body["draft_styled"] is True
    # абзацы — через пустую строку: Word ищет таймкод и говорящего по абзацам
    assert body["draft_text"] == ("[00:00] Говорящий 1: Добрый день.\n\n[00:45] Говорящий 2: Готовы.\n\n"
                                  "Текст жирныйссылка")
    # таймкод в начале абзаца хранится <span class="tc"> (Quill класс снимает — сервер возвращает)
    assert body["draft_html"].startswith('<p><span class="tc">[00:00]</span> <strong>Говорящий 1:</strong>'
                                         ' Добрый <em>день</em>.</p><p><span class="tc">[00:45]</span> Говорящий 2:')
    no_xss(body["draft_html"])
    no_xss(raw_row(rid)["draft_html"])                        # и в базе ничего опасного
    # старая версия — 409, текст в базе не тронут
    late = client.put(f"/api/records/{rid}/text", headers=EMP, json={"html": "<p>чужое</p>", "version": version})
    assert late.status_code == 409 and "Обновите карточку" in late.json()["message"]
    assert client.get(f"/api/records/{rid}", headers=EMP).json()["draft_text"].startswith("[00:00] Говорящий 1")


def test_transkripciya_prostym_tekstom_kak_ranshe(client):
    record = record_after_upload(client)
    rid = record["id"]
    saved = client.put(f"/api/records/{rid}/text", headers=EMP,
                       json={"html": "<p><b>с оформлением</b></p>", "version": record["draft_version"]}).json()
    assert saved["draft_styled"] is True
    plain = client.put(f"/api/records/{rid}/text", headers=EMP,
                       json={"text": "[00:10] Говорящий 1: простой текст", "version": saved["draft_version"]})
    assert plain.status_code == 200
    body = plain.json()
    assert body["draft_styled"] is False and raw_row(rid)["draft_html"] == ""
    assert body["draft_html"] == '<p><span class="tc">[00:10]</span> <b>Говорящий 1:</b> простой текст</p>'


def test_taymkod_hranitsya_tc_i_seryy_v_word(client):
    """Quill класс `tc` не хранит: правка приходит с таймкодом простым текстом — сервер помечает его
    снова; пометка, пришедшая с браузера, переживает очистку; таймкод не в начале абзаца — текст."""
    record = record_after_upload(client)
    rid = record["id"]
    html = ('<p>[00:00] <strong>Говорящий 1:</strong> Добрый день.</p>'
            '<p><span class="tc">[00:45]</span> Говорящий 2: Да.</p>'
            '<p>Сказано в [01:00] — не таймкод.</p>')
    body = client.put(f"/api/records/{rid}/text", headers=EMP,
                      json={"html": html, "version": record["draft_version"]}).json()
    stored = raw_row(rid)["draft_html"]
    assert stored == ('<p><span class="tc">[00:00]</span> <strong>Говорящий 1:</strong> Добрый день.</p>'
                      '<p><span class="tc">[00:45]</span> Говорящий 2: Да.</p>'
                      '<p>Сказано в [01:00] — не таймкод.</p>')
    assert body["draft_text"].startswith("[00:00] Говорящий 1: Добрый день.\n\n[00:45] Говорящий 2: Да.")
    paras = _doc(client, rid, "transcript").paragraphs[3:]
    assert paras[0].runs[0].text == "[00:00] " and str(paras[0].runs[0].font.color.rgb) == "6B7280"
    assert paras[1].runs[0].text == "[00:45] " and paras[1].runs[1].bold   # как раньше: говорящий полужирным


def test_novaya_rasshifrovka_sbrasyvaet_oformlenie(client):
    record = record_after_upload(client)
    rid = record["id"]
    client.put(f"/api/records/{rid}/text", headers=EMP,
               json={"html": "<p><u>правка</u></p>", "version": record["draft_version"]})
    assert raw_row(rid)["draft_html"]
    db.set_result(rid, [{"start": 1, "end": 2, "text": "Новый текст «Искры»"}], 2, "iskra", state="done")
    after = client.get(f"/api/records/{rid}", headers=EMP).json()
    assert raw_row(rid)["draft_html"] == "" and after["draft_styled"] is False
    assert after["draft_html"] == '<p><span class="tc">[00:01]</span> Новый текст «Искры»</p>'


# ── саммари ──────────────────────────────────────────────────────────────────

def test_summari_html_versiya_i_sbros_pri_peresborke(client, monkeypatch):
    record = record_after_upload(client)
    rid, version = record["id"], record["summary_version"]
    html = '<h2>Итоги</h2><p>Решили <span style="background-color:#ffff00">готовить</span></p><ul><li>пункт</li></ul>' + XSS
    r = client.put(f"/api/records/{rid}/summary", headers=EMP, json={"html": html, "version": version})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["summary"] == "Итоги\nРешили готовить\nпункт\nТекст жирныйссылка"
    assert body["summary_styled"] is True
    assert body["summary_html"].startswith(
        '<h2>Итоги</h2><p>Решили <span style="background-color:#ffff00">готовить</span></p><ul><li>пункт</li></ul>')
    no_xss(raw_row(rid)["summary_html"])
    assert client.put(f"/api/records/{rid}/summary", headers=EMP,
                      json={"html": "<p>x</p>", "version": version}).status_code == 409
    # «Пересобрать» — «Искра» пишет новое саммари, оформление прежней правки сбрасывается
    monkeypatch.setattr(main.summarizer, "summarize",
                        lambda text, cancelled=None: ("Свежее саммари", {"in": 1, "out": 1, "total": 2}))
    job = finish(client, client.post(f"/api/records/{rid}/rebuild", headers=EMP))
    assert job["status"] == "done"
    after = client.get(f"/api/records/{rid}", headers=EMP).json()
    assert after["summary"] == "Свежее саммари"
    assert after["summary_styled"] is False and after["summary_html"] == "<p>Свежее саммари</p>"
    assert raw_row(rid)["summary_html"] == ""


def test_pustoy_html_pustoy_tekst(client):
    record = record_after_upload(client)
    rid = record["id"]
    body = client.put(f"/api/records/{rid}/summary", headers=EMP,
                      json={"html": "<p><br></p>", "version": record["summary_version"]}).json()
    assert body["summary"] == "" and body["summary_styled"] is False and body["summary_html"] == ""


# ── поручения ────────────────────────────────────────────────────────────────

def test_porucheniya_s_html(client):
    record = record_after_upload(client)
    rid, version = record["id"], record["orders_version"]
    items = [
        {"html": '<p>Минфину России — <strong>проработать</strong> финансирование</p>' + XSS,
         "assignee": "А.Г. Силуанов", "deadline": "2026-09-01"},
        {"html": record["orders"][1]["html"], "assignee": "", "deadline": ""},   # не тронутое — как было
        {"html": "<p><br></p>", "assignee": "пусто", "deadline": ""},            # пустое — прочь
    ]
    r = client.put(f"/api/records/{rid}/orders", headers=EMP, json={"orders": items, "version": version})
    assert r.status_code == 200, r.text
    got = r.json()["orders"]
    assert len(got) == 2
    first = got[0]
    assert first["text"] == "Минфину России – проработать финансирование Текст жирныйссылка"   # тире короткое
    assert first["deadline"] == "01.09.2026" and first["assignee"] == "А.Г. Силуанов"
    assert first["html"].startswith("<p>Минфину России – <strong>проработать</strong> финансирование</p>")
    no_xss(first["html"])
    stored = json.loads(raw_row(rid)["orders"])
    no_xss(stored[0]["html"])
    assert got[1]["text"] == "Рособрнадзору проработать доступ"
    # Word перечня — по тексту (оформление в Word не идёт, эталон orders_docx.py не меняем)
    doc = "\n".join(p.text for p in Document(io.BytesIO(
        client.get(f"/api/records/{rid}/export/orders.docx", headers=EMP).content)).paragraphs)
    assert "1. Минфину России – проработать финансирование" in doc
    assert "Срок – до 1 сентября 2026 г." in doc
    assert client.put(f"/api/records/{rid}/orders", headers=EMP,
                      json={"orders": items, "version": version}).status_code == 409


def test_otvet_iskry_html_ne_neset_dlinnyy_html_bez_oformleniya():
    assert orders.check_schema({"orders": [{"text": "Сделать", "html": "<p><b>чужое</b></p>"}]}) == [
        {"text": "Сделать", "assignee": "", "deadline": ""}]
    long_html = "<p>" + "<b>а</b> " * 3000 + "</p>"
    got = orders.from_human([{"html": long_html}])[0]
    assert len(got["text"]) <= orders.MAX_TEXT
    assert got["html"] == richtext_common.from_text(got["text"])    # оформление снято, текст абзацем
    # старый вызов — текстом, как раньше
    assert orders.from_human([{"text": "Текст", "assignee": "", "deadline": ""}]) == [
        {"text": "Текст", "assignee": "", "deadline": ""}]


# ── Word из HTML ─────────────────────────────────────────────────────────────

def _doc(client, rid, kind):
    r = client.get(f"/api/records/{rid}/export/{kind}.docx", headers=EMP)
    assert r.status_code == 200
    return Document(io.BytesIO(r.content))


def test_word_transkripcii_iz_html(client):
    record = record_after_upload(client)
    rid = record["id"]
    html = ('<p><span class="tc">[00:00]</span> <strong>Говорящий 1:</strong> Добрый день.</p>'
            '<p>[00:45] Говорящий 2: Обычный абзац.</p>'
            '<p>[01:30] Говорящий 1: Тут <strong>важное</strong> и <s>лишнее</s>.</p>'
            '<ul><li>пункт списка</li></ul>')
    client.put(f"/api/records/{rid}/text", headers=EMP, json={"html": html, "version": record["draft_version"]})
    paras = _doc(client, rid, "transcript").paragraphs[3:]          # после заголовка, строки записи, пустой
    texts = [p.text for p in paras]
    assert texts == ["[00:00] Говорящий 1: Добрый день.", "[00:45] Говорящий 2: Обычный абзац.",
                     "[01:30] Говорящий 1: Тут важное и лишнее.", "пункт списка"]
    # абзац без оформления — как раньше: таймкод серым, говорящий полужирным
    first = paras[0].runs
    assert first[0].text == "[00:00] " and str(first[0].font.color.rgb) == "6B7280"
    assert first[1].text == "Говорящий 1: " and first[1].bold
    assert [r.text for r in paras[1].runs if r.bold] == ["Говорящий 2: "]
    # абзац с оформлением человека — по кускам
    third = paras[2].runs
    assert [r.text for r in third if r.bold] == ["важное"]
    assert [r.text for r in third if r.font.strike] == ["лишнее"]
    assert paras[3].style.name == "List Bullet"


def test_word_summari_iz_html(client):
    record = record_after_upload(client)
    rid = record["id"]
    html = ('<h1>Итоги — кратко</h1>'
            '<p class="ql-align-center"><em>курсив</em> <u>подчёркнуто</u> '
            '<span style="color:#e60000">красным</span> <span style="background-color:#cce8cc">заливка</span></p>'
            '<ol><li>первый</li><li class="ql-indent-1">вложенный</li></ol>')
    client.put(f"/api/records/{rid}/summary", headers=EMP, json={"html": html, "version": record["summary_version"]})
    paras = _doc(client, rid, "summary").paragraphs[3:]
    assert [p.text for p in paras] == ["Итоги – кратко", "курсив подчёркнуто красным заливка", "первый", "вложенный"]
    head = paras[0].runs[0]
    assert head.bold and head.font.size.pt == 16
    styled = paras[1]
    assert styled.alignment == WD_ALIGN_PARAGRAPH.CENTER
    runs = {r.text: r for r in styled.runs}
    assert runs["курсив"].italic and runs["подчёркнуто"].underline
    assert str(runs["красным"].font.color.rgb) == "E60000"
    assert runs["заливка"].font.highlight_color == WD_COLOR_INDEX.YELLOW
    assert [p.style.name for p in paras[2:]] == ["List Number", "List Number 2"]


def test_word_bez_html_kak_ranshe(client):
    rid = record_after_upload(client)["id"]
    texts = [p.text for p in _doc(client, rid, "summary").paragraphs[3:]]
    assert texts == ["Обсудили систему показателей.", "Решили готовить доклад."]


# ── «Совещанию» ─────────────────────────────────────────────────────────────

def test_sovesh_chaniyu_sohranennyy_html(client):
    record = record_after_upload(client)
    rid = record["id"]
    client.put(f"/api/records/{rid}/text", headers=EMP,
               json={"html": "<p>[00:00] <strong>Говорящий 1:</strong> <em>Добрый</em></p>",
                     "version": record["draft_version"]})
    client.put(f"/api/records/{rid}/summary", headers=EMP,
               json={"html": "<p><u>Подчёркнуто</u></p>", "version": record["summary_version"]})
    client.put(f"/api/records/{rid}/orders", headers=EMP, json={"version": record["orders_version"], "orders": [
        {"html": "<p><strong>Важное</strong> поручение</p>", "assignee": "", "deadline": "25.06.2026"},
        {"html": record["orders"][1]["html"], "assignee": "", "deadline": ""}]})
    docs = {d["kind"]: d for d in client.get(f"/api/records/{rid}/results", headers=EMP).json()["documents"]}
    assert docs["transcript"]["html"] == ('<p><span class="tc">[00:00]</span> <strong>Говорящий 1:</strong>'
                                          ' <em>Добрый</em></p>')
    assert docs["transcript"]["text"] == "[00:00] Говорящий 1: Добрый"
    assert docs["summary"]["html"] == "<p><u>Подчёркнуто</u></p>"
    assert docs["decisions"]["html"] == ("<ol><li><p><strong>Важное</strong> поручение</p><p>Срок – до 25 июня 2026 г.</p></li>"
                                         "<li>Рособрнадзору проработать доступ</li></ol>")


# ── ревью RC-2026-10-13-2 ────────────────────────────────────────────────────

def test_k1_postoyannyy_id_porucheniya(client):
    """К-1: у поручения постоянный id; правка его сохраняет, новому выдаётся, повтор — заменяется."""
    record = record_after_upload(client)
    rid = record["id"]
    first, second = record["orders"]
    assert first["id"] != second["id"]
    assert json.loads(raw_row(rid)["orders"])[0]["id"] == first["id"]          # «Искра» — id записан
    items = [{"id": second["id"], "html": second["html"], "assignee": "", "deadline": ""},
             {"html": "<p>Новое</p>", "assignee": "", "deadline": ""},
             {"id": second["id"], "html": "<p>Повтор id</p>", "assignee": "", "deadline": ""},
             {"id": "<script>", "html": "<p>Плохой id</p>", "assignee": "", "deadline": ""}]
    got = client.put(f"/api/records/{rid}/orders", headers=EMP,
                     json={"orders": items, "version": record["orders_version"]}).json()["orders"]
    ids = [o["id"] for o in got]
    assert ids[0] == second["id"]                    # первое убрали — у второго id прежний
    assert len(set(ids)) == 4 and all(orders.ORDER_ID.match(i) for i in ids)
    assert client.get(f"/api/records/{rid}", headers=EMP).json()["orders"] == got   # id не меняется при чтении


def test_k1_starye_porucheniya_bez_id_po_mestu(client):
    record = record_after_upload(client)
    rid = record["id"]
    conn = sqlite3.connect(config.DB_PATH)
    conn.execute("UPDATE records SET orders = ? WHERE id = ?",
                 (json.dumps([{"text": "А", "assignee": "", "deadline": ""},
                              {"text": "Б", "assignee": "", "deadline": ""}], ensure_ascii=False), rid))
    conn.commit()
    conn.close()
    one = client.get(f"/api/records/{rid}", headers=EMP).json()
    two = client.get(f"/api/records/{rid}", headers=EMP).json()
    assert [o["id"] for o in one["orders"]] == [o["id"] for o in two["orders"]] == ["o1", "o2"]
    # первое сохранение записывает их — дальше они постоянные
    items = [{"id": o["id"], "html": o["html"], "assignee": "", "deadline": ""} for o in one["orders"][1:]]
    client.put(f"/api/records/{rid}/orders", headers=EMP, json={"orders": items, "version": one["orders_version"]})
    assert [o["id"] for o in json.loads(raw_row(rid)["orders"])] == ["o2"]


def _old_code_write(rid: int, **cols) -> None:
    """Так пишет код до 13.10 (после отката): текст меняется, колонки html он не знает."""
    conn = sqlite3.connect(config.DB_PATH)
    for col, value in cols.items():
        conn.execute(f"UPDATE records SET {col} = ? WHERE id = ?", (value, rid))
    conn.commit()
    conn.close()


def test_n1_tekst_izmenen_v_obhod_html_pokazyvaetsya_tekst(client):
    """Н-1: откат → правка на старом коде → снова новый код: карточка, Word и «Совещание» берут
    более новый текст, а не устаревшее оформление."""
    record = record_after_upload(client)
    rid = record["id"]
    client.put(f"/api/records/{rid}/text", headers=EMP,
               json={"html": "<p><b>Правка А</b></p>", "version": record["draft_version"]})
    client.put(f"/api/records/{rid}/summary", headers=EMP,
               json={"html": "<p><u>Саммари А</u></p>", "version": record["summary_version"]})
    rec = client.get(f"/api/records/{rid}", headers=EMP).json()
    client.put(f"/api/records/{rid}/orders", headers=EMP, json={"version": rec["orders_version"], "orders": [
        {"id": rec["orders"][0]["id"], "html": "<p><b>Поручение А</b></p>", "assignee": "", "deadline": ""}]})
    stored_orders = json.loads(raw_row(rid)["orders"])
    stored_orders[0]["text"] = "Поручение Б (старый код)"
    _old_code_write(rid, draft=json.dumps([{"start": 0, "end": 0, "text": "Правка Б (старый код)"}], ensure_ascii=False),
                    summary="Саммари Б (старый код)", orders=json.dumps(stored_orders, ensure_ascii=False))
    after = client.get(f"/api/records/{rid}", headers=EMP).json()
    assert after["draft_styled"] is False and after["draft_html"] == "<p>Правка Б (старый код)</p>"
    assert after["summary_styled"] is False and after["summary_html"] == "<p>Саммари Б (старый код)</p>"
    assert after["orders"][0]["html"] == "<p>Поручение Б (старый код)</p>"
    words = [p.text for p in _doc(client, rid, "transcript").paragraphs]
    assert "Правка Б (старый код)" in words and "Правка А" not in words
    assert "Саммари Б (старый код)" in [p.text for p in _doc(client, rid, "summary").paragraphs]
    docs = {d["kind"]: d for d in client.get(f"/api/records/{rid}/results", headers=EMP).json()["documents"]}
    assert "Правка Б" in docs["transcript"]["html"] and "Саммари Б" in docs["summary"]["html"]
    assert "Поручение Б" in docs["decisions"]["html"] and "Поручение А" not in docs["decisions"]["html"]


def test_n3_predel_html_dokumenta(client):
    record = record_after_upload(client)
    rid = record["id"]
    huge = "<p>" + "<b>ж</b><i>к</i>" * 70_000 + "</p>"            # больше 1 000 000 знаков разметки
    assert len(huge) > main.MAX_DOC_HTML
    for kind, version in (("text", record["draft_version"]), ("summary", record["summary_version"])):
        r = client.put(f"/api/records/{rid}/{kind}", headers=EMP, json={"html": huge, "version": version})
        assert r.status_code == 413 and "слишком большой" in r.json()["message"], r.text
    after = client.get(f"/api/records/{rid}", headers=EMP).json()
    assert after["draft_version"] == record["draft_version"] and after["summary_version"] == record["summary_version"]


def test_n4_predel_html_porucheniya_pri_korotkom_tekste():
    """Текст в пределе, а разметка длиннее MAX_HTML — оформление снимается, текст остаётся."""
    html = "<p>" + "<strong>а</strong><em>б</em>" * 700 + "</p>"
    got = orders.from_human([{"html": html}])[0]
    assert len(got["text"]) == 1400 < orders.MAX_TEXT
    assert len(richtext_common.clean(html)) > orders.MAX_HTML
    assert got["html"] == richtext_common.from_text(got["text"])


MODERATOR_RICHTEXT = SERVICE.parent / "moderator" / "app" / "richtext.py"


@pytest.mark.skipif(not MODERATOR_RICHTEXT.exists(), reason="moderator не рядом")
def test_n4_skvozno_cherez_ochistku_soveshchaniya(client):
    """`/results` → `clean_html` «Совещания»: таймкод `span.tc`, жирный говорящий и поручение
    с абзацами доходят, текст не теряется."""
    import importlib.util
    spec = importlib.util.spec_from_file_location("moderator_richtext", MODERATOR_RICHTEXT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    record = record_after_upload(client)
    rid = record["id"]
    client.put(f"/api/records/{rid}/text", headers=EMP,
               json={"html": "<p>[00:00] <strong>Говорящий 1:</strong> Добрый день.</p>", "version": record["draft_version"]})
    client.put(f"/api/records/{rid}/orders", headers=EMP, json={"version": record["orders_version"], "orders": [
        {"id": record["orders"][0]["id"], "html": "<p><strong>Важное</strong> поручение</p>", "assignee": "",
         "deadline": "25.06.2026"}]})
    docs = {d["kind"]: d for d in client.get(f"/api/records/{rid}/results", headers=EMP).json()["documents"]}
    tr = mod.clean_html(docs["transcript"]["html"])
    assert '<span class="tc">[00:00]</span>' in tr and "<strong>Говорящий 1:</strong>" in tr
    dec = mod.clean_html(docs["decisions"]["html"])
    assert "<strong>Важное</strong> поручение" in dec and "Срок – до 25 июня 2026 г." in dec
