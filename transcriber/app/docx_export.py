"""Выгрузка записи в Word: транскрипция и саммари — по документу на каждое.

Решение пользователя 02.10: после разбора три документа, каждый скачивается в Word.
Транскрипция и саммари — один простой стиль (заголовок, дата записи, абзацы; Times New Roman,
по ширине); перечень поручений — общий модуль `orders_docx.py` по эталону
(`portal/shared/orders_docx.py`).

В транскрипции таймкод «[05:12]» печатается серым, подпись говорящего — полужирным.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor

from . import orders_docx

FONT = "Times New Roman"
_PARA = re.compile(r"^\[(\d{1,2}:\d{2}(?::\d{2})?)\]\s*(.*)$", re.DOTALL)
_SPEAKER = re.compile(
    r"^((?:Говорящий|Говорящая|Спикер|Участник|Участница|Собеседник|Голос)\s*№?\s*\d{1,2}"
    r"|[А-ЯЁ]\.\s?(?:[А-ЯЁ]\.\s?)?[А-ЯЁ][а-яё-]+"
    r"|[А-ЯЁ][а-яё-]+\s+[А-ЯЁ]\.\s?(?:[А-ЯЁ]\.)?"
    r"|[А-ЯЁ][а-яё-]+(?:\s+[А-ЯЁ][а-яё-]+){1,2}):\s+(.+)$", re.DOTALL)


def _stamp(seconds: float) -> str:
    seconds = max(0, int(seconds or 0))
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def record_date(record: dict[str, Any]) -> date | None:
    """Дата совещания из карточки (ДД.ММ.ГГГГ), иначе — день загрузки записи."""
    parsed = orders_docx.parse_deadline(record.get("meeting_date") or "")
    if parsed:
        return parsed
    stamp = str(record.get("created_at") or "")[:10]
    try:
        return datetime.strptime(stamp, "%Y-%m-%d").date()
    except ValueError:
        return None


def date_ru(record: dict[str, Any]) -> str:
    when = record_date(record)
    return f"{when:%d.%m.%Y}" if when else ""


def _doc(heading: str, record: dict[str, Any]):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin, section.right_margin = Mm(25), Mm(15)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(13)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(heading)
    run.bold = True
    run.font.size = Pt(15)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bits = [record.get("title") or "Запись без названия"]
    when = date_ru(record)
    if when:
        bits.append(f"запись от {when}")
    if record.get("duration"):
        bits.append(f"длительность {_stamp(record['duration'])}")
    sub_run = sub.add_run(" · ".join(bits))
    sub_run.italic = True
    sub_run.font.size = Pt(11)
    doc.add_paragraph()
    return doc


def _para(doc, text: str = ""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if text:
        p.add_run(text)
    return p


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def transcript_paragraph(doc, para: str) -> None:
    """«[05:12] Говорящий 1: текст» → таймкод серым, говорящий полужирным, дальше текст."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body = para
    m = _PARA.match(para)
    if m:
        stamp_run = p.add_run(f"[{m.group(1)}] ")
        stamp_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        body = m.group(2)
    sm = _SPEAKER.match(body)
    if sm:
        who = p.add_run(f"{sm.group(1)}: ")
        who.bold = True
        body = sm.group(2)
    p.add_run(body)


def build_transcript(record: dict[str, Any]) -> bytes:
    """Дословная транскрипция («слово в слово») с таймкодами и говорящими."""
    doc = _doc("Транскрипция", record)
    draft = (record.get("draft_text") or "").strip()
    if not draft:
        _para(doc, "Текст ещё не распознан.")
    for para in re.split(r"\n\s*\n", draft):
        if para.strip():
            transcript_paragraph(doc, " ".join(para.split()))
    return _save(doc)


def build_summary(record: dict[str, Any]) -> bytes:
    doc = _doc("Саммари", record)
    summary = orders_docx.dash(record.get("summary") or "").strip()
    if not summary:
        _para(doc, "Саммари ещё нет.")
    for para in re.split(r"\n\s*\n|\n(?=\s*(?:[-–•*]|\d+[.)])\s)", summary):
        if para.strip():
            _para(doc, " ".join(para.split()))
    return _save(doc)


def build(record: dict[str, Any]) -> bytes:
    """Совместимость с первой версией: «запись в Word» — это транскрипция."""
    return build_transcript(record)


def build_orders(record: dict[str, Any]) -> bytes:
    return orders_docx.build(record.get("orders") or [], record.get("chair") or "",
                             record.get("chair_role") or "", record_date(record))


def _clean(title: str) -> str:
    clean = re.sub(r'[\\/:*?"<>|]', " ", (title or "запись").strip())
    clean = re.sub(r"\.(mp3|wav|m4a|aac|ogg|opus|flac|wma|mp4|m4v|mov|webm|mkv|avi)$", "",
                   clean, flags=re.IGNORECASE)
    return " ".join(clean.split())[:120] or "запись"


def file_name(title: str) -> str:
    return f"{_clean(title)}.docx"


def names(record: dict[str, Any]) -> dict[str, str]:
    """Имена трёх файлов: «Транскрипция ДД.ММ.ГГГГ – <запись>», «Саммари ДД.ММ.ГГГГ – <запись>»,
    «Перечень поручений ДД.ММ.ГГГГ» (как в эталоне)."""
    when = date_ru(record)
    tail = f" {when}" if when else ""
    title = _clean(record.get("title") or "")
    return {
        "transcript": f"Транскрипция{tail} – {title}.docx",
        "summary": f"Саммари{tail} – {title}.docx",
        "orders": orders_docx.file_name(record_date(record)),
    }
