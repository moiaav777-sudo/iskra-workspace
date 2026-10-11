"""Снять вёрстку с образца оформления в шаблон `app/template.docx`.

Образец (`docs/Пример оформления.docx`) — рабочий документ пользователя: его текст
в репозиторий не кладётся. В шаблоне остаётся только оформление: размер страницы,
поля, тема и шрифты, список с маркером, стили. Текста в нём нет.

Запуск после замены образца:
    .venv/bin/python tools/make_template.py
"""
from __future__ import annotations

import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

from app import rules  # noqa: E402

TEMPLATE = BASE / "app" / "template.docx"
PPR = BASE / "app" / "template_bullet.xml"


def main() -> int:
    import docx
    from docx.oxml.ns import qn

    samples = rules.samples()
    if not samples:
        print("В docs/ нет образца оформления — шаблон не меняем.")
        return 1
    source = rules.DOCS_DIR / samples[0]
    doc = docx.Document(str(source))

    first = doc.paragraphs[0] if doc.paragraphs else None
    if first is not None and first._p.find(qn("w:pPr")) is not None:
        import copy
        from lxml import etree
        ppr = copy.deepcopy(first._p.find(qn("w:pPr")))
        PPR.write_bytes(etree.tostring(ppr, pretty_print=True))
        print(f"оформление буллита снято → {PPR.name}")

    body = doc.element.body
    for node in list(body):
        if node.tag != qn("w:sectPr"):
            body.remove(node)
    doc.save(str(TEMPLATE))
    _clean_properties(TEMPLATE)
    print(f"шаблон без текста → {TEMPLATE.name} ({TEMPLATE.stat().st_size} байт)")
    return 0


def _clean_properties(path: Path) -> None:
    """Убрать из шаблона свойства чужого документа: автор, время правки, счётчики.

    В репозитории лежит только вёрстка — ни текста, ни следов людей, которые
    готовили образец.
    """
    import re
    import zipfile

    core = (
        "<?xml version='1.0' encoding='UTF-8' standalone='yes'?>\n"
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/'
        'metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/">'
        "<dc:title>Шаблон тезисов</dc:title><dc:creator>Тезисы</dc:creator>"
        "<cp:lastModifiedBy>Тезисы</cp:lastModifiedBy></cp:coreProperties>"
    )
    source = path.read_bytes()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        with zipfile.ZipFile(__import__("io").BytesIO(source)) as src:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == "docProps/core.xml":
                    data = core.encode("utf-8")
                elif item.filename == "docProps/app.xml":
                    text = data.decode("utf-8")
                    for tag in ("Pages", "Words", "Characters", "Lines", "Paragraphs",
                                "CharactersWithSpaces", "TotalTime"):
                        text = re.sub(f"<{tag}>.*?</{tag}>", f"<{tag}>0</{tag}>", text)
                    data = text.encode("utf-8")
                out.writestr(item, data)


if __name__ == "__main__":
    raise SystemExit(main())
