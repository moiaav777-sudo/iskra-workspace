"""Перечень поручений в Word — один вид документа на всё пространство.

Эталон — `moderator/docs/примеры/Перечень поручений 25.05.2026 (пример).docx` и решения
пользователя 02.10 (`TRACK.md` → «Модератор» → задача «Поручения»):

- A4 книжный, поля: верх 20, право 15, низ 20, лево 20 мм;
- Times New Roman 16, межстрочный 1,15, красная строка 12,5 мм, выравнивание по ширине
  (абзацы стиля «No Spacing» — без отступов до и после, как в эталоне);
- шапка по центру полужирно в две строки — должность председателя заглавными
  («ЗАМЕСТИТЕЛЬ ПРЕДСЕДАТЕЛЯ ПРАВИТЕЛЬСТВА / РОССИЙСКОЙ ФЕДЕРАЦИИ»), под ней черта (как в
  эталоне — нижняя граница пустого абзаца) и ещё одна пустая строка: две пустые строки;
- каждое поручение — «N. текст» одним абзацем, следующей строкой «Срок – до 25 июня 2026 г.»
  (дата ДД.ММ.ГГГГ печатается словами); **даты нет — строки «Срок» нет вовсе** (решение
  пользователя 02.10: заготовку с пробелами из эталона не воспроизводим); пустая строка между
  поручениями;
- в конце три пустые строки и справа «И.О. Фамилия» председателя;
- названия совещания и даты в теле нет; имя файла — «Перечень поручений ДД.ММ.ГГГГ.docx»;
- только короткое тире «–».

Должности нет — нет шапки (и пустых строк под ней); председателя нет — нет подписи.

Сознательные отступления от эталона: под шапкой черта и одна пустая строка (в эталоне черта и
две — по постановке «две пустые строки»); пометка «ПРОЕКТ» из колонтитула первой страницы не
воспроизводится.

ЭТОТ ФАЙЛ ПРАВИТСЯ ТОЛЬКО В `portal/shared/orders_docx.py`.
В сервисы он раскладывается скриптом `portal/scripts/sync_shared.sh` — правка копии
будет затёрта. Зависимость одна — `python-docx`.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, Twips

FONT = "Times New Roman"
SIZE = Pt(16)
LINE = 1.15
# в твипах, как в эталоне: 709 = 12,5 мм, 1134 = 20 мм, 851 = 15 мм (Mm() округлил бы 851 до 850)
INDENT = Twips(709)
MARGINS = {"top": Twips(1134), "right": Twips(851), "bottom": Twips(1134), "left": Twips(1134)}
STYLE = "No Spacing"

MONTHS = ("января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
          "сентября", "октября", "ноября", "декабря")

_DASHES = re.compile(r"[—―−]")          # длинное тире, черта, минус → «–»
_LEADING_NUM = re.compile(r"^\s*\d{1,3}\s*[.)]\s+")     # «3. текст» — номер ставим сами
_DMY = re.compile(r"^\s*(\d{1,2})\.(\d{1,2})\.(\d{4})\s*(?:г\.?)?\s*$")
_ISO = re.compile(r"^\s*(\d{4})-(\d{2})-(\d{2})\s*$")
_INITIALS_FIRST = re.compile(r"^[А-ЯЁA-Z]\.\s?(?:[А-ЯЁA-Z]\.\s?)?[А-ЯЁA-Z][\w-]+$")
_SURNAME_FIRST = re.compile(r"^([А-ЯЁA-Z][\w-]+)\s+([А-ЯЁA-Z])\.\s?(?:([А-ЯЁA-Z])\.)?$")


def dash(text: str) -> str:
    """Только короткое тире; «-» между пробелами (как «Срок - 1 декабря» в эталоне) — тоже тире."""
    text = _DASHES.sub("–", str(text or ""))
    return re.sub(r"(?<=\s)-(?=\s)", "–", text)


def parse_deadline(value: Any) -> date | None:
    """ДД.ММ.ГГГГ (или ГГГГ-ММ-ДД из поля даты браузера) → дата; иное — None."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    m = _DMY.match(text)
    if m:
        day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
    else:
        m = _ISO.match(text)
        if not m:
            return None
        year, month, day = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        return date(year, month, day)
    except ValueError:
        return None


def date_words(when: date) -> str:
    """25.06.2026 → «25 июня 2026 г.»"""
    return f"{when.day} {MONTHS[when.month - 1]} {when.year} г."


def deadline_line(value: Any) -> str:
    """Строка срока или пусто, если срока нет. Дата — «Срок – до 25 июня 2026 г.»;
    срок словами («начиная с 1 сентября 2026 г.») печатается как есть."""
    text = dash(str(value or "")).strip()
    if not text:
        return ""
    parsed = parse_deadline(text)
    if parsed:
        return f"Срок – до {date_words(parsed)}"
    text = re.sub(r"^\s*срок\s*[–:]?\s*", "", text, flags=re.IGNORECASE).strip()
    return f"Срок – {text}" if text else ""


def signature(chair: str) -> str:
    """Подпись председателя «И.О. Фамилия»: «Чернышенко Дмитрий Николаевич» и
    «Чернышенко Д.Н.» приводятся к «Д.Н. Чернышенко»; «Д.Н. Чернышенко» — как есть."""
    name = " ".join(dash(chair).split())
    if not name:
        return ""
    if _INITIALS_FIRST.match(name):
        return re.sub(r"\.\s+(?=[А-ЯЁA-Z]\.)", ".", name)
    m = _SURNAME_FIRST.match(name)
    if m:
        return f"{m.group(2)}.{(m.group(3) + '.') if m.group(3) else ''} {m.group(1)}"
    words = name.split(" ")
    if len(words) in (2, 3) and all(w[:1].isupper() and "." not in w for w in words):
        surname, rest = words[0], words[1:]
        return "".join(w[0] + "." for w in rest) + " " + surname
    return name


def role_lines(role: str) -> list[str]:
    """Должность — заглавными, в две строки (как в эталоне).

    Перенос задан руками («\\n» или « / ») — по нему; «… Российской Федерации» — хвост на
    вторую строку; длинная должность — делится по словам примерно пополам; короткая — одна строка.
    """
    text = " ".join(dash(role).replace("\r", "").split(" ")).strip()
    if not text:
        return []
    if "\n" in text or " / " in text:
        parts = [p.strip() for p in re.split(r"\n| / ", text)]
        return [p.upper() for p in parts if p]
    text = " ".join(text.split())
    m = re.match(r"^(.+?)\s+(Российской\s+Федерации)$", text, flags=re.IGNORECASE)
    if m:
        return [m.group(1).upper(), m.group(2).upper()]
    words = text.split(" ")
    if len(text) <= 32 or len(words) < 2:
        return [text.upper()]
    best, cut = None, 1
    for i in range(1, len(words)):
        diff = abs(len(" ".join(words[:i])) - len(" ".join(words[i:])))
        if best is None or diff < best:
            best, cut = diff, i
    return [" ".join(words[:cut]).upper(), " ".join(words[cut:]).upper()]


def clean_orders(orders: Iterable[Any]) -> list[dict[str, str]]:
    """Поручения без текста отбрасываются; номер в начале текста снимается — ставим свой."""
    out: list[dict[str, str]] = []
    for item in orders or []:
        if isinstance(item, str):
            item = {"text": item}
        if not isinstance(item, dict):
            continue
        text = " ".join(dash(str(item.get("text") or "")).split())
        text = _LEADING_NUM.sub("", text).strip()
        if not text:
            continue
        out.append({"text": text,
                    "assignee": " ".join(dash(str(item.get("assignee") or "")).split()),
                    "deadline": " ".join(dash(str(item.get("deadline") or "")).split())})
    return out


# ── сборка документа ───────────────────────────────────────────────────────────

def _font(run) -> None:
    run.font.name = FONT
    run.font.size = SIZE
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def _para(doc, text: str = "", align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent: bool = True,
          bold: bool = False):
    p = doc.add_paragraph(style=STYLE)
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = LINE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = INDENT if indent else None
    if text:
        run = p.add_run(text)
        run.bold = bold or None
        _font(run)
    else:
        # пустая строка той же высоты, что текст: размер шрифта — у знака абзаца
        rpr = p._p.get_or_add_pPr()
        mark = OxmlElement("w:rPr")
        for tag in ("w:sz", "w:szCs"):
            el = OxmlElement(tag)
            el.set(qn("w:val"), str(int(SIZE.pt * 2)))
            mark.append(el)
        rpr.append(mark)
    return p


def _bottom_rule(paragraph) -> None:
    """Черта под шапкой — как в эталоне: нижняя граница пустого абзаца."""
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "12"), ("w:space", "1"), ("w:color", "auto")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    # порядок детей pPr по схеме: pStyle … pBdr … spacing, ind, jc — ставим сразу за pStyle
    style = ppr.find(qn("w:pStyle"))
    if style is not None:
        style.addnext(borders)
    else:
        ppr.insert(0, borders)


def _setup(doc) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = MARGINS["top"]
    section.right_margin = MARGINS["right"]
    section.bottom_margin = MARGINS["bottom"]
    section.left_margin = MARGINS["left"]
    section.header_distance = Twips(709)
    section.footer_distance = Twips(709)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    style = doc.styles[STYLE] if STYLE in [s.name for s in doc.styles] else None
    if style is not None:
        style.font.name = FONT
        style.font.size = SIZE


def build(orders: list[dict], chair: str = "", chair_role: str = "",
          when: date | None = None) -> bytes:
    """Перечень поручений → байты .docx по эталону.

    `orders` — [{"text", "assignee", "deadline"}]: текст идёт в документ как есть (его пишет
    «Искра» или человек в официальной форме «Кому (И.О. Фамилия) … сделать что …»),
    `assignee` в документ не печатается — это подпись для реестра, `deadline` — ДД.ММ.ГГГГ
    (печатается словами), срок словами (печатается как есть) или пусто (строки «Срок» нет).
    `when` в тело не печатается — только в свойства файла (название документа).
    """
    doc = Document()
    _setup(doc)
    body = doc.element.body
    for p in list(body.iterchildren(qn("w:p"))):
        body.remove(p)

    lines = role_lines(chair_role)
    if lines:
        head = _para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
        for i, line in enumerate(lines):
            run = head.add_run(line + (" " if i < len(lines) - 1 else ""))
            run.bold = True
            _font(run)
            if i < len(lines) - 1:
                run.add_break()
        _bottom_rule(_para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False))
        _para(doc, "")

    items = clean_orders(orders)
    for n, order in enumerate(items, start=1):
        if n > 1:
            _para(doc, "")
        _para(doc, f"{n}. {order['text']}")
        term = deadline_line(order["deadline"])
        if term:
            _para(doc, term)

    who = signature(chair)
    if who:
        for _ in range(3):          # три пустые строки перед подписью — как в эталоне (абзацы 27–29)
            _para(doc, "")
        _para(doc, who, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.core_properties.title = file_name(when)[:-len(".docx")]
    doc.core_properties.author = ""
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def file_name(when: date | None = None) -> str:
    """«Перечень поручений ДД.ММ.ГГГГ.docx»; без даты — «Перечень поручений.docx»."""
    if isinstance(when, datetime):
        when = when.date()
    return f"Перечень поручений {when:%d.%m.%Y}.docx" if when else "Перечень поручений.docx"
