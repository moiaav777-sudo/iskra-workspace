"""«00. Реестр материалов.docx» — что лежит в папке мероприятия.

Вёрстка: A4, Times New Roman 16, интервал 1,15. Кегль крупнее, чем в остальных документах
пространства, — реестр читают с листа, держа папку в руках.
Реестр пересобирается после каждой перекладки, поэтому правки руками в нём не живут —
это отчёт о папке, а не документ для редактирования.
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from . import folder, naming, tree

NAME = naming.REGISTRY
FONT = "Times New Roman"
SIZE = 16        # решение пользователя от 20.09: реестр читают распечатанным


def _setup(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(SIZE)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(0)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(2), Cm(1.5)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)


def _para(doc: Document, text: str, *, bold: bool = False, center: bool = False,
          size: int = SIZE, after: int = 0, italic: bool = False, hanging: float = 0):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    if hanging:
        # Название длиннее строки переносится под название, а не под номер:
        # иначе список рассыпается и читается как два разных документа.
        p.paragraph_format.left_indent = Cm(hanging)
        p.paragraph_format.first_line_indent = Cm(-hanging)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def build(event: tree.Event, items: list[folder.Item]) -> Document:
    """Заголовок и нумерованный список — ровно как в папках, которые человек вёл руками.

    Ни подписи о мероприятии, ни служебных пометок: реестр раздают вместе с материалами,
    и лишняя строка в нём — это вопрос, на который придётся отвечать. Чего не хватает,
    видно в карточке мероприятия; в документе, который раздают людям, этому места нет.
    """
    doc = Document()
    _setup(doc)
    _para(doc, "Реестр материалов", bold=True, center=True, after=14)

    listed = [i for i in items if i.parsed.num is not None and i.name != NAME]
    if listed:
        for item in listed:
            title = item.parsed.title
            if item.parsed.question:
                title = f"{title} (к вопросу {item.parsed.question})"
            _para(doc, f"{item.parsed.number}. {title}", hanging=1.4)
    else:
        _para(doc, "Пока пусто: документы ещё не разложены.", italic=True)
    return doc


def write(event: tree.Event, items: list[folder.Item]) -> None:
    """Собрать реестр и положить его в папку. Ошибку наверх не пускаем:
    из-за неудавшегося отчёта не должна падать загрузка документов."""
    try:
        build(event, items).save(event.path / NAME)
    except (OSError, ValueError):
        pass


def refresh(event: tree.Event) -> list[folder.Item]:
    write(event, folder.items(event))
    return folder.items(event)
