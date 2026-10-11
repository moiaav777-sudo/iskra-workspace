"""Путь человека: завёл папку → бросил документы → получил готовый пакет и реестр."""
import io
import zipfile

from conftest import headers
from docx import Document


def _docx(paragraphs: list[str]) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _upload(client, emp, eid, *files):
    payload = [("files", (name, text.encode(), "text/plain")) for name, text in files]
    return client.post(f"/api/events/{eid}/files", headers=emp, files=payload).json()


def test_папка_заводится_по_привычному_пути(client, event):
    assert event["folder"] == "2026 год/09. Сентябрь/02.09.2026 Встреча с Музаевым"


def test_реестр_появляется_сразу(client, emp, event):
    card = client.get(f"/api/events/{event['id']}", headers=emp).json()
    assert [f["name"] for f in card["files"]] == ["00. Реестр материалов.docx"]
    assert card["state"] == {"count": 0, "have": []}


def test_загруженный_документ_раскладывается_сразу(client, emp, event):
    data = _docx(["Приложение № 2", "Информация по докладу Минстроя России", "Далее текст"])
    r = client.post(f"/api/events/{event['id']}/files", headers=emp,
                    files=[("files", ("Инф_справка_ТН.docx", data,
                                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))])
    card = r.json()
    document = [f for f in card["files"] if f["number"] != "00"][0]
    assert document["name"] == "05. Информация по докладу Минстроя России.docx"
    assert document["origin"] == "Инф_справка_ТН.docx"
    # первичка лежит нетронутой
    assert [a["name"] for a in card["archive"]] == ["Инф_справка_ТН.docx"]
    assert card["archive"][0]["in_package"] is True


def test_у_документа_видна_дата_изменения(client, emp, event):
    """Пакет к совещанию читают соседи (вкладка «Пакет») — mtime не хранится
    отдельно, берём его с диска, как размер и имя."""
    data = _docx(["Приложение № 2", "Информация по докладу Минстроя России", "Далее текст"])
    r = client.post(f"/api/events/{event['id']}/files", headers=emp,
                    files=[("files", ("Инф_справка_ТН.docx", data,
                                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))])
    card = r.json()
    document = [f for f in card["files"] if f["number"] != "00"][0]
    assert isinstance(document["mtime"], float) and document["mtime"] > 0
    registry = [f for f in card["files"] if f["number"] == "00"][0]
    assert isinstance(registry["mtime"], float) and registry["mtime"] > 0


def test_один_бриф_встаёт_на_свой_номер(client, emp, event):
    """Номера 01–04 закреплены (решение пользователя 07.10): нет повестки и модератора —
    бриф всё равно 04, а материалы идут с 05."""
    card = _upload(client, emp, event["id"], ("бриф.txt", "Бриф по ЯНАО"),
                   ("сводка.txt", "Сводка по школам"))
    assert [f["name"] for f in card["files"]] == [
        "00. Реестр материалов.docx", "04. Бриф по ЯНАО.txt", "05. Сводка по школам.txt"]
    assert card["state"] == {"count": 2, "have": ["Бриф"]}


def test_пришедшая_повестка_встаёт_на_01_и_номера_не_сдвигаются(client, emp, event):
    eid = event["id"]
    _upload(client, emp, eid, ("бриф.txt", "Бриф по ЯНАО"), ("сводка.txt", "Сводка по школам"))
    card = _upload(client, emp, eid, ("п.txt", "Повестка совещания"))
    # повестка так и называется «Повестка»: что в ней — и так понятно
    assert [f["name"] for f in card["files"]] == [
        "00. Реестр материалов.docx",
        "01. Повестка.txt",
        "04. Бриф по ЯНАО.txt",
        "05. Сводка по школам.txt",
    ]


def test_материалы_идут_по_вопросам_а_общие_впереди(client, emp, event):
    eid = event["id"]
    _upload(client, emp, eid, ("а.txt", "Позиция Минстроя"), ("б.txt", "Справка Рособрнадзора"),
            ("в.txt", "Общие сведения о регионе"))
    def named(card, title):
        return next(f["name"] for f in card["files"] if f["title"] == title)

    card = client.get(f"/api/events/{eid}", headers=emp).json()
    card = client.post(f"/api/events/{eid}/files/{named(card, 'Позиция Минстроя')}/title",
                       headers=emp, json={"title": "Позиция Минстроя", "question": 2}).json()
    card = client.post(f"/api/events/{eid}/files/{named(card, 'Справка Рособрнадзора')}/title",
                       headers=emp, json={"title": "Справка Рособрнадзора", "question": 1}).json()
    assert [f["name"] for f in card["files"]] == [
        "00. Реестр материалов.docx",
        "05. Общие сведения о регионе.txt",
        "06. Справка Рособрнадзора (к вопросу 1).txt",
        "07. Позиция Минстроя (к вопросу 2).txt",
    ]


def test_реестр_перечисляет_документы(client, emp, event):
    eid = event["id"]
    _upload(client, emp, eid, ("сводка.txt", "Сводка по школам"))
    r = client.get(f"/api/events/{eid}/files/00. Реестр материалов.docx/download", headers=emp)
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "Реестр материалов" in text and "05. Сводка по школам" in text
    # в документе только то, что лежит в папке: ни подписи о мероприятии, ни пометок
    assert event["title"] not in text and "Чего не хватает" not in text


def test_поиск_находит_по_тексту_внутри_документа(client, emp, event):
    eid = event["id"]
    _upload(client, emp, eid, ("справка.txt", "Речь о молодых педагогах"))
    items = client.get("/api/search", headers=emp, params={"q": "молодых педагогах"}).json()["items"]
    assert items and items[0]["url"] == f"/materials/#event={eid}"
    assert items[0]["kind"] == "материалы"


def test_убрать_из_пакета_и_вернуть(client, emp, event):
    eid = event["id"]
    _upload(client, emp, eid, ("лишнее.txt", "Лишняя справка"))
    card = client.delete(f"/api/events/{eid}/files/05. Лишняя справка.txt", headers=emp).json()
    assert [f["number"] for f in card["files"]] == ["00"]
    assert card["archive"] == [{"name": "лишнее.txt", "size": 27, "in_package": False}]

    card = client.post(f"/api/events/{eid}/archive/лишнее.txt/restore", headers=emp).json()
    assert [f["name"] for f in card["files"]] == [
        "00. Реестр материалов.docx", "05. Лишняя справка.txt"]
    assert [a["name"] for a in card["archive"]] == ["лишнее.txt"]


def test_выгрузка_отдаёт_пакет_без_первички(client, emp, event):
    eid = event["id"]
    _upload(client, emp, eid, ("Сводка_v7.txt", "Сводка по школам"))
    r = client.get(f"/api/events/{eid}/export/package.zip", headers=emp)
    names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert sorted(names) == ["00. Реестр материалов.docx", "05. Сводка по школам.txt"]


def test_переименование_мероприятия_переносит_папку(client, emp, event):
    r = client.put(f"/api/events/{event['id']}", headers=emp, json={"title": "Встреча с А.А. Музаевым"})
    assert r.json()["event"]["folder"].endswith("02.09.2026 Встреча с А.А. Музаевым")
    assert client.get(f"/api/events/{event['id']}", headers=emp).status_code == 404


def test_дата_мероприятия_меняется_и_папка_переезжает(client, emp, event):
    """Совещание перенесли — папка уходит в свой месяц, пустой месяц убирается."""
    eid = event["id"]
    _upload(client, emp, eid, ("Инф_сводка.txt", "Сводка по школам"))

    r = client.put(f"/api/events/{eid}", headers=emp, json={"date": "2026-10-14"})
    assert r.status_code == 200, r.text
    assert r.json()["event"]["folder"] == "2026 год/10. Октябрь/14.10.2026 Встреча с Музаевым"

    from app import tree
    assert not (tree.root() / "2026 год/09. Сентябрь").exists()
    # исходное имя файла пережило переезд: ключи кэша уехали вместе с папкой
    files = {f["name"]: f for f in r.json()["files"]}
    assert files["05. Сводка по школам.txt"]["origin"] == "Инф_сводка.txt"


def test_наружу_из_хранилища_не_выйти(client, emp, event):
    r = client.get(f"/api/events/{event['id']}/files/..%2F..%2Fsecret.txt/download", headers=emp)
    assert r.status_code == 404


def test_новая_редакция_встаёт_на_место_прежней(client, emp, event):
    eid = event["id"]
    _upload(client, emp, eid, ("б.txt", "Бриф по ЯНАО"), ("Справка.txt", "Справка о школе"))
    # человек поправил название руками — новая редакция его не сбрасывает
    client.post(f"/api/events/{eid}/files/05. Справка о школе.txt/title", headers=emp,
                json={"title": "Справка о школе № 1"})
    card = client.post(f"/api/events/{eid}/files", headers=emp, data={"mode": "replace"},
                       files=[("files", ("Справка.txt", "Справка о школе, уточнённая".encode(),
                                         "text/plain"))]).json()
    assert [f["name"] for f in card["files"]] == [
        "00. Реестр материалов.docx", "04. Бриф по ЯНАО.txt", "05. Справка о школе № 1.txt"]
    # прежняя редакция цела и подписана датой, новая лежит под исходным именем
    from datetime import date
    stamp = date.today().strftime("%d.%m.%Y")
    assert sorted(a["name"] for a in card["archive"]) == [
        f"Справка (версия от {stamp}).txt", "Справка.txt", "б.txt"]

    r = client.get(f"/api/events/{eid}/files/05. Справка о школе № 1.txt/download", headers=emp)
    assert r.text == "Справка о школе, уточнённая"


def test_без_замены_приходит_вторым_документом(client, emp, event):
    eid = event["id"]
    _upload(client, emp, eid, ("Справка.txt", "Справка о школе"))
    card = _upload(client, emp, eid, ("Справка.txt", "Справка о школе"))
    assert len(card["archive"]) == 2
    assert len([f for f in card["files"] if f["number"] != "00"]) == 2


def test_повестка_узнаётся_по_имени_файла(client, emp, event):
    """Повестка часто начинается сразу с первого вопроса — слова «повестка» в ней нет."""
    eid = event["id"]
    _upload(client, emp, eid, ("сводка.txt", "Сводка по школам"))
    card = _upload(client, emp, eid,
                   ("Повестка.txt", "Создание эталонной системы показателей"))
    assert [f["name"] for f in card["files"]] == [
        "00. Реестр материалов.docx",
        "01. Повестка.txt",
        "05. Сводка по школам.txt",
    ]
    assert card["state"]["have"] == ["Повестка"]


def test_вид_узнаётся_по_первым_строкам_документа(client, emp, event):
    eid = event["id"]
    _upload(client, emp, eid, ("сводка.txt", "Сводка по школам"))
    card = _upload(client, emp, eid, ("док.txt", "Бриф\nСистема образования округа"))
    # жанр из первых строк встаёт впереди сути
    assert [f["name"] for f in card["files"]][1] == "04. Бриф – Система образования округа.txt"
    assert card["state"]["have"] == ["Бриф"]


def test_адресат_письма_не_становится_названием(client, emp, event):
    """В письме сверху штамп «кому», и заголовок идёт после него."""
    eid = event["id"]
    card = _upload(client, emp, eid, ("письмо.txt",
                                      "Заместителю Председателя Правительства\n"
                                      "Российской Федерации\n"
                                      "Д.Н. Чернышенко\n"
                                      "О развитии среднего профессионального образования"))
    assert [f["name"] for f in card["files"]][1] == \
        "05. О развитии среднего профессионального образования.txt"


def test_пересобрать_реестр_заодно_наводит_порядок(client, emp, event):
    """Папку могли поправить руками или поменялись правила — одна кнопка чинит обе вещи."""
    eid = event["id"]
    _upload(client, emp, eid, ("сводка.txt", "Сводка по школам"))
    from app import tree
    folder = tree.root() / event["folder"]
    (folder / "07. Повестка.txt").write_text("Повестка совещания", encoding="utf-8")

    card = client.post(f"/api/events/{eid}/registry", headers=emp).json()
    assert [f["name"] for f in card["files"]] == [
        "00. Реестр материалов.docx", "01. Повестка.txt", "05. Сводка по школам.txt"]


def test_реестр_набран_шестнадцатым_кеглем(client, emp, event):
    """Решение пользователя от 20.09: реестр читают с листа."""
    eid = event["id"]
    _upload(client, emp, eid, ("сводка.txt", "Сводка по школам"))
    r = client.get(f"/api/events/{eid}/files/00. Реестр материалов.docx/download", headers=emp)
    doc = Document(io.BytesIO(r.content))
    assert doc.styles["Normal"].font.size.pt == 16
    sizes = {run.font.size.pt for p in doc.paragraphs for run in p.runs if run.font.size}
    assert sizes == {16}


def test_настоящее_название_не_подменяется_жанром(client, emp, event):
    """Слово «участник» внутри фразы не делает документ списком участников."""
    eid = event["id"]
    card = _upload(client, emp, eid,
                   ("а.txt", "Информация об участниках долевого строительства"),
                   ("б.txt", "Материалы к брифу Министра"))
    assert [f["title"] for f in card["files"]][1:] == [
        "Информация об участниках долевого строительства", "Материалы к брифу Министра"]
    assert card["state"]["have"] == []


def test_переезд_папки_с_подчёркиванием_не_ломает_кэш(client, emp):
    """«_» в LIKE — подстановочный знак: соседняя папка переписывалась чужими ключами."""
    made = []
    for title in ("План_А", "План-А"):
        r = client.post("/api/events", headers=headers(), json={"date": "2026-09-02", "title": title})
        made.append(r.json()["event"]["id"])
        client.post(f"/api/events/{made[-1]}/files", headers=headers(),
                    files=[("files", (f"{title}.txt", title.encode(), "text/plain"))])
    r = client.put(f"/api/events/{made[0]}", headers=headers(), json={"date": "2026-10-14"})
    assert r.status_code == 200, r.text
    # у соседней папки всё на месте
    card = client.get(f"/api/events/{made[1]}", headers=headers()).json()
    assert [f["name"] for f in card["files"]][1] == "05. План-А.txt"


def test_поиск_не_находит_сам_реестр(client, emp, event):
    eid = event["id"]
    _upload(client, emp, eid, ("а.txt", "Позиция Минстроя России"))
    items = client.get("/api/search", headers=emp, params={"q": "Позиция"}).json()["items"]
    assert items and "Реестр материалов" not in items[0]["title"]


def test_мероприятие_убирается_и_папка_цела(client, emp, event):
    """Завести по ошибке легко, а в папке бывают единственные экземпляры документов."""
    eid = event["id"]
    _upload(client, emp, eid, ("сводка.txt", "Сводка по школам"))
    assert client.post(f"/api/events/{eid}/drop", headers=emp).status_code == 200

    assert client.get(f"/api/events/{eid}", headers=emp).status_code == 404
    assert client.get("/api/events", headers=emp).json()["items"] == []

    from app import tree
    убрано = tree.root() / "Убранное" / "02.09.2026 Встреча с Музаевым"
    assert (убрано / "05. Сводка по школам.txt").is_file()
    assert (убрано / "Архив" / "сводка.txt").is_file()
    # пустой месяц из дерева убрался
    assert not (tree.root() / "2026 год" / "09. Сентябрь").exists()


def test_стереть_мероприятие_может_только_админ(client, emp, admin, event):
    eid = event["id"]
    _upload(client, emp, eid, ("сводка.txt", "Сводка по школам"))
    assert client.delete(f"/api/events/{eid}", headers=emp).status_code == 403
    assert client.delete(f"/api/events/{eid}", headers=admin).status_code == 200

    from app import tree
    assert not (tree.root() / "2026 год" / "09. Сентябрь").exists()
    assert client.get("/api/events", headers=emp).json()["items"] == []


def test_повестка_приходит_второй_раз_и_заменяет_прежнюю(client, emp, event):
    """Сервис прислал свой документ снова — это новая редакция, а не второй документ."""
    eid = event["id"]
    _upload(client, emp, eid, ("Повестка.txt", "Повестка совещания"),
            ("сводка.txt", "Сводка по школам"))
    card = _upload(client, emp, eid, ("Повестка.txt", "Повестка совещания, уточнённая"))
    assert [f["name"] for f in card["files"]] == [
        "00. Реестр материалов.docx", "01. Повестка.txt", "05. Сводка по школам.txt"]

    got = client.get(f"/api/events/{eid}/files/01. Повестка.txt/download", headers=emp)
    assert got.text == "Повестка совещания, уточнённая"
    # прежняя редакция цела и подписана датой
    from datetime import date
    stamp = date.today().strftime("%d.%m.%Y")
    assert f"01. Повестка (версия от {stamp}).txt" in [a["name"] for a in card["archive"]]


def test_другой_формат_того_же_жанра_не_портит_файл(client, emp, event):
    """Ревьюер RC-2026-09-28-6, К1: docx от «Повестки», следом html от «Совещания» —
    расширение рабочей копии должно поехать за новым форматом, а не остаться прежним,
    иначе файл выглядит как .docx, а внутри лежит html (открыть в Word нельзя)."""
    eid = event["id"]
    docx = _docx(["Повестка совещания", "1. Первый вопрос"])
    client.post(f"/api/events/{eid}/files", headers=emp,
               files=[("files", ("Повестка_14.10.2026.docx", docx,
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))])
    from app import tree
    folder = tree.root() / event["folder"]
    assert (folder / "01. Повестка.docx").is_file()

    html = "<html><body>Повестка дня</body></html>".encode()
    card = client.post(f"/api/events/{eid}/files", headers=emp,
                       files=[("files", ("Повестка.html", html, "text/html"))]).json()

    assert [f["name"] for f in card["files"]][1:] == ["01. Повестка.html"]
    assert not (folder / "01. Повестка.docx").is_file()          # старое расширение не осталось
    assert (folder / "01. Повестка.html").read_bytes() == html    # это настоящий html, не обёртка

    from datetime import date
    stamp = date.today().strftime("%d.%m.%Y")
    kept = folder / "Архив" / f"01. Повестка (версия от {stamp}).docx"
    assert kept.is_file()
    Document(io.BytesIO(kept.read_bytes()))                       # прежний docx остался docx'ом


def test_замена_с_другим_расширением_у_первички_тоже_не_портит_файл(client, emp, event):
    """`mode=replace` находит прежний документ по имени первички: если у файла без
    расширения формат содержимого сменился (был pdf, стал docx), рабочая копия должна
    получить новое расширение, а не принять чужие байты под старым (К1, RC-2026-09-28-6)."""
    eid = event["id"]
    pdf = b"%PDF-1.4\n%..."
    client.post(f"/api/events/{eid}/files", headers=emp,
               files=[("files", ("Документ", pdf, "application/octet-stream"))])
    from app import tree
    folder = tree.root() / event["folder"]
    assert (folder / "05. Документ.pdf").is_file()

    docx = _docx(["Уточнённая версия документа"])
    card = client.post(f"/api/events/{eid}/files", headers=emp, data={"mode": "replace"},
                       files=[("files", ("Документ", docx,
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))]).json()

    names = [f["name"] for f in card["files"] if f["number"] != "00"]
    assert names == ["05. Документ.docx"]
    assert not (folder / "05. Документ.pdf").is_file()
    Document(io.BytesIO((folder / "05. Документ.docx").read_bytes()))


def test_обычный_материал_второй_раз_ложится_рядом(client, emp, event):
    """Осторожность касается только повестки, участников, модератора и брифа."""
    eid = event["id"]
    _upload(client, emp, eid, ("а.txt", "Сводка по школам"))
    card = _upload(client, emp, eid, ("б.txt", "Позиция Минстроя"))
    assert len([f for f in card["files"] if f["number"] != "00"]) == 2


# ── закреплённые номера 01–04 (решение пользователя 07.10) ───────────────

def _numbered(card) -> dict[str, str]:
    """Номер → название: так короче сверять, что где стоит."""
    return {f["number"]: f["title"] for f in card["files"] if f["number"] != "00"}


def test_модератор_всегда_встаёт_третьим(client, emp, event):
    """«Модератор всегда должен вставать п. 3»: повестки и участников ещё нет, а модератор
    уже 03, материал — 05; пришла повестка — встала на 01, номера материалов не сдвинулись."""
    eid = event["id"]
    card = _upload(client, emp, eid, ("Модератор_И61.txt", "Модератор\nСценарий совещания"),
                   ("Тезисы.txt", "Тезисы доклада Минпросвещения России"))
    numbers = _numbered(card)
    assert sorted(numbers) == ["03", "05"]
    assert numbers["03"].startswith("Модератор")
    assert numbers["05"] == "Тезисы доклада Минпросвещения России"
    before = {f["number"]: f["name"] for f in card["files"]}

    card = _upload(client, emp, eid, ("п.txt", "Повестка совещания"))
    after = {f["number"]: f["name"] for f in card["files"]}
    assert after["01"] == "01. Повестка.txt"
    assert after["03"] == before["03"] and after["05"] == before["05"]
    assert sorted(after) == ["00", "01", "03", "05"]
    # реестр перечисляет по номерам, пустые 02 и 04 не выдумывает
    r = client.get(f"/api/events/{eid}/files/00. Реестр материалов.docx/download", headers=emp)
    lines = [p.text for p in Document(io.BytesIO(r.content)).paragraphs][1:]
    assert [line[:3] for line in lines] == ["01.", "03.", "05."]


def test_файл_без_номера_встаёт_на_закреплённый_номер(client, emp, event):
    """Положили мимо сервиса — «Привести в порядок» ставит модератора на 03, а не в хвост."""
    eid = event["id"]
    _upload(client, emp, eid, ("сводка.txt", "Сводка по школам"))
    from app import tree
    folder = tree.root() / event["folder"]
    (folder / "Модератор_И61.txt").write_text("Модератор совещания", encoding="utf-8")
    card = client.post(f"/api/events/{eid}/registry", headers=emp).json()
    assert [f["name"] for f in card["files"]] == [
        "00. Реестр материалов.docx", "03. Модератор совещания.txt", "05. Сводка по школам.txt"]


def test_второй_документ_одной_ступени_уходит_к_материалам(client, emp, event):
    """Двух модераторов быть не должно, но если человек положил второго руками — 03 один."""
    eid = event["id"]
    _upload(client, emp, eid, ("сводка.txt", "Сводка по школам"))
    from app import tree
    folder = tree.root() / event["folder"]
    (folder / "03. Модератор.txt").write_text("Модератор", encoding="utf-8")
    (folder / "Модератор запасной.txt").write_text("Модератор, запасной", encoding="utf-8")
    card = client.post(f"/api/events/{eid}/registry", headers=emp).json()
    assert [f["name"] for f in card["files"]] == [
        "00. Реестр материалов.docx", "03. Модератор.txt", "05. Модератор, запасной.txt",
        "06. Сводка по школам.txt"]


def test_замена_после_новой_редакции_не_плодит_второй_документ(client, emp, event):
    """07.10: модератор загружен дважды (вторая загрузка — новая редакция, первичка стала
    «Модератор_И61 (2)»), третья — `mode=replace`. Прежний документ по имени первички
    не находился, и рядом ложился второй модератор вместо замены."""
    eid = event["id"]
    post = lambda text, **kw: client.post(  # noqa: E731
        f"/api/events/{eid}/files", headers=emp, **kw,
        files=[("files", ("Модератор_И61.txt", text.encode(), "text/plain"))]).json()
    post("Модератор\nРедакция первая")
    card = post("Модератор\nРедакция вторая")
    assert "Модератор_И61 (2).txt" in [a["name"] for a in card["archive"]]
    card = post("Модератор\nРедакция третья", data={"mode": "replace"})

    docs = [f for f in card["files"] if f["number"] != "00"]
    assert len(docs) == 1, [f["name"] for f in docs]
    assert docs[0]["number"] == "03"
    got = client.get(f"/api/events/{eid}/files/{docs[0]['name']}/download", headers=emp)
    assert got.text == "Модератор\nРедакция третья"


def test_длинное_название_не_срывает_пересчёт(client, emp, event, monkeypatch):
    """Имя на сервере — не больше 255 байт, кириллица — два байта на букву. Временная
    приставка «~перекладка~» к длинному имени выходила за предел, первый проход падал,
    и пересчёт откатывался целиком — молча, а новый документ оставался без номера."""
    from app import tree
    # предел пути Windows (240 знаков) тут ни при чём: временная папка теста длинная сама
    monkeypatch.setattr(tree, "MAX_PATH", 1000)
    eid = event["id"]
    long = ("Тезисы доклада заместителя Министра просвещения Российской Федерации о ходе "
            "реализации программы капитального ремонта общеобразовательных организаций")
    card = _upload(client, emp, eid, ("тезисы.txt", long))
    name = next(f["name"] for f in card["files"] if f["number"] == "05")
    card = client.post(f"/api/events/{eid}/files/{name}/title", headers=emp,
                       json={"title": long, "question": 12}).json()
    name = next(f["name"] for f in card["files"] if f["number"] == "05")
    assert len(name.encode()) <= 255

    # материал к вопросу 1 встаёт впереди — длинному надо переехать с 05 на 06
    card = _upload(client, emp, eid, ("справка.txt", "Справка о школах"))
    short = next(f["name"] for f in card["files"] if f["title"] == "Справка о школах")
    card = client.post(f"/api/events/{eid}/files/{short}/title", headers=emp,
                       json={"title": "Справка о школах", "question": 1}).json()
    card = _upload(client, emp, eid, ("Модератор_И61.txt", "Модератор совещания"))
    names = [f["name"] for f in card["files"]]
    assert [n[:3] for n in names] == ["00.", "03.", "05.", "06."], names
    assert names[2] == "05. Справка о школах (к вопросу 1).txt"
    assert names[3].startswith("06. Тезисы доклада") and names[3].endswith("(к вопросу 12).txt")


def test_оборванная_перекладка_возвращает_документ(client, emp, event):
    """Процесс остановили между двумя проходами — документ остался в «~перекладка~/».
    Он настоящий и единственный: следующий пересчёт возвращает его в пакет с номером."""
    eid = event["id"]
    _upload(client, emp, eid, ("сводка.txt", "Сводка по школам"))
    from app import tree
    folder = tree.root() / event["folder"]
    (folder / "~перекладка~").mkdir()
    (folder / "~перекладка~" / "Модератор_И61.txt").write_text("Модератор", encoding="utf-8")
    card = client.post(f"/api/events/{eid}/registry", headers=emp).json()
    assert [f["name"] for f in card["files"]] == [
        "00. Реестр материалов.docx", "03. Модератор И61.txt", "05. Сводка по школам.txt"]
    assert not (folder / "~перекладка~").exists()


def test_файл_унесли_во_время_пересчёта_остальные_номера_ставятся(client, emp, event, monkeypatch):
    """Два пересчёта одной папки разом (загрузка и правка названия): второй не находил
    файл, который уже унёс первый, — и откатывал всё, новый документ оставался без номера."""
    eid = event["id"]
    _upload(client, emp, eid, ("а.txt", "Сводка по школам"), ("б.txt", "Позиция Минстроя"))
    from app import folder as folder_mod, tree
    path = tree.root() / event["folder"]
    (path / "Модератор_И61.txt").write_text("Модератор", encoding="utf-8")
    (path / "Справка.txt").write_text("Справка о школах", encoding="utf-8")
    real_inside = tree.inside

    def gone(p):
        # между проверкой и переименованием файл унёс соседний запрос
        if p.name == "Справка.txt" and p.exists():
            p.rename(path / "Справка (унесли).txt")
        return real_inside(p)

    monkeypatch.setattr(tree, "inside", gone)
    event_obj = tree.find(eid)
    done = {m.src: m.dst for m in folder_mod.renumber(event_obj)}
    assert done.get("Модератор_И61.txt") == "03. Модератор И61.txt"
