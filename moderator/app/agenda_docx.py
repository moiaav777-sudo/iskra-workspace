"""Word-экспорт «Повестки» и «Списка участников» по утверждённому пользователем образцу
(решение пользователя 18.09/28.09, перенос из сервиса «Повестка» — `agenda/app/docx_export.py`;
тот сервис со временем закроют, детали в `../CLAUDE.md` и `../TRACK.md` → «Модератор»).

Вёрстка переносится **как есть, без переосмысления** (`agenda/CLAUDE.md`: «вёрстку без слова
пользователя не менять»):

- A4 книжная, Times New Roman 14 пт, межстрочный 1,15 в повестке и одинарный в списке
  участников, номер страницы в верхнем колонтитуле по центру 12 пт, на первой странице
  номера нет;
- шапка: заголовок и подзаголовок по центру жирным, ниже справа дата, время и место;
- люди — в таблице без границ: ФАМИЛИЯ заглавными и «Имя Отчество» двумя строками в одной
  ячейке, посередине тире, справа должность;
- отбивка строки 6 пт сверху и снизу — и в повестке, и в списке.

**Отличия переноса от `agenda`, вызванные разницей моделей** (не вёрстки):
- у `agenda` докладчик — объект `Speaker` со своим `kind` (main/co); у «Модератора» докладчики
  лежат параллельными списками без этого признака. Порядок и так означает то же самое везде
  в интерфейсе (`static/app.js`: первый — докладчик, остальные — содокладчики,
  `AgendaImportItemIn.speaker_kinds` тоже документирует это как решённое: «у нас деления нет,
  порядок уже верный») — здесь тот же приём: докладчик = первый в списке, остальные = содокладчики;
- председателя с должностью «Модератор» раньше не хранил (в `Meeting.chair` только имя) —
  добавлено `Meeting.chair_role`, печатается в блоке «Вступительное слово» так же, как любой
  докладчик (ФИО – должность);
- организация участника в Word не печатается (решение пользователя 07.10: «пишется только должность,
  принадлежность к организации через запятую не указываем»); до 07.10 она дописывалась к должности
  через запятую. В ячейке — только должность; если должности нет совсем, печатается организация,
  чтобы строка не осталась пустой.
"""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH as ALIGN
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .models import Meeting, Participant, _is_closing, _is_opening

FONT = "Times New Roman"
SIZE = 14
PAGE_NUMBER_SIZE = 12
DASH = "–"                     # тире между человеком и должностью, как в образце

MARGINS = dict(left=2.0, right=1.5, top=2.0, bottom=2.0)

AGENDA_COLS = (6.5, 1.2, 9.8)              # ФИО | тире | должность
PARTICIPANT_COLS = (1.0, 6.2, 0.8, 9.5)    # номер | ФИО | тире | должность

LINE_AGENDA = 1.15
LINE_LIST = 1.0
ROW_SPACE = 6                              # отбивка строки сверху и снизу, пт — одна на оба документа
ITEM_GAP = 16                              # отбивка перед вопросом

FIRST_LINE_ITEM = 1.5                      # красная строка у пункта повестки
FIRST_LINE_BLOCK = 1.25                    # у «Вступительного слова» и «Подведения итогов»
INDENT_SPEAKER_TITLE = 1.1                 # отступ подзаголовка «Докладчики:»

SPEAKER_TITLES = {"main": ("Докладчик:", "Докладчики:"), "co": ("Содокладчик:", "Содокладчики:")}
MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня",
          "июля", "августа", "сентября", "октября", "ноября", "декабря"]

# несверенный документ помечен в нём самом (../CLAUDE.md п. 5а) — тот же приём и цвет, что
# у документа ведущего (app/docx_export.py)
UNCHECKED_MARK = "ФИО не сверены с реестром людей"
RED = RGBColor(0xC0, 0x00, 0x00)


def _mark_unchecked(doc) -> None:
    section = doc.sections[0]
    for footer in (section.footer, section.first_page_footer):
        p = footer.add_paragraph()
        p.alignment = ALIGN.CENTER
        run = p.add_run(UNCHECKED_MARK)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RED
    first = doc.paragraphs[0].insert_paragraph_before()
    run = first.add_run(UNCHECKED_MARK + " — проверьте фамилии перед выгрузкой")
    run.bold = True
    run.font.name = FONT
    run.font.size = Pt(SIZE)
    run.font.color.rgb = RED


def split_name(full: str) -> tuple[str, str]:
    """«Иванов Иван Иванович» → («ИВАНОВ», «Иван Иванович»)."""
    parts = (full or "").strip().split()
    if not parts:
        return "", ""
    return parts[0].upper(), " ".join(parts[1:])


class _Row:
    """ФИО + должность одной строкой — общий вид для докладчика (параллельные списки
    `AgendaItem`) и участника (`Participant` подходит и так, у него уже есть `.name`/`.role`)."""

    __slots__ = ("name", "role")

    def __init__(self, name: str, role: str = "") -> None:
        self.name = name
        self.role = role


def _date_line(m: Meeting) -> str:
    """«7 сентября 2026 г., 16:00». Пустые части опускаются. Дата хранится ДД.ММ.ГГГГ."""
    text = ""
    match = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})$", (m.date or "").strip())
    if match:
        d, mo, y = match.groups()
        try:
            text = f"{int(d)} {MONTHS[int(mo) - 1]} {y} г."
        except IndexError:
            text = m.date
    elif m.date:
        text = m.date
    if m.time_start:
        text = f"{text}, {m.time_start}" if text else m.time_start
    return text


def _place_lines(m: Meeting) -> list[str]:
    return [x for x in (m.location_primary, m.location_secondary) if (x or "").strip()]


def _role_cell(role: str, organization: str = "") -> str:
    """Только должность (решение пользователя 07.10: организацию через запятую в Word не указываем;
    до того она дописывалась). Нет должности — организация, чтобы ячейка не была пустой."""
    role = (role or "").strip()
    return role or (organization or "").strip()


# ── подзаголовок «СПИСОК участников …» / «ПОВЕСТКА …» в родительном падеже ────────
# Пользователь 07.10: в Word вышло «СПИСОК / участников Совещание» — подзаголовок карточки стоял в
# именительном с заглавной. Под словом «СПИСОК» и «ПОВЕСТКА» нужен родительный: «участников совещания …».
_GENITIVE_HEAD = {"совещание": "совещания", "заседание": "заседания", "встреча": "встречи",
                  "совещания": "совещания", "заседания": "заседания", "встречи": "встречи",
                  "совещанию": "совещания", "заседанию": "заседания", "встречу": "встречи"}
_ADJ_GEN = {"neuter": [("ого", "ого"), ("его", "его"), ("ое", "ого"), ("ее", "его"), ("ый", "ого"), ("ий", "его"),
                       ("ом", "ого"), ("ем", "его"), ("ому", "ого"), ("ему", "его")],
            "fem": [("ой", "ой"), ("ей", "ей"), ("ая", "ой"), ("яя", "ей"), ("ую", "ой"), ("юю", "ей")]}


_HUSHING_GEN = set("жшчщ")   # после шипящих женский родительный «-ей»: «рабочая» → «рабочей»; но «краткая» → «краткой»


def _adj_genitive(word: str, gender: str) -> str:
    low = word.lower()
    for tail, gen in _ADJ_GEN[gender]:
        if low.endswith(tail):
            stem = low[: -len(tail)]
            if gen == "ой" and stem[-1:] in _HUSHING_GEN:
                gen = "ей"
            return stem + gen
    return low


def event_genitive(subtitle: str) -> str:
    """«Совещание у Министра …» → «совещания у Министра …»; пусто → «совещания».

    Склоняются только вид мероприятия и одно-два определения перед ним (как в `models.full_title_of`,
    только в обратную сторону); остальной текст не трогаем. Подзаголовок, не начинающийся с вида
    мероприятия («в режиме ВКС»), возвращается как есть со строчной буквы."""
    text = " ".join((subtitle or "").split())
    if not text:
        return "совещания"
    words = text.split(" ")

    def _split(w: str) -> tuple[str, str]:          # «совещание,» → («совещание», ","); знак препинания сохраняется
        core = w.rstrip(",.;:!?)»")
        return core, w[len(core):]

    head_at = next((i for i, w in enumerate(words[:3]) if _split(w)[0].lower().split("-", 1)[0] in _GENITIVE_HEAD), None)
    if head_at is None:
        # аббревиатуру в начале («ВКС-совещание», «VIP-встреча») не трогаем — строчная только у обычного слова
        first = words[0].split("-", 1)[0]
        return text if (len(first) > 1 and first.isupper()) else text[:1].lower() + text[1:]
    core, punct = _split(words[head_at])
    head, dash, tail = core.partition("-")
    gender = "fem" if head.lower().startswith("встреч") else "neuter"
    words[head_at] = _GENITIVE_HEAD[head.lower()] + (dash + tail if dash else "") + punct
    for i in range(head_at):
        words[i] = _adj_genitive(words[i], gender)
    return " ".join(words)


# ── низкоуровневые помощники (общая docx-плотность, без привязки к модели) ─

def _fmt(p, text="", *, bold=False, italic=False, underline=False, align=None,
         first_line=None, left=None, right=None, size=SIZE, space_before=0,
         space_after=0, line=LINE_AGENDA):
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    if left is not None:
        pf.left_indent = Cm(left)
    if right is not None:
        pf.right_indent = Cm(right)
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if text:
        run = p.add_run(text)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if underline:
            run.underline = True
        run.font.size = Pt(size)
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return p


def _p(doc, text="", **kw):
    return _fmt(doc.add_paragraph(), text, **kw)


def _cell_p(cell, text="", *, first=False, **kw):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    return _fmt(p, text, **kw)


def _new_doc() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(SIZE)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = Cm(MARGINS["left"])
    section.right_margin = Cm(MARGINS["right"])
    section.top_margin = Cm(MARGINS["top"])
    section.bottom_margin = Cm(MARGINS["bottom"])
    _page_number(section)
    return doc


def _page_number(section) -> None:
    """Номер страницы по центру вверху; на первой странице номера нет."""
    section.different_first_page_header_footer = True
    section.first_page_header.is_linked_to_previous = False
    p = section.header.paragraphs[0]
    p.alignment = ALIGN.CENTER
    run = p.add_run()
    run.font.size = Pt(PAGE_NUMBER_SIZE)
    run.font.name = FONT
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, end):
        run._r.append(node)


def _table(doc, widths: tuple[float, ...]):
    table = doc.add_table(rows=0, cols=len(widths))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    return table


def _row(table, widths: tuple[float, ...]):
    row = table.add_row()
    for cell, width in zip(row.cells, widths):
        cell.width = Cm(width)
    return row


def _merged_row(table, widths: tuple[float, ...]):
    row = _row(table, widths)
    return row.cells[0].merge(row.cells[-1])


def _person_row(table, widths: tuple[float, ...], person, *, number: str = "",
                line: float = LINE_AGENDA, row_space: int = 0) -> None:
    """[номер] ФАМИЛИЯ / Имя Отчество | – | должность (организация в Word не печатается, 07.10)."""
    cells = _row(table, widths).cells
    both = dict(space_before=row_space, space_after=row_space, line=line)
    offset = 0
    if number:
        _cell_p(cells[0], number, first=True, **both)
        offset = 1
    surname, given = split_name(person.name)
    _cell_p(cells[offset], surname, first=True, space_before=row_space, line=line)
    _cell_p(cells[offset], given, space_after=row_space, line=line)
    _cell_p(cells[offset + 1], DASH, first=True, align=ALIGN.CENTER, **both)
    role = _role_cell(person.role, getattr(person, "organization", ""))
    _cell_p(cells[offset + 2], role, first=True, align=ALIGN.JUSTIFY, **both)


# ── повестка ─────────────────────────────────────────────────────────────

def _heading(doc, m: Meeting, title: str, subtitle: str) -> None:
    _p(doc, title, bold=True, align=ALIGN.CENTER, space_after=0)
    if subtitle:
        _p(doc, subtitle, bold=True, align=ALIGN.CENTER, space_after=0)
    _p(doc, space_after=0)
    date_line = _date_line(m)
    if date_line:
        _p(doc, date_line, bold=True, align=ALIGN.RIGHT, space_after=0)
    for line in _place_lines(m):
        _p(doc, line, align=ALIGN.RIGHT, space_after=0)
    _p(doc, space_after=0)


def _speaker_rows(item) -> list[_Row]:
    return [_Row(name, item.speaker_roles[i] if i < len(item.speaker_roles) else "")
            for i, name in enumerate(item.speakers) if name]


def _speaker_block(doc, item) -> None:
    """Таблица выступающих под пунктом повестки: подзаголовок и строки людей.

    Кто докладчик, а кто содокладчик — по месту в списке (первый — докладчик, остальные —
    содокладчики), как и везде в интерфейсе (см. заголовок модуля).
    """
    rows = _speaker_rows(item)
    if not rows:
        return
    main, co = rows[:1], rows[1:]
    table = _table(doc, AGENDA_COLS)
    for kind, group in (("main", main), ("co", co)):
        if not group:
            continue
        single, plural = SPEAKER_TITLES[kind]
        cell = _merged_row(table, AGENDA_COLS)
        _cell_p(cell, single if len(group) == 1 else plural, first=True,
                bold=True, underline=True, left=INDENT_SPEAKER_TITLE,
                space_before=ROW_SPACE, space_after=ROW_SPACE)
        for person in group:
            _person_row(table, AGENDA_COLS, person, row_space=ROW_SPACE)


def build_agenda(m: Meeting, unchecked: bool = False) -> Document:
    doc = _new_doc()
    _heading(doc, m, "ПОВЕСТКА", event_genitive(m.subtitle) if (m.subtitle or "").strip() else "")

    opening = next((a for a in m.agenda if _is_opening(a.title)), None)
    closing = next((a for a in m.agenda if _is_closing(a.title)), None)
    numbered = [a for a in m.agenda if a is not opening and a is not closing]

    if opening is not None and (m.chair or m.chair_role):
        _p(doc, opening.title or "Вступительное слово", bold=True, first_line=FIRST_LINE_BLOCK, space_after=0)
        _p(doc, space_after=0)
        table = _table(doc, AGENDA_COLS)
        _person_row(table, AGENDA_COLS, _Row(m.chair, m.chair_role), row_space=ROW_SPACE)

    # Отбивка перед вопросом — отступ самого абзаца, а не пустая строка (см. `agenda/app/docx_export.py`):
    # пустой абзац после таблицы Word показывает не всегда, и вопросы расходились по-разному.
    for item in numbered:
        num = item.num or ""
        _p(doc, f"{num}. {item.title}" if num else item.title, bold=True, align=ALIGN.JUSTIFY,
           left=0, first_line=FIRST_LINE_ITEM, space_before=ITEM_GAP, space_after=ROW_SPACE)
        _speaker_block(doc, item)

    if closing is not None:
        _p(doc, closing.title or "Подведение итогов", bold=True, first_line=FIRST_LINE_BLOCK,
           space_before=ITEM_GAP, space_after=0)
    if unchecked:
        _mark_unchecked(doc)
    return doc


# ── список участников ────────────────────────────────────────────────────

ATTENDANCE_TITLES = {"in_person": "Участвуют очно:", "remote": "Участвуют по видеосвязи:"}


def _participant_sections(m: Meeting) -> list[tuple[str, list[Participant]]]:
    """Разделы списка («Участвуют очно:» / «Участвуют по видеосвязи:») — только когда в списке
    есть и очные, и по видеосвязи; иначе единый список без заголовков, как в образце. Порядок —
    ровно тот, что уже хранит совещание (сортировку по старшинству и её отключение вручную
    ведёт сам сервис, см. `participants_manual` — здесь не пересортировываем заново)."""
    ordered = list(m.participants)
    in_person = [p for p in ordered if p.attendance == "in_person"]
    remote = [p for p in ordered if p.attendance == "remote"]
    if not in_person or not remote:
        return [("", ordered)]
    unset = [p for p in ordered if p.attendance not in ("in_person", "remote")]
    sections = [("", unset)] if unset else []
    sections.append((ATTENDANCE_TITLES["in_person"], in_person))
    sections.append((ATTENDANCE_TITLES["remote"], remote))
    return sections


def build_participants(m: Meeting, unchecked: bool = False) -> Document:
    doc = _new_doc()
    _heading(doc, m, "СПИСОК", f"участников {event_genitive(m.subtitle)}")

    table = _table(doc, PARTICIPANT_COLS)
    number = 0
    for section_title, people in _participant_sections(m):
        if section_title:
            cell = _merged_row(table, PARTICIPANT_COLS)
            _cell_p(cell, section_title, first=True, bold=True, italic=True, underline=True,
                    space_before=ROW_SPACE, space_after=ROW_SPACE, line=LINE_LIST)
        for person in people:
            number += 1
            _person_row(table, PARTICIPANT_COLS, person, number=f"{number}.",
                       line=LINE_LIST, row_space=ROW_SPACE)
    if unchecked:
        _mark_unchecked(doc)
    return doc


# ── сохранение ───────────────────────────────────────────────────────────

def to_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _file_name(base: str, unchecked: bool) -> str:
    return base + (" (ФИО не сверены)" if unchecked else "") + ".docx"


def agenda_filename(unchecked: bool = False) -> str:
    """Ровно «Повестка.docx» (без даты и темы) — как и раньше у HTML-варианта, по названию
    заголовка «Материалы» узнают жанр документа и сами ставят номер, дату в имя не добавляют
    (`../CLAUDE.md` → «Повестка и Список участников в пакет»)."""
    return _file_name("Повестка", unchecked)


def participants_filename(unchecked: bool = False) -> str:
    return _file_name("Список участников", unchecked)
