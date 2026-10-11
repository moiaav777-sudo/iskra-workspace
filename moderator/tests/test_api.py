"""Тесты API: вход через портал, права, журнал, лимиты, выгрузка.

Запуск из корня проекта:  .venv/bin/python -m pytest -q
Искра и портал в тестах не вызываются — их функции подменяются.
"""
from __future__ import annotations

import io
import json
import zipfile

import pytest
from fastapi.testclient import TestClient

from app import deps, llm_extract, main, people_flow, portal, storage

REAL_PEOPLE_CALL = portal._people_call   # до подмены фикстурой: для проверки на настоящем HTTP

SECRET = "portal-test-secret"
ROLES = {"admin": "admin", "alice": "employee", "bob": "employee"}


@pytest.fixture()
def audit(monkeypatch):
    events = []
    monkeypatch.setattr(portal, "send_audit", lambda **e: events.append(e))
    return events


@pytest.fixture()
def client(tmp_path, monkeypatch, audit):
    monkeypatch.setenv("PORTAL_SECRET", SECRET)
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    monkeypatch.setattr(storage, "DB_PATH", tmp_path / "meetings.db")
    return TestClient(main.app)


class FakePeople:
    """Реестр людей портала (/internal/people) в памяти — с теми же правилами, что в портале."""

    def __init__(self):
        self.rows: dict[str, dict] = {}
        self.posts: list[dict] = []
        self.checks = 0                      # сколько раз спрашивали /internal/people/check (кэш сверки)
        self.down = False

    @staticmethod
    def key(name):
        return " ".join((name or "").lower().split())

    def call(self, method, path, body=None, **_kw):
        from urllib.parse import parse_qs, urlsplit
        if self.down:
            raise OSError("портал молчит")
        url = urlsplit(path)
        qs = {k: v[0] for k, v in parse_qs(url.query).items()}
        if method == "GET" and url.path == "/internal/people":
            q = qs.get("q", "").lower()
            return {"items": [r for r in self.rows.values()
                              if q in " ".join((r["name"], r["role"], r["organization"])).lower()]}
        if method == "GET" and url.path == "/internal/people/one":
            return self.rows.get(self.key(qs.get("name")))
        if method == "POST" and url.path == "/internal/people":
            # как в портале: без force — только вопрос, реестр не меняется (../CLAUDE.md п. 5а)
            self.posts.append(body)
            row = self.rows.get(self.key(body["name"]))
            if body["force"]:
                raise AssertionError("сервис не пишет в реестр сам — только человек из браузера")
            if not row:
                return {"summary": "", "suggest": "new"}
            differs = (body["role"] and body["role"] != row["role"]) or \
                      (body["organization"] and body["organization"] != row["organization"])
            return {"summary": "", "suggest": "differs" if differs else "same"}
        if method == "POST" and url.path == "/internal/people/check":
            self.checks += 1
            out = []
            for n in body["names"]:
                if self.key(n) in self.rows:
                    out.append({"name": n, "status": "exact", "person": self.rows[self.key(n)]})
                    continue
                alike = [r for r in self.rows.values()
                         if r["name"].split()[0][:3] == n.split()[0][:3] and r["name"].split()[1:] == n.split()[1:]]
                # как в портале с 02.10: короткая запись («Кравцов», «Кравцов Сергей») → close с полной
                short = n.split()
                alike += [r for r in self.rows.values() if r not in alike
                          and r["name"].split()[0] == short[0] and len(r["name"].split()) > len(short)
                          and r["name"].split()[1:len(short)] == short[1:]]
                out.append({"name": n, "status": "close", "suggestions": alike} if alike
                           else {"name": n, "status": "unknown"})
            return {"items": out}
        raise AssertionError(f"неожиданный запрос {method} {path}")


@pytest.fixture(autouse=True)
def people(monkeypatch):
    fake = FakePeople()
    monkeypatch.setattr(portal, "_people_call", fake.call)
    portal._people_cache.clear()
    people_flow._checked.clear()          # кэш сверки: в каждом тесте своя база, а ключ совещания тот же
    portal._request.set(None)             # контекст запроса (бюджет реестра) не протекает из теста в тест
    monkeypatch.setattr(portal, "_people_down_until", 0.0)
    return fake


def login(client: TestClient, who: str, secret: str = SECRET) -> TestClient:
    """Имитация nginx: заголовки пользователя от портала."""
    client.headers.update({"X-Portal-Secret": secret, "X-Portal-User": who, "X-Portal-Role": ROLES.get(who, "employee")})
    return client


def create(client: TestClient, **kw) -> str:
    r = client.post("/api/meetings", json={"title": "Тест", **kw})
    assert r.status_code == 200, r.text
    return r.json()["meeting"]["id"]


def test_healthz_without_login(client):
    assert client.get("/healthz").text == "ok"


def test_api_requires_portal_headers(client):
    assert client.get("/api/meetings").status_code == 401
    r = client.get("/", follow_redirects=False)
    assert r.status_code == 302 and r.headers["location"].startswith("/login")


def test_forged_headers_without_secret_rejected(client):
    login(client, "admin", secret="wrong")
    assert client.get("/api/meetings").status_code == 401
    client.headers.pop("X-Portal-Secret")
    assert client.get("/api/me").status_code == 401


def test_me_from_portal(client):
    login(client, "alice")
    client.headers["X-Portal-Name"] = "%D0%90%D0%BB%D0%B8%D1%81%D0%B0"
    me = client.get("/api/me").json()
    assert {k: me[k] for k in ("login", "name", "is_admin", "team")} == \
        {"login": "alice", "name": "Алиса", "is_admin": False, "team": []}
    # адреса соседей отдаёт сервер: в самой странице абсолютных адресов нет
    assert me["materials"]["package"].startswith("/materials/api/events/")


def test_everyone_sees_and_edits_all_meetings(client):
    login(client, "alice")
    mid = create(client)
    login(client, "bob")
    assert mid in [m["id"] for m in client.get("/api/meetings").json()["meetings"]]
    assert client.get(f"/api/meetings/{mid}").status_code == 200
    assert client.post(f"/api/meetings/{mid}/orders", json={"text": "Поручение Боба"}).status_code == 200
    assert client.delete(f"/api/meetings/{mid}").status_code == 200          # архив — можно всем
    assert client.post(f"/api/meetings/{mid}/restore").status_code == 200
    assert client.get(f"/api/meetings/{mid}/export").status_code == 200


def test_permanent_delete_admin_only(client, audit):
    login(client, "alice")
    mid = create(client)
    assert client.delete(f"/api/meetings/{mid}/permanent").status_code == 403
    login(client, "admin")
    assert client.delete(f"/api/meetings/{mid}/permanent").status_code == 200
    assert audit[-1]["action"] == "Совещание удалено навсегда" and audit[-1]["object"] == f"{mid} · Тест"


def test_changes_go_to_audit(client, audit):
    login(client, "alice")
    mid = create(client, title="Совещание по школам")
    client.put(f"/api/meetings/{mid}/header", json={"title": "Совещание по школам и детсадам", "date": "20.09.2026"})
    client.post(f"/api/meetings/{mid}/orders", json={"text": "Подготовить справку"})
    client.get(f"/api/meetings/{mid}/export")
    client.get(f"/api/meetings/{mid}")                                        # чтение в журнал не пишется
    actions = [e["action"] for e in audit]
    assert actions == ["Создано совещание", "Изменены данные совещания", "Изменены поручения", "Выгружен модератор-бриф"]
    assert audit[0]["object"] == f"{mid} · Совещание по школам"
    header = audit[1]
    assert header["login"] == "alice" and mid in header["object"]
    assert "название: «Совещание по школам» → «Совещание по школам и детсадам»" in header["details"]
    assert "поручения: 0→1" in audit[2]["details"]
    assert client.get(f"/api/meetings/{mid}/history").json()


def test_failed_request_not_audited(client, audit):
    login(client, "alice")
    client.delete("/api/meetings/нет-такого/permanent")
    assert audit == []


def test_share_endpoints_removed(client):
    login(client, "alice")
    mid = create(client)
    assert client.post(f"/api/meetings/{mid}/share", json={"login": "bob"}).status_code in (404, 405)
    assert client.get("/api/admin/users").status_code == 404


def test_analyze_unknown_meeting_before_iskra(client, monkeypatch):
    def boom(*a, **kw):
        raise AssertionError("Искра не должна вызываться")
    monkeypatch.setattr(llm_extract, "extract", boom)
    monkeypatch.setattr(llm_extract, "extract_with_file", boom)
    login(client, "alice")
    assert client.post("/api/analyze", json={"text": "x", "meeting_id": "нет-такого"}).status_code == 404
    r = client.post("/api/analyze-file?meeting_id=нет-такого", files={"file": ("a.docx", b"x")})
    assert r.status_code == 404


def test_unknown_item_key_rejected(client):
    login(client, "alice")
    mid = create(client)
    client.post(f"/api/meetings/{mid}/orders", json={"text": "Поручение"})
    for kind in ("orders", "scenario", "glossary", "materials", "participants", "agenda"):
        assert client.delete(f"/api/meetings/{mid}/{kind}/-1").status_code == 409
    assert len(client.get(f"/api/meetings/{mid}").json()["meeting"]["orders"]) == 1


def test_edits_follow_the_item_not_the_position(client):
    """Пока одна вкладка правит, другая удаляет пункт выше — правка должна остаться на своём пункте."""
    login(client, "alice")
    mid = create(client)
    for text in ("Первое", "Второе", "Третье"):
        client.post(f"/api/meetings/{mid}/orders", json={"text": text})
    orders = client.get(f"/api/meetings/{mid}").json()["meeting"]["orders"]
    third = orders[2]["id"]
    # «вторая вкладка» удалила первое поручение — позиции сдвинулись
    assert client.delete(f"/api/meetings/{mid}/orders/{orders[0]['id']}").status_code == 200
    assert client.delete(f"/api/meetings/{mid}/orders/{third}").status_code == 200
    left = client.get(f"/api/meetings/{mid}").json()["meeting"]["orders"]
    assert [o["text"] for o in left] == ["Второе"]
    # повторный запрос по тому же id — 409, а не удаление соседа
    assert client.delete(f"/api/meetings/{mid}/orders/{third}").status_code == 409


def test_speaker_position_checked_by_name(client):
    login(client, "alice")
    mid = create(client)
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "Вопрос"})
    key = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][0]["id"]
    for name in ("Иванов Иван Иванович", "Петров Пётр Петрович"):
        client.post(f"/api/meetings/{mid}/agenda/{key}/speakers", json={"name": name})
    # ФИО на этой позиции уже другое — правку не принимаем
    assert client.delete(f"/api/meetings/{mid}/agenda/{key}/speakers/1?name=Сидоров").status_code == 409
    assert client.put(f"/api/meetings/{mid}/agenda/{key}/minutes",
                      json={"minutes": [5, 7], "speakers": ["Иванов Иван Иванович"]}).status_code == 409
    r = client.delete(f"/api/meetings/{mid}/agenda/{key}/speakers/1?name=Петров Пётр Петрович")
    assert r.status_code == 200 and r.json()["meeting"]["agenda"][0]["speakers"] == ["Иванов Иван Иванович"]


def test_edit_agenda_title(client, audit):
    """Формулировку уже добавленного вопроса можно поправить, не трогая номер и докладчиков."""
    login(client, "alice")
    mid = create(client)
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О жилье",
                "speakers": ["Иванов Иван Иванович"], "speaker_roles": ["Министр"]})
    key = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][0]["id"]

    r = client.put(f"/api/meetings/{mid}/agenda/{key}", json={"title": "О вводе жилья для сирот"})
    assert r.status_code == 200, r.text
    item = r.json()["meeting"]["agenda"][0]
    assert item["title"] == "О вводе жилья для сирот" and item["num"] == "1"
    assert item["speakers"] == ["Иванов Иван Иванович"]
    assert audit[-1]["action"] == "Изменена повестка"

    # пустая формулировка не проходит
    r = client.put(f"/api/meetings/{mid}/agenda/{key}", json={"title": "   "})
    assert r.status_code == 400 and "формулировк" in r.json()["detail"]
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][0]["title"] == "О вводе жилья для сирот"

    # несуществующий/чужой ключ — 409, как у соседних маршрутов агенды (`deps.at`)
    assert client.put(f"/api/meetings/{mid}/agenda/no-such-key", json={"title": "Что угодно"}).status_code == 409

    # неизвестное совещание — обычный 404
    assert client.put(f"/api/meetings/no-such-meeting/agenda/{key}", json={"title": "Что угодно"}).status_code == 404


def test_edit_agenda_title_optimistic_lock(client):
    """Двое открыли один пункт повестки одновременно: кто сохранил вторым — получает 409, а не
    молча затирает правку первого (ревьюер RC-2026-09-28-6, некритическое замечание — как у
    `header`/`materials`, `base` — то, что редактор видел при открытии формы)."""
    login(client, "alice")
    mid = create(client)
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О жилье"})
    key = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][0]["id"]

    # base не прислали — работает как раньше, без проверки (старые клиенты не ломаются)
    assert client.put(f"/api/meetings/{mid}/agenda/{key}", json={"title": "О жилье (правка 1)"}).status_code == 200

    # base устарела — коллега успел поправить формулировку первым
    r = client.put(f"/api/meetings/{mid}/agenda/{key}",
                   json={"title": "Моя формулировка", "base": {"title": "О жилье"}})
    assert r.status_code == 409
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][0]["title"] == "О жилье (правка 1)"

    # base совпадает с текущим значением — правка проходит
    r = client.put(f"/api/meetings/{mid}/agenda/{key}",
                   json={"title": "О жилье (правка 2)", "base": {"title": "О жилье (правка 1)"}})
    assert r.status_code == 200
    assert r.json()["meeting"]["agenda"][0]["title"] == "О жилье (правка 2)"


def test_editors_warning_shows_colleague_not_self(client):
    login(client, "alice")
    mid = create(client)
    client.post(f"/api/meetings/{mid}/orders", json={"text": "Поручение Алисы"})
    assert client.get(f"/api/meetings/{mid}").json()["editors"] == []      # себя не показываем
    login(client, "bob")
    client.headers["X-Portal-Name"] = "%D0%91%D0%BE%D1%80%D0%B8%D1%81"
    client.post(f"/api/meetings/{mid}/orders", json={"text": "Поручение Бориса"})
    login(client, "alice")
    editors = client.get(f"/api/meetings/{mid}").json()["editors"]
    assert [e["login"] for e in editors] == ["bob"] and editors[0]["name"] == "Борис"
    assert client.get(f"/api/meetings/{mid}/editors").json()["editors"][0]["login"] == "bob"
    # чтение совещания правкой не считается
    login(client, "admin")
    client.get(f"/api/meetings/{mid}")
    login(client, "alice")
    assert [e["login"] for e in client.get(f"/api/meetings/{mid}").json()["editors"]] == ["bob"]


def test_long_edit_does_not_overwrite_colleague(client):
    """Двое открыли один блок сценария: второе сохранение не затирает первое молча."""
    login(client, "alice")
    mid = create(client)
    client.post(f"/api/meetings/{mid}/scenario", json={"section": "opening", "kind": "opening", "text": "Черновик"})
    key = client.get(f"/api/meetings/{mid}").json()["meeting"]["scenario"][0]["id"]
    login(client, "bob")
    client.headers["X-Portal-Name"] = "%D0%91%D0%BE%D1%80%D0%B8%D1%81"
    assert client.patch(f"/api/meetings/{mid}/scenario/{key}",
                        json={"text": "Версия Бориса", "base": {"text": "Черновик"}}).status_code == 200
    login(client, "alice")
    r = client.patch(f"/api/meetings/{mid}/scenario/{key}",
                     json={"text": "Версия Алисы", "base": {"text": "Черновик"}})
    assert r.status_code == 409 and "Борис" in r.json()["detail"]
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["scenario"][0]["text"] == "Версия Бориса"
    # осознанная перезапись поверх свежей версии проходит
    assert client.patch(f"/api/meetings/{mid}/scenario/{key}",
                        json={"text": "Версия Алисы", "base": {"text": "Версия Бориса"}}).status_code == 200
    # правка без base (старая вкладка) работает по-прежнему
    assert client.patch(f"/api/meetings/{mid}/scenario/{key}", json={"text": "Без сверки"}).status_code == 200


def test_header_and_material_conflicts(client):
    login(client, "alice")
    mid = create(client, title="Совещание")
    r = client.put(f"/api/meetings/{mid}/header", json={"title": "Новое", "base": {"title": "Другое"}})
    assert r.status_code == 409
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["title"] == "Совещание"
    assert client.put(f"/api/meetings/{mid}/header",
                      json={"title": "Новое", "base": {"title": "Совещание"}}).status_code == 200
    client.post(f"/api/meetings/{mid}/materials", json={"title": "Справка", "proposal": "Согласовать"})
    mat = client.get(f"/api/meetings/{mid}").json()["meeting"]["materials"][0]
    r = client.patch(f"/api/meetings/{mid}/materials/{mat['id']}",
                     json={"title": "Справка 2", "base": {"proposal": "Что-то другое"}})
    assert r.status_code == 409
    assert client.patch(f"/api/meetings/{mid}/materials/{mat['id']}",
                        json={"title": "Справка 2", "base": {"proposal": "Согласовать"}}).status_code == 200


AGENDA_PAYLOAD = {
    "source": {"service": "agenda", "meeting_id": "3f9a2b7c1d04"},
    "target_id": "",
    "title": "Совещание по отопительному сезону",
    "subtitle": "совещания у Заместителя Председателя Правительства",
    "chair": "Ольховой Виктор Петрович",
    "date": "2026-10-14",
    "time_start": "16:00",
    "location_primary": "Дом Правительства",
    "agenda": [
        {"num": "", "title": "Вступительное слово", "speakers": [], "speaker_minutes": [], "minutes": 5},
        {"num": "I", "title": "О готовности зданий",
         "speakers": ["Северцев Андрей Игоревич", "Гриневич Мария Олеговна"],
         "speaker_roles": ["Заместитель Министра строительства", "Статс-секретарь – заместитель Министра просвещения"],
         "speaker_minutes": [10, 10], "speaker_kinds": ["main", "co"], "minutes": 0},
    ],
    "participants": [
        {"name": "Кораблёва Ольга Дмитриевна", "role": "помощник Заместителя Председателя Правительства",
         "organization": "", "is_speaker": False, "attendance": "in_person"},
    ],
}


def test_import_from_agenda_creates_then_updates(client, audit):
    login(client, "alice")
    r = client.post("/api/meetings/import-from-agenda", json=AGENDA_PAYLOAD)
    assert r.status_code == 200, r.text
    out = r.json()
    mid = out["id"]
    assert out["created"] is True and out["url"] == f"/moderator/#meeting={mid}"
    assert audit[-1]["action"] == "Импорт повестки из сервиса «Повестка»"
    assert mid in audit[-1]["object"] and "повестка: 0→2" in audit[-1]["details"]

    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert m["date"] == "14.10.2026" and m["chair"] == "Ольховой Виктор Петрович"
    assert [a["title"] for a in m["agenda"]] == ["Вступительное слово", "О готовности зданий"]
    assert m["agenda"][1]["speaker_minutes"] == [10, 10]
    assert m["participants"][0]["protocol_rank"] > 0        # ранг считаем сами по должности
    assert m["source"] == {"service": "agenda", "meeting_id": "3f9a2b7c1d04"}

    # работа, сделанная в «Модераторе», при повторной передаче не теряется
    client.post(f"/api/meetings/{mid}/orders", json={"text": "Поручение из Модератора"})
    client.post(f"/api/meetings/{mid}/glossary", json={"term": "ОО", "definition": "образовательная организация"})
    again = dict(AGENDA_PAYLOAD, title="Совещание по отопительному сезону (уточнено)")
    again["agenda"] = AGENDA_PAYLOAD["agenda"][:1]
    r2 = client.post("/api/meetings/import-from-agenda", json=again)
    assert r2.status_code == 200 and r2.json() == {"id": mid, "url": out["url"], "created": False}
    m2 = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert m2["title"].endswith("(уточнено)") and len(m2["agenda"]) == 1
    assert [o["text"] for o in m2["orders"]] == ["Поручение из Модератора"]
    assert [g["term"] for g in m2["glossary"]] == ["ОО"]
    assert len(client.get("/api/meetings").json()["meetings"]) == 1   # дубля не появилось


FULL = ("совещания у Заместителя Председателя Правительства Российской Федерации "
        "по вопросам обеспечения жильем детей-сирот")


def test_full_title_from_agenda_is_shown_and_key_stays(client):
    """К2 RC-2026-09-22-1: полное наименование из «Повестки» — в subtitle, карточка показывает его,
    правка человеком пишет в subtitle; title (ключ find-or-create) не меняется — второго совещания нет."""
    login(client, "alice")
    payload = dict(AGENDA_PAYLOAD, title="Отопительный сезон", subtitle=FULL)
    mid = client.post("/api/meetings/import-from-agenda", json=payload).json()["id"]
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    NOM = "Совещание" + FULL[len("совещания"):]
    assert m["full_title"] == NOM and m["title"] == "Отопительный сезон"
    assert m["subtitle"] == FULL                                   # данные не переписаны

    # человек поправил полное наименование в карточке — меняется subtitle, ключ тот же
    fixed = NOM + " (уточнено)"
    r = client.put(f"/api/meetings/{mid}/header", json={
        "title": m["title"], "subtitle": fixed, "chair": m["chair"], "date": m["date"],
        "location_primary": m["location_primary"], "location_secondary": m["location_secondary"],
        "time_start": m["time_start"]})
    assert r.status_code == 200, r.text
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert m["full_title"] == fixed and m["title"] == "Отопительный сезон"
    again = client.post("/api/meetings/find-or-create", json={"date": m["date"], "title": "Отопительный сезон"}).json()
    assert again["id"] == mid and again["created"] is False

    # в документе заголовок — полное наименование
    html = client.get(f"/api/meetings/{mid}/export?unchecked=1").text
    assert "(уточнено)" in html


def test_document_title_from_agenda_is_nominative_and_data_untouched(client):
    """К1 RC-2026-09-22-2: «Повестка» шлёт шапку в родительном со строчной — в документе заголовок
    в именительном с заглавной, а сохранённый subtitle не переписывается."""
    login(client, "alice")
    payload = dict(AGENDA_PAYLOAD, title="Отопительный сезон", subtitle=FULL)
    mid = client.post("/api/meetings/import-from-agenda", json=payload).json()["id"]
    html = client.get(f"/api/meetings/{mid}/export?unchecked=1").text
    assert "Совещание у Заместителя Председателя" in html
    assert "совещания у Заместителя" not in html and ">Отопительный сезон<" not in html
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["subtitle"] == FULL
    # шапку без правки полного наименования сохраняем как есть — родительный остаётся в данных
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    client.put(f"/api/meetings/{mid}/header", json={
        "title": m["title"], "subtitle": m["subtitle"], "chair": "", "date": m["date"],
        "location_primary": "", "location_secondary": "", "time_start": ""})
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["subtitle"] == FULL


def test_plain_subtitle_keeps_title_as_document_heading(client):
    """К1 RC-2026-09-22-2: подзаголовок «в режиме ВКС» — не полное наименование: заголовок документа —
    прежнее название, а подзаголовок печатается отдельной строкой, как раньше."""
    login(client, "alice")
    mid = create(client, title="Отопительный сезон в школах", date="14.10.2026", subtitle="в режиме ВКС")
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert m["full_title"] == "Отопительный сезон в школах" and m["subtitle_line"] == "в режиме ВКС"
    html = client.get(f"/api/meetings/{mid}/export?unchecked=1").text
    assert "Отопительный сезон в школах" in html and "в режиме ВКС" in html
    from app.models import Meeting
    from app import exporter
    plain = exporter.render(Meeting(title="Отопительный сезон в школах", subtitle="в режиме ВКС"))
    assert "<h1>Отопительный сезон в школах</h1>" in plain and "в режиме ВКС" in plain
    plain = exporter.render(Meeting(title="Отопительный сезон", subtitle=FULL))
    assert "<h1>Совещание у Заместителя" in plain and "совещания у" not in plain


def test_full_title_rule_knows_adjectives_and_subtitle_is_not_erased(client):
    """Н2/Н3 RC-2026-09-22-3: «Рабочее совещание …» — полное наименование (в именительном при показе);
    повторный приём без subtitle его не стирает; при target_id ключ title не переписывается."""
    from app.models import full_title_of, is_full_title
    assert full_title_of("Т", "рабочего совещания у Заместителя по вопросам жилья") == \
        "Рабочее совещание у Заместителя по вопросам жилья"
    assert full_title_of("Т", "Расширенное заседание штаба") == "Расширенное заседание штаба"
    assert full_title_of("Т", "селекторного совещания с регионами") == "Селекторное совещание с регионами"
    assert full_title_of("Т", "рабочей встречи с губернаторами") == "Рабочая встреча с губернаторами"
    assert full_title_of("Т", "в режиме ВКС") == "Т" and not is_full_title("(закрытое)") and not is_full_title("2-е заседание")

    login(client, "alice")
    payload = dict(AGENDA_PAYLOAD, title="Отопительный сезон", subtitle="рабочего совещания у Заместителя")
    mid = client.post("/api/meetings/import-from-agenda", json=payload).json()["id"]
    # повторная передача той же повестки без subtitle — полное наименование остаётся
    again = client.post("/api/meetings/import-from-agenda", json=dict(AGENDA_PAYLOAD, title="Отопительный сезон", subtitle="")).json()
    assert again["id"] == mid
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert m["subtitle"] == "рабочего совещания у Заместителя" and m["full_title"] == "Рабочее совещание у Заместителя"
    # «Переименовать» в «Повестке» по-прежнему работает: та же повестка (source) — новое название доезжает
    client.post("/api/meetings/import-from-agenda", json=dict(AGENDA_PAYLOAD, title="Отопительный сезон 2026", subtitle=""))
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["title"] == "Отопительный сезон 2026"

    # человек выбрал совещание из списка (target_id): ключ title остаётся, повестка и участники приезжают
    other = create(client, title="Совещание, заведённое здесь", date="14.10.2026", subtitle="в режиме ВКС")
    r = client.post("/api/meetings/import-from-agenda", json=dict(AGENDA_PAYLOAD, target_id=other, source={"service": "agenda", "meeting_id": "zz"}, title="Чужое имя", subtitle=""))
    assert r.status_code == 200
    m = client.get(f"/api/meetings/{other}").json()["meeting"]
    assert m["title"] == "Совещание, заведённое здесь" and m["subtitle"] == "в режиме ВКС" and len(m["agenda"]) == 2
    assert client.post("/api/meetings/find-or-create", json={"date": "14.10.2026", "title": "Совещание, заведённое здесь"}).json()["id"] == other


def test_full_title_with_hyphen_does_not_crash(client):
    """К1 RC-2026-09-27-2: вид мероприятия с дефисом («совещания-семинара») ронял шапку и карточку — 500."""
    from app.models import full_title_of
    cases = {
        "Совещание-семинар по вопросам ЖКХ": "Совещание-семинар по вопросам ЖКХ",
        "совещания-семинара у министра": "Совещание-семинар у министра",
        "рабочего совещания-семинара": "Рабочее совещание-семинар",
        "Заседание-совещание штаба": "Заседание-совещание штаба",
        "встреча-брифинг": "Встреча-брифинг",
        "встречи-брифинга с журналистами": "Встреча-брифинг с журналистами",
    }
    for src, want in cases.items():
        assert full_title_of("Т", src) == want, src
    login(client, "alice")
    mid = create(client, title="ЖКХ", date="14.10.2026")
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    for src in cases:
        r = client.put(f"/api/meetings/{mid}/header", json={
            "title": m["title"], "subtitle": src, "chair": "", "date": m["date"],
            "location_primary": "", "location_secondary": "", "time_start": ""})
        assert r.status_code == 200, (src, r.text)
        card = client.get(f"/api/meetings/{mid}")
        assert card.status_code == 200 and card.json()["meeting"]["full_title"] == cases[src]
    assert client.get(f"/api/meetings/{mid}/export?unchecked=1").status_code == 200


def test_subtitle_equal_to_title_is_not_printed_twice():
    """Н1 RC-2026-09-22-3: подзаголовок, равный названию, отдельной строкой не печатается."""
    from app import exporter
    from app.models import Meeting
    html = exporter.render(Meeting(title="Отопительный сезон", subtitle="Отопительный сезон"))
    assert html.count("Отопительный сезон") == 2          # <title> и <h1>, без строки подзаголовка
    assert '<div class="subtitle">' not in html


def test_parsed_agenda_full_title_does_not_rewrite_key(client, monkeypatch):
    """Разбор повестки «Искрой»: полное наименование — в subtitle, а не в title (иначе — двойник)."""
    login(client, "alice")
    mid = create(client, title="Детям-сиротам жильё", date="14.10.2026")
    result = {"doc_type": "повестка", "confidence": "высокая", "summary": "",
              "data": {"title": "Совещание " + FULL, "agenda": [{"num": "1", "title": "О жилье"}]}}
    assert client.post(f"/api/meetings/{mid}/apply", json={"result": result}).status_code == 200
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert m["title"] == "Детям-сиротам жильё" and m["full_title"] == "Совещание " + FULL
    again = client.post("/api/meetings/find-or-create", json={"date": "14.10.2026", "title": "Детям-сиротам жильё"}).json()
    assert again["id"] == mid and again["created"] is False


def test_import_from_agenda_needs_login_and_known_target(client):
    r = client.post("/api/meetings/import-from-agenda", json=AGENDA_PAYLOAD)   # без заголовков портала
    assert r.status_code == 401
    login(client, "alice")
    r = client.post("/api/meetings/import-from-agenda", json=dict(AGENDA_PAYLOAD, target_id="нет-такого"))
    assert r.status_code == 404
    r = client.post("/api/meetings/import-from-agenda", json=AGENDA_PAYLOAD,
                    headers={"Origin": "https://evil.example"})
    assert r.status_code == 403


def test_find_or_create_is_idempotent(client, audit):
    """Совещание — общая ось: соседний сервис не заводит своё, а просит у нас (../CLAUDE.md п. 9)."""
    login(client, "alice")
    r = client.post("/api/meetings/find-or-create",
                    json={"date": "2026-10-14", "title": "  Отопительный   сезон ",
                          "source": {"service": "agenda", "id": "a-1"}})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["created"] is True and out["url"] == f"/moderator/#meeting={out['id']}" and out["folder"] == ""
    assert audit[-1]["action"] == "Совещание найдено или заведено"
    m = client.get(f"/api/meetings/{out['id']}").json()["meeting"]
    assert m["date"] == "14.10.2026" and m["title"] == "Отопительный сезон"
    assert m["source"] == {"service": "agenda", "meeting_id": "a-1"}

    # тот же день и то же название из другого сервиса и от другого человека — то же совещание
    login(client, "bob")
    again = client.post("/api/meetings/find-or-create",
                        json={"date": "14.10.2026", "title": "отопительный сезон",
                              "source": {"service": "theses", "id": "t-9"}}).json()
    assert again["id"] == out["id"] and again["created"] is False
    assert len(client.get("/api/meetings").json()["meetings"]) == 1      # дубля не появилось

    # другая дата — другое совещание
    other = client.post("/api/meetings/find-or-create",
                        json={"date": "15.10.2026", "title": "Отопительный сезон"}).json()
    assert other["id"] != out["id"] and other["created"] is True


def test_find_or_create_needs_date_and_title(client):
    r = client.post("/api/meetings/find-or-create", json={"date": "14.10.2026", "title": "Без входа"})
    assert r.status_code == 401
    login(client, "alice")
    assert client.post("/api/meetings/find-or-create", json={"title": "Без даты"}).status_code == 400
    assert client.post("/api/meetings/find-or-create", json={"date": "14.10.2026"}).status_code == 400
    r = client.post("/api/meetings/find-or-create", json={"date": "14.10.2026", "title": "Чужой сайт"},
                    headers={"Origin": "https://evil.example"})
    assert r.status_code == 403


LONG_TITLE = ("Совещание у Заместителя Председателя Правительства Российской Федерации "
              "Д.Н. Чернышенко по вопросам обеспечения жильем детей-сирот")


def _header(**kw):
    body = {"title": "", "chair": "", "date": "", "location_primary": "", "location_secondary": "",
            "time_start": ""}
    body.update(kw)
    return body


def test_short_title_everywhere(client):
    """Короткое название — тема совещания, одно правило на пространство (../CLAUDE.md п. 9)."""
    import datetime as dt
    login(client, "alice")
    soon = (dt.date.today() + dt.timedelta(days=3)).strftime("%d.%m.%Y")
    out = client.post("/api/meetings/find-or-create", json={"date": soon, "title": LONG_TITLE}).json()
    assert out["short_title"] == "Обеспечение жильем детей-сирот"
    mid = out["id"]

    card = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert card["short_title"] == card["short_title_auto"] == "Обеспечение жильем детей-сирот"
    row = [r for r in client.get("/api/meetings").json()["meetings"] if r["id"] == mid][0]
    assert row["short_title"] == "Обеспечение жильем детей-сирот" and row["title"] == LONG_TITLE
    up = [r for r in client.get("/api/upcoming").json()["meetings"] if r["id"] == mid][0]
    assert up["short_title"] == "Обеспечение жильем детей-сирот"


def test_short_title_can_be_fixed_by_hand(client, audit):
    login(client, "alice")
    mid = create(client, title="Совещание по вопросу итогового сочинения")
    auto = client.get(f"/api/meetings/{mid}").json()["meeting"]["short_title"]
    assert auto == "По вопросу итогового сочинения"      # правило честно не переписало незнакомое слово

    r = client.put(f"/api/meetings/{mid}/header", json=_header(
        title="Совещание по вопросу итогового сочинения", short_title="  Итоговое   сочинение "))
    assert r.status_code == 200, r.text
    assert "короткое название" in audit[-1]["details"]
    row = [x for x in client.get("/api/meetings").json()["meetings"] if x["id"] == mid][0]
    assert row["short_title"] == "Итоговое сочинение"

    # поле очистили — снова считается правилом; вписали ровно то же, что считает правило, — не закрепляем
    client.put(f"/api/meetings/{mid}/header", json=_header(
        title="Совещание по вопросу итогового сочинения", short_title=""))
    card = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert card["short_title"] == auto and card["short_title_manual"] == ""
    client.put(f"/api/meetings/{mid}/header", json=_header(
        title="Совещание по вопросу итогового сочинения", short_title=auto))
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["short_title_manual"] == ""

    # старый интерфейс без поля ничего не сбрасывает
    client.put(f"/api/meetings/{mid}/header", json=_header(title="Совещание по вопросу итогового сочинения",
                                                           short_title="Сочинение"))
    client.put(f"/api/meetings/{mid}/header", json=_header(title="Совещание по вопросу итогового сочинения"))
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["short_title"] == "Сочинение"


def test_folder_of_meeting(client, audit):
    """Папку ведут «Материалы», у нас лежит только её id — «Пакет к совещанию» ведёт туда."""
    login(client, "alice")
    mid = create(client)
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["folder"] == ""
    folder = "2026/10/14.10.2026 Отопительный сезон"
    r = client.put(f"/api/meetings/{mid}/folder", json={"folder": folder})
    assert r.status_code == 200 and r.json() == {"folder": folder, "changed": True}
    assert audit[-1]["action"] == "Привязана папка материалов"
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["folder"] == folder
    assert client.put(f"/api/meetings/{mid}/folder", json={"folder": folder}).json()["changed"] is False
    assert client.put(f"/api/meetings/{mid}/folder", json={"folder": "  "}).status_code == 400


def _meeting_for_word(client, time_start="14:00"):
    """Совещание с повесткой из двух вопросов и сценарием, как его собирает интерфейс."""
    mid = create(client, title="Совещание по вопросам обеспечения жильем детей-сирот",
                 date="14.10.2026", time_start=time_start)
    for num, title, speaker, role in (("1", "О вводе жилья", "Иванов Иван Иванович", "министр строительства"),
                                      ("2", "О расселении", "Петрова Анна Ивановна", "")):
        client.post(f"/api/meetings/{mid}/agenda", json={"num": num, "title": title, "speakers": [speaker],
                                                        "speaker_roles": [role], "speaker_minutes": [10]})
    client.post(f"/api/meetings/{mid}/scenario/generate")
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    opening = [b for b in m["scenario"] if b["section"] == "opening"][0]
    client.patch(f"/api/meetings/{mid}/scenario/{opening['id']}", json={
        "text": "<p><strong>Добрый день, коллеги!</strong></p><p>Сегодня <em>два вопроса</em>.</p>",
        "speaker": "", "minutes": 5})
    closing = [b for b in m["scenario"] if b["section"] == "closing"][0]
    client.patch(f"/api/meetings/{mid}/scenario/{closing['id']}", json={
        "text": "<ul><li>Спасибо за работу.</li><li>Протокол — в течение трёх дней.</li></ul>",
        "speaker": "", "minutes": 5})
    return mid


def test_moderator_docx_layout(client, audit):
    """Вёрстка по образцу пользователя: альбомный A4, поля 2 см, Times New Roman 16, по ширине."""
    import io
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    login(client, "alice")
    mid = _meeting_for_word(client)
    r = client.get(f"/api/meetings/{mid}/export.docx")
    assert r.status_code == 200, r.text
    assert "Модератор_" in __import__("urllib.parse").parse.unquote(r.headers["x-filename"])
    assert audit[-1]["action"] == "Выгружен модератор в Word"

    d = docx.Document(io.BytesIO(r.content))
    s = d.sections[0]
    assert s.orientation == WD_ORIENT.LANDSCAPE and s.page_width > s.page_height
    assert round(s.page_width.cm, 1) == 29.7 and round(s.page_height.cm, 1) == 21.0
    assert all(round(getattr(s, f).cm, 1) == 2.0 for f in ("left_margin", "right_margin", "top_margin", "bottom_margin"))
    assert s.different_first_page_header_footer                  # номер страницы — со второй
    normal = d.styles["Normal"]
    assert normal.font.name == "Times New Roman" and normal.font.size.pt == 16
    assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    paras = [p for p in d.paragraphs if p.text.strip()]
    texts = [p.text for p in paras]
    # шапки нет — документ начинается сразу со вступительного слова, со временем части
    assert texts[0] == "Вступительное слово · 14:00–14:05"
    # строки частей с часами — синим курсивом
    heads = [p for p in paras if p.text.startswith(("Вступительное слово · ", "Вопрос ", "Заключительное слово · "))]
    assert len(heads) == 4
    assert all(r.italic and str(r.font.color.rgb) == "0070C0" for p in heads for r in p.runs)
    # речь ведущего — маркированный список
    speech = [p for p in paras if p.style.name == "List Bullet"]
    assert [p.text for p in speech][:2] == ["Добрый день, коллеги!", "Сегодня два вопроса."]
    # длинного тире в документе нет: «—» → «–»
    assert [p.text for p in speech][-2:] == ["Спасибо за работу.", "Протокол – в течение трёх дней."]
    assert not any("—" in p.text for p in d.paragraphs)
    # жирный — как в редакторе; курсива в речи нет, кроме должности в передаче слова — синим курсивом
    assert speech[0].runs[0].bold
    handoff = [p for p in speech if p.text.startswith("По первому вопросу прошу доложить Иванова Ивана Ивановича")][0]
    role_runs = [r for r in handoff.runs if r.italic]
    assert [r.text for r in role_runs] == ["Министра строительства"]
    assert str(role_runs[0].font.color.rgb) == "0070C0"
    assert not any(r.italic for p in speech if p is not handoff for r in p.runs)
    # ремарки: кому слово и когда
    # доклад — синим курсивом: «Доклад И.И. Иванова, Министра строительства.»
    ivanov = [p for p in paras if p.text == "Доклад И.И. Иванова, Министра строительства."][0]
    assert all(r.italic and str(r.font.color.rgb) == "0070C0" for r in ivanov.runs)
    assert "Доклад А.И. Петровой." in texts
    # части разделены синим «///»: вступление, два вопроса, заключение — три разделителя
    seps = [p for p in paras if p.text == "///"]
    assert len(seps) == 3 and all(str(p.runs[0].font.color.rgb) == "0070C0" for p in seps)
    assert any(t.startswith("Заключительное слово · ") for t in texts)


def test_moderator_docx_needs_start_time(client):
    """Без времени начала документ не собирается и «5 минут» не подставляет."""
    login(client, "alice")
    mid = _meeting_for_word(client, time_start="")
    r = client.get(f"/api/meetings/{mid}/export.docx")
    assert r.status_code == 409 and "время начала" in r.json()["detail"]
    # время поставили в карточку — документ собирается
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    client.put(f"/api/meetings/{mid}/header", json={
        "title": m["title"], "chair": "", "date": m["date"], "location_primary": "", "location_secondary": "",
        "time_start": "10:30"})
    assert client.get(f"/api/meetings/{mid}/export.docx").status_code == 200


def test_moderator_docx_report_line():
    """Строка доклада — как у пользователя: инициалы и должность в родительном падеже."""
    from app import docx_export
    assert docx_export.report_line(
        "Корнеев Андрей Алексеевич", "Статс-секретарь – заместитель Министра просвещения Российской Федерации") \
        == "Доклад А.А. Корнеева, Статс-секретаря – заместителя Министра просвещения Российской Федерации."


def test_moderator_docx_filename_is_short_and_safe():
    from app import docx_export
    from app.models import Meeting
    m = Meeting(title='Совещание по вопросам "итогов"/сочинения: ' + "очень длинное " * 10)
    name = docx_export.filename(m)
    assert name.startswith("Модератор_") and name.endswith(".docx")
    assert not any(ch in name for ch in '\\/:*?"<>|') and len(name) <= 80


# ── «Повестка» и «Список участников» в пакет (../CLAUDE.md п. 9) ────────
# Материалы узнают вид документа по имени файла: ровно «Повестка» и «Список участников»
# (materials/app/wording.py → WHOLE), без имени совещания — поэтому имя фиксированное.

def test_export_agenda_document(client, audit):
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О вводе жилья",
                "speakers": ["Иванов Иван Иванович"], "speaker_roles": ["Министр строительства"]})
    r = client.get(f"/api/meetings/{mid}/export/agenda.html")
    assert r.status_code == 200, r.text
    assert r.headers["x-names-check"] == "done"
    import urllib.parse
    assert urllib.parse.unquote(r.headers["x-filename"]) == "Повестка.html"
    body = r.text
    assert "О вводе жилья" in body and "Иванов Иван Иванович" in body and "Министр строительства" in body
    assert audit[-1]["action"] == "Выгружена повестка (для пакета совещания)"


def test_export_participants_document(client, audit):
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/participants", json={
        "name": "Сидоров Сидор Сидорович", "role": "Директор департамента", "organization": "Минстрой России"})
    r = client.get(f"/api/meetings/{mid}/export/participants.html")
    assert r.status_code == 200, r.text
    assert r.headers["x-names-check"] == "done"
    import urllib.parse
    assert urllib.parse.unquote(r.headers["x-filename"]) == "Список участников.html"
    body = r.text
    assert "Сидоров Сидор Сидорович" in body and "Директор департамента" in body and "Минстрой России" in body
    assert audit[-1]["action"] == "Выгружен список участников (для пакета совещания)"


def test_export_agenda_checks_only_agenda_speakers(client, people):
    """У «Повестки» своя сверка: участник, который не докладчик, её не задерживает."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О вводе жилья",
                "speakers": ["Иванов Иван Иванович"], "speaker_roles": [""]})
    client.post(f"/api/meetings/{mid}/participants", json={"name": "Совсем Неизвестный Человек"})
    items = client.post(f"/api/meetings/{mid}/people/check-agenda").json()["items"]
    assert [i["name"] for i in items] == ["Иванов Иван Иванович"]
    r = client.get(f"/api/meetings/{mid}/export/agenda.html")
    assert r.status_code == 200


def test_export_participants_checks_everyone_not_just_speakers(client, people):
    """У «Списка участников» сверяются все, не только докладчики — иначе опечатку не заметят."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/participants", json={"name": "Совсем Неизвестный Человек"})
    items = client.post(f"/api/meetings/{mid}/people/check-participants").json()["items"]
    assert [i["name"] for i in items] == ["Совсем Неизвестный Человек"]
    assert items[0]["status"] == "unknown"


def test_export_agenda_blocks_close_name(client, people):
    """Опечатка в докладчике повестки — в документ не выпускаем, как у Word."""
    people.rows["кадочников павел анатольевич"] = {"name": "Кадочников Павел Анатольевич",
                                                   "role": "Министр", "organization": ""}
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О жилье",
                "speakers": ["Кадошников Павел Анатольевич"], "speaker_roles": ["Министр"]})
    r = client.get(f"/api/meetings/{mid}/export/agenda.html")
    assert r.status_code == 409 and "Кадошников" in r.json()["detail"]


def test_export_participants_blocks_close_name(client, people):
    """Опечатка в участнике — «Список участников» не собираем, пока не поправили ФИО."""
    people.rows["кадочников павел анатольевич"] = {"name": "Кадочников Павел Анатольевич",
                                                   "role": "Министр", "organization": ""}
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/participants", json={"name": "Кадошников Павел Анатольевич"})
    r = client.get(f"/api/meetings/{mid}/export/participants.html")
    assert r.status_code == 409 and "Кадошников" in r.json()["detail"]


def test_export_participants_without_registry_needs_flag(client, people):
    """Реестр молчит — документ не уходит молча, только с флагом `unchecked` и пометкой в самом файле."""
    people.down = True
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/participants", json={"name": "Сидоров Сидор Сидорович"})
    assert client.get(f"/api/meetings/{mid}/export/participants.html").status_code == 409
    r = client.get(f"/api/meetings/{mid}/export/participants.html?unchecked=1")
    assert r.status_code == 200 and r.headers["x-names-check"] == "skipped"
    assert "ФИО не сверены с реестром людей" in r.text


# ── «Повестка»/«Список участников» в Word по утверждённому образцу (решение пользователя
# 18.09/28.09, перенос из сервиса «Повестка», app/agenda_docx.py) — заменяют HTML в пакете ──


def _docx_text(d) -> str:
    """Весь читаемый текст документа: обычные абзацы и абзацы внутри ячеек таблиц —
    `Document.paragraphs` таблицы не видит вовсе, а у нас вся «шапка» людей — в таблицах."""
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def test_export_agenda_docx_layout_and_chair_block(client, people, audit):
    """Вёрстка по утверждённому образцу: книжная A4, Times New Roman 14; «Вступительное слово» —
    председатель строкой ФИО – должность, как любой докладчик; докладчик/содокладчик — по месту
    в списке (первый — докладчик, остальные — содокладчики, ../CLAUDE.md)."""
    import io
    import urllib.parse

    import docx
    from docx.enum.section import WD_ORIENT
    login(client, "alice")
    mid = create(client, date="14.10.2026", time_start="10:00")
    client.put(f"/api/meetings/{mid}/header", json={
        "chair": "Смирнов Пётр Ильич", "chair_role": "Заместитель Министра",
        "location_primary": "г. Москва, ул. Ленина, 1"})
    client.post(f"/api/meetings/{mid}/agenda", json={"title": "Вступительное слово",
               "speakers": ["Смирнов Пётр Ильич"], "speaker_roles": [""]})
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О ремонте школ",
               "speakers": ["Иванов Иван Иванович", "Петров Пётр Петрович"],
               "speaker_roles": ["Министр", "Замминистра"], "speaker_minutes": [10, 5]})
    client.post(f"/api/meetings/{mid}/agenda", json={"title": "Подведение итогов"})

    r = client.get(f"/api/meetings/{mid}/export/agenda.docx")
    assert r.status_code == 200, r.text
    assert r.headers["x-names-check"] == "done"
    assert urllib.parse.unquote(r.headers["x-filename"]) == "Повестка.docx"
    assert audit[-1]["action"] == "Выгружена повестка в Word (для пакета совещания)"

    d = docx.Document(io.BytesIO(r.content))
    section = d.sections[0]
    assert section.orientation == WD_ORIENT.PORTRAIT
    assert round(section.page_width.cm, 1) == 21.0 and round(section.page_height.cm, 1) == 29.7
    normal = d.styles["Normal"]
    assert normal.font.name == "Times New Roman" and normal.font.size.pt == 14
    assert section.different_first_page_header_footer   # номер страницы — со второй, как в образце

    text = _docx_text(d)
    assert "ПОВЕСТКА" in text
    assert "г. Москва, ул. Ленина, 1" in text
    assert "Вступительное слово" in text
    assert "СМИРНОВ" in text and "Пётр Ильич" in text and "Заместитель Министра" in text
    assert "1. О ремонте школ" in text
    assert "Докладчик:" in text and "Содокладчик:" in text   # первый — докладчик, второй — содокладчик
    assert "ИВАНОВ" in text and "ПЕТРОВ" in text
    assert "Подведение итогов" in text


def test_export_participants_docx_sections_by_attendance(client, people, audit):
    """Разделы «Участвуют очно:»/«Участвуют по видеосвязи:» — только когда в списке есть обе
    формы (как в образце); ведомство в Word не печатается (решение пользователя 07.10) — только должность."""
    import io
    import urllib.parse

    import docx
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/participants", json={
        "name": "Сидоров Сидор Сидорович", "role": "Директор департамента",
        "organization": "Минстрой России", "attendance": "in_person"})
    client.post(f"/api/meetings/{mid}/participants", json={
        "name": "Кузнецова Анна Павловна", "role": "Начальник отдела", "attendance": "remote"})

    r = client.get(f"/api/meetings/{mid}/export/participants.docx")
    assert r.status_code == 200, r.text
    assert urllib.parse.unquote(r.headers["x-filename"]) == "Список участников.docx"
    assert audit[-1]["action"] == "Выгружен список участников в Word (для пакета совещания)"

    d = docx.Document(io.BytesIO(r.content))
    text = _docx_text(d)
    assert "Участвуют очно:" in text and "Участвуют по видеосвязи:" in text
    assert "СИДОРОВ" in text and "Директор департамента" in text
    assert "Минстрой России" not in text and "Директор департамента," not in text
    assert "КУЗНЕЦОВА" in text


def test_role_cell_only_position():
    """07.10: в Word только должность; без должности — организация, чтобы ячейка не пустовала."""
    from app.agenda_docx import _role_cell
    assert _role_cell("Директор департамента", "Минстрой России") == "Директор департамента"
    assert _role_cell("", "Минстрой России") == "Минстрой России"
    assert _role_cell("  Министр  ", "") == "Министр"


def test_export_agenda_docx_blocks_close_name(client, people):
    """Та же обязательная сверка ФИО, что у HTML-варианта: опечатка не выпускается в Word."""
    people.rows["кадочников павел анатольевич"] = {"name": "Кадочников Павел Анатольевич",
                                                   "role": "Министр", "organization": ""}
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О жилье",
               "speakers": ["Кадошников Павел Анатольевич"], "speaker_roles": ["Министр"]})
    r = client.get(f"/api/meetings/{mid}/export/agenda.docx")
    assert r.status_code == 409 and "Кадошников" in r.json()["detail"]


def test_export_agenda_docx_blocks_close_name_in_chair(client, people):
    """Ревьюер RC-2026-09-28-14, К2: сверка раньше смотрела только докладчиков вопросов с
    номером и пропускала председателя — а Word печатает его строкой «Вступительного слова».
    Опечатка в председателе теперь тоже 409 close, как у обычного докладчика."""
    people.rows["кадочников павел анатольевич"] = {"name": "Кадочников Павел Анатольевич",
                                                   "role": "Министр", "organization": ""}
    login(client, "alice")
    mid = create(client, date="14.10.2026", time_start="10:00")
    client.put(f"/api/meetings/{mid}/header", json={"chair": "Кадошников Павел Анатольевич",
               "chair_role": "Министр"})
    client.post(f"/api/meetings/{mid}/agenda", json={"title": "Вступительное слово"})
    r = client.get(f"/api/meetings/{mid}/export/agenda.docx")
    assert r.status_code == 409 and "Кадошников" in r.json()["detail"]


def test_export_agenda_docx_blocks_close_name_in_unnumbered_question(client, people):
    """Вопрос, добавленный вручную, хранится с `num=''`, пока его не перенумеровали
    (`add_agenda` не зовёт `renumber_agenda`) — Word всё равно печатает его докладчика,
    и сверка обязана его увидеть, а не только вопросы с номером (ревьюер RC-2026-09-28-14, К2)."""
    people.rows["кадочников павел анатольевич"] = {"name": "Кадочников Павел Анатольевич",
                                                   "role": "Министр", "organization": ""}
    login(client, "alice")
    mid = create(client, date="14.10.2026", time_start="10:00")
    r = client.post(f"/api/meetings/{mid}/agenda", json={"title": "О ремонте школ",
               "speakers": ["Кадошников Павел Анатольевич"], "speaker_roles": ["Министр"]})
    assert r.json()["meeting"]["agenda"][0]["num"] == ""       # ещё не перенумерован
    r2 = client.get(f"/api/meetings/{mid}/export/agenda.docx")
    assert r2.status_code == 409 and "Кадошников" in r2.json()["detail"]


def test_export_participants_docx_without_registry_needs_flag_and_marks_document(client, people):
    """Реестр молчит — без сверки Word не уходит молча: 409, затем `unchecked=1` с пометкой,
    видной в самом документе (первая строка и колонтитул — тот же приём, что у документа
    ведущего, app/docx_export.py)."""
    import io

    import docx
    people.down = True
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/participants", json={"name": "Сидоров Сидор Сидорович"})
    assert client.get(f"/api/meetings/{mid}/export/participants.docx").status_code == 409
    r = client.get(f"/api/meetings/{mid}/export/participants.docx?unchecked=1")
    assert r.status_code == 200 and r.headers["x-names-check"] == "skipped"
    import urllib.parse
    assert "(ФИО не сверены)" in urllib.parse.unquote(r.headers["x-filename"])
    d = docx.Document(io.BytesIO(r.content))
    first = [p for p in d.paragraphs if p.text.strip()][0]
    assert first.text.startswith("ФИО не сверены с реестром людей") and first.runs[0].bold
    assert any("ФИО не сверены" in p.text for p in d.sections[0].footer.paragraphs)


def test_header_chair_role_roundtrip(client):
    """Должность председателя — новое поле шапки, нужно только для Word-повестки/списка."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    r = client.put(f"/api/meetings/{mid}/header", json={
        "chair": "Смирнов Пётр Ильич", "chair_role": "Заместитель Министра"})
    assert r.status_code == 200
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert m["chair"] == "Смирнов Пётр Ильич" and m["chair_role"] == "Заместитель Министра"


THESES_PAYLOAD = {
    "text": "Коллеги, сегодня обсуждаем готовность школ к отопительному сезону.\n"
            "Из 120 зданий паспорта готовности получили 118.\n"
            "Два здания остаются без паспорта — прошу доложить сроки.",
    "minutes": 7,
    "source": {"service": "theses", "id": "th-100500"},
}


def test_opening_from_theses(client, audit):
    """Тезисы из сервиса «Тезисы» ложатся во вступительное слово сценария."""
    login(client, "alice")
    mid = create(client)
    r = client.post(f"/api/meetings/{mid}/scenario/opening-from-theses", json=THESES_PAYLOAD)
    assert r.status_code == 200, r.text
    assert r.json() == {"ok": True, "url": f"/moderator/#meeting={mid}", "replaced": False}
    assert audit[-1]["action"] == "Вступительное слово из сервиса «Тезисы»"
    assert mid in audit[-1]["object"] and "th-100500" in audit[-1]["details"]
    assert "пунктов 3" in audit[-1]["details"]

    scenario = client.get(f"/api/meetings/{mid}").json()["meeting"]["scenario"]
    opening = [b for b in scenario if b["section"] == "opening"]
    assert len(opening) == 1 and opening[0]["minutes"] == 7
    # сценарий хранится rich-текстом: пункты тезисов — пункты списка, а не одна строка
    assert opening[0]["text"].count("<li>") == 3 and "паспорта готовности получили 118" in opening[0]["text"]
    html = client.get(f"/api/meetings/{mid}/export").text
    assert "Два здания остаются без паспорта" in html

    # повторная передача заменяет текст целиком, второго вступительного слова не появляется
    again = dict(THESES_PAYLOAD, text="Одна строка вместо трёх.", minutes=0)
    r2 = client.post(f"/api/meetings/{mid}/scenario/opening-from-theses", json=again)
    assert r2.status_code == 200 and r2.json()["replaced"] is True
    scenario = client.get(f"/api/meetings/{mid}").json()["meeting"]["scenario"]
    opening = [b for b in scenario if b["section"] == "opening"]
    assert len(opening) == 1
    assert opening[0]["text"] == "<p>Одна строка вместо трёх.</p>"
    assert opening[0]["minutes"] == 7          # minutes = 0 — время не трогаем


def test_opening_from_theses_keeps_scenario_and_checks_input(client):
    login(client, "alice")
    mid = create(client)
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О готовности зданий"})
    client.post(f"/api/meetings/{mid}/scenario/generate")
    before = len(client.get(f"/api/meetings/{mid}").json()["meeting"]["scenario"])
    assert client.post(f"/api/meetings/{mid}/scenario/opening-from-theses",
                       json=THESES_PAYLOAD).status_code == 200
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert len(m["scenario"]) == before                      # блоки сценария на месте
    assert m["scenario"][0]["section"] == "opening" and "<li>" in m["scenario"][0]["text"]

    # пустые тезисы не затирают уже написанное
    assert client.post(f"/api/meetings/{mid}/scenario/opening-from-theses",
                       json={"text": "   \n  ", "minutes": 5}).status_code == 400
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["scenario"][0]["text"].count("<li>") == 3


def test_opening_from_theses_needs_login_and_known_meeting(client):
    r = client.post("/api/meetings/x/scenario/opening-from-theses", json=THESES_PAYLOAD)
    assert r.status_code == 401                              # без заголовков портала
    login(client, "alice")
    mid = create(client)
    assert client.post("/api/meetings/нет-такого/scenario/opening-from-theses",
                       json=THESES_PAYLOAD).status_code == 404
    r = client.post(f"/api/meetings/{mid}/scenario/opening-from-theses", json=THESES_PAYLOAD,
                    headers={"Origin": "https://evil.example"})
    assert r.status_code == 403                              # чужой сайт не отправит тезисы за пользователя


def test_iskra_answer_checked_against_schema(client, monkeypatch):
    """Ответ Искры не по схеме не должен ни ронять запрос, ни класть мусор в бриф."""
    crooked = {
        "doc_type": "ПОВЕСТКА",
        "confidence": "высокая",
        "summary": "Повестка совещания",
        "data": {
            "title": "Совещание по школам",
            "agenda": [
                {"num": 1, "title": "О ремонте", "speakers": "Иванов Иван Иванович",
                 "speaker_roles": "Министр", "speaker_minutes": "15 мин"},
                {"num": 2, "speakers": ["Без темы"]},          # нет формулировки — отбрасываем
                "Строкой вместо объекта",
            ],
            "лишнее": {"что-то": 1},
        },
    }
    monkeypatch.setattr(llm_extract, "_call_iskra", lambda *a, **kw: json.dumps(crooked, ensure_ascii=False))
    login(client, "alice")
    mid = create(client)
    r = client.post("/api/analyze", json={"text": "текст повестки", "meeting_id": mid})
    assert r.status_code == 200, r.text
    res = r.json()["result"]
    assert res["doc_type"] == "повестка" and res["confidence"] == "низкая"
    assert res["schema_problems"]                      # пользователю скажем, что не взяли
    items = res["data"]["agenda"]
    assert [i["title"] for i in items] == ["О ремонте", "Строкой вместо объекта"]
    assert items[0]["num"] == "1" and items[0]["speakers"] == ["Иванов Иван Иванович"]
    assert items[0]["speaker_minutes"] == [15]         # «15 мин» → 15
    assert "лишнее" not in res["data"]
    # и такой результат применяется без ошибок
    assert client.post(f"/api/meetings/{mid}/apply", json={"result": res}).status_code == 200
    agenda = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"]
    assert "О ремонте" in [a["title"] for a in agenda]


def test_iskra_unknown_type_and_broken_shape(client, monkeypatch):
    monkeypatch.setattr(llm_extract, "_call_iskra",
                        lambda *a, **kw: '{"doc_type": "накладная", "data": ["не объект"], "confidence": "высокая"}')
    login(client, "alice")
    mid = create(client)
    res = client.post("/api/analyze", json={"text": "х", "meeting_id": mid}).json()["result"]
    assert res["doc_type"] == "неизвестно" and res["data"] == {}
    assert any("накладная" in p for p in res["schema_problems"])
    assert client.post(f"/api/meetings/{mid}/apply", json={"result": res}).status_code == 200

    monkeypatch.setattr(llm_extract, "_call_iskra", lambda *a, **kw: "Извините, не могу разобрать")
    r = client.post("/api/analyze", json={"text": "х", "meeting_id": mid})
    assert r.status_code == 502 and "Искра" in r.json()["detail"]


def test_material_answer_checked_against_schema(client, monkeypatch):
    monkeypatch.setattr(
        llm_extract, "_call_iskra_raw",
        lambda *a, **kw: json.dumps({"question_num": 2, "theses": "Один тезис строкой",
                                     "orders": [{"text": "объект вместо строки"}, "Поручение"],
                                     "proposal": "Согласовать", "summary": None}, ensure_ascii=False))
    login(client, "alice")
    mid = create(client)
    res = client.post(f"/api/meetings/{mid}/materials/analyze", json={"text": "текст справки"}).json()["result"]
    assert res["question_num"] == "2" and res["theses"] == ["Один тезис строкой"]
    # с 02.10 поручения — объекты (окно с галочками); строка старого ответа становится объектом
    assert res["orders"] == [
        {"id": "o1", "text": "объект вместо строки", "assignee": "", "deadline": ""},
        {"id": "o2", "text": "Поручение", "assignee": "", "deadline": ""}]
    assert res["summary"] == "" and res["type"] == "Материал к вопросу"
    assert client.post(f"/api/meetings/{mid}/materials/apply", json={"result": res}).status_code == 200


def test_registry_only_by_human(client, people):
    """В реестр пишет только человек: разбор спрашивает и ставит людей в очередь окна (../CLAUDE.md п. 5а)."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    doc = {"doc_type": "участники", "confidence": "высокая", "summary": "",
           "data": {"participants": [{"name": "Иванов Иван Иванович", "role": "Директор департамента",
                                      "organization": "Минпросвещения"}]}}
    m = client.post(f"/api/meetings/{mid}/apply", json={"result": doc}).json()["meeting"]
    assert people.rows == {}                                     # реестр не тронут
    assert people.posts[-1]["force"] is False and people.posts[-1]["source"] == "moderator"
    assert m["registry_pending"] == [{
        "name": "Иванов Иван Иванович", "role": "Директор департамента", "organization": "Минпросвещения",
        "suggest": "new", "document": "Список участников к совещанию «Тест» от 14.10.2026",
        "registry_role": "", "registry_org": ""}]

    # записан иначе — в очередь с тем, что в реестре; одинаково — окна нет
    people.rows["петров пётр петрович"] = {"name": "Петров Пётр Петрович", "role": "Министр", "organization": ""}
    people.rows["сидоров сидор сидорович"] = {"name": "Сидоров Сидор Сидорович", "role": "Директор", "organization": ""}
    doc2 = {"doc_type": "участники", "confidence": "высокая", "summary": "",
            "data": {"participants": [{"name": "Петров Пётр Петрович", "role": "Заместитель министра"},
                                      {"name": "Сидоров Сидор Сидорович", "role": "Директор"}]}}
    m = client.post(f"/api/meetings/{mid}/apply", json={"result": doc2}).json()["meeting"]
    pending = {p["name"]: p for p in m["registry_pending"]}
    assert pending["Петров Пётр Петрович"]["suggest"] == "differs"
    assert pending["Петров Пётр Петрович"]["registry_role"] == "Министр"
    assert "Сидоров Сидор Сидорович" not in pending

    # человек решил в окне — из очереди уходит
    r = client.post(f"/api/meetings/{mid}/people/pending/resolve", json={"name": "Иванов Иван Иванович", "added": True})
    assert [p["name"] for p in r.json()["pending"]] == ["Петров Пётр Петрович"]
    assert client.get("/api/me").json()["people"]["add"] == "/api/people"   # «Добавить» — из браузера в портал


def test_import_from_agenda_queues_people(client, people):
    login(client, "alice")
    out = client.post("/api/meetings/import-from-agenda", json=AGENDA_PAYLOAD).json()
    m = client.get(f"/api/meetings/{out['id']}").json()["meeting"]
    assert {p["suggest"] for p in m["registry_pending"]} == {"new"}
    assert all(p["document"].startswith("Повестка к совещанию") for p in m["registry_pending"])
    assert people.rows == {}


def test_manual_speaker_and_participant_add_queue_registry(client, people):
    """Докладчика и участника вписали руками (не из подсказки реестра) — не тихо: то же окно
    «Добавить в реестр?», что у разбора и импорта (../CLAUDE.md п. 5а)."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О жилье"})
    key = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][0]["id"]

    # новый докладчик, вписанный вручную, — в очередь окна с указанием документа «Повестка»
    r = client.post(f"/api/meetings/{mid}/agenda/{key}/speakers",
                    json={"name": "Совсем Новый Докладчик", "role": "Министр"})
    m = r.json()["meeting"]
    pending = {p["name"]: p for p in m["registry_pending"]}
    assert pending["Совсем Новый Докладчик"]["suggest"] == "new"
    assert pending["Совсем Новый Докладчик"]["document"].startswith("Повестка к совещанию")
    client.post(f"/api/meetings/{mid}/people/pending/resolve",
               json={"name": "Совсем Новый Докладчик", "added": False})

    # записан в реестре иначе — тоже в очередь, с тем, что там записано
    people.rows["петров пётр петрович"] = {"name": "Петров Пётр Петрович", "role": "Министр", "organization": ""}
    r = client.post(f"/api/meetings/{mid}/agenda/{key}/speakers",
                    json={"name": "Петров Пётр Петрович", "role": "Заместитель министра"})
    pending = {p["name"]: p for p in r.json()["meeting"]["registry_pending"]}
    assert pending["Петров Пётр Петрович"]["suggest"] == "differs"
    assert pending["Петров Пётр Петрович"]["registry_role"] == "Министр"
    client.post(f"/api/meetings/{mid}/people/pending/resolve",
               json={"name": "Петров Пётр Петрович", "added": False})

    # тот же человек, та же запись — реестр говорит «same», окна нет
    people.rows.clear()
    people.rows["сидоров сидор сидорович"] = {"name": "Сидоров Сидор Сидорович",
                                              "role": "Директор", "organization": "Минстрой"}
    r = client.post(f"/api/meetings/{mid}/participants",
                    json={"name": "Сидоров Сидор Сидорович", "role": "Директор", "organization": "Минстрой"})
    assert r.json()["meeting"]["registry_pending"] == []

    # новый участник, вписанный вручную, — в очередь с документом «Список участников»
    r = client.post(f"/api/meetings/{mid}/participants", json={"name": "Совсем Новый Участник"})
    pending = {p["name"]: p for p in r.json()["meeting"]["registry_pending"]}
    assert pending["Совсем Новый Участник"]["suggest"] == "new"
    assert pending["Совсем Новый Участник"]["document"].startswith("Список участников к совещанию")


# ── список изменений с галочками (решение пользователя 28.09, перенос из сервиса «Повестка»,
# app/agenda_changes.py) — «Заменить целиком» для непустой повестки/участников стирало
# материалы, поручения и сценарий, привязанные к номерам вопросов ──────────────────────────


def _mock_iskra(monkeypatch, doc: dict) -> None:
    monkeypatch.setattr(llm_extract, "_call_iskra", lambda *a, **kw: json.dumps(doc, ensure_ascii=False))


def test_analyze_offers_checklist_for_nonempty_agenda_not_full_replace(client, monkeypatch, people):
    """Непустая повестка — «Искра» предлагает список изменений, а не «Заменить целиком»."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О ремонте школ",
                "speakers": ["Иванов Иван Иванович"], "speaker_roles": ["Министр"], "speaker_minutes": [10]})
    doc = {"doc_type": "повестка", "confidence": "высокая", "summary": "",
           "data": {"agenda": [
               {"title": "О ремонте школ", "speakers": ["Петров Пётр Петрович"],
                "speaker_roles": ["Замминистра"], "speaker_minutes": [15]},
               {"title": "Новый вопрос", "speakers": [], "speaker_roles": [], "speaker_minutes": []},
           ]}}
    _mock_iskra(monkeypatch, doc)
    r = client.post("/api/analyze", json={"text": "текст повестки", "meeting_id": mid})
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["changes"] is not None            # непустая повестка — список изменений, а не None
    kinds = {op["kind"] for op in data["changes"]}
    assert kinds == {"replace_speaker", "add_question"}
    replace_op = next(op for op in data["changes"] if op["kind"] == "replace_speaker")
    add_op = next(op for op in data["changes"] if op["kind"] == "add_question")
    assert replace_op["checked"] is True and add_op["checked"] is True   # добавления — по умолчанию отмечены

    # внесли только замену докладчика — «Новый вопрос» человек не отмечал
    r2 = client.post(f"/api/meetings/{mid}/agenda/apply-changes", json={"ops": [replace_op]})
    assert r2.status_code == 200, r2.text
    m = r2.json()["meeting"]
    assert len(m["agenda"]) == 1 and m["agenda"][0]["speakers"] == ["Петров Пётр Петрович"]
    # докладчик из точечной правки тоже идёт в очередь «Добавить в реестр?» (та же очередь,
    # что у ручного добавления, ../CLAUDE.md п. 5а)
    assert "Петров Пётр Петрович" in {p["name"] for p in m["registry_pending"]}


def test_apply_changes_remove_question_cascades_materials_and_orders(client, monkeypatch):
    """Убрать вопрос по галочке — вместе с ним уходят его материалы и поручения: без этого они
    молча остались бы висеть на чужом вопросе после пересчёта номеров (renumber_agenda)."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О ремонте школ"})
    num = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][0]["num"]
    client.post(f"/api/meetings/{mid}/orders", json={"text": "Подготовить смету", "agenda_num": num})

    doc = {"doc_type": "повестка", "confidence": "высокая", "summary": "", "data": {"agenda": []}}
    _mock_iskra(monkeypatch, doc)
    r = client.post("/api/analyze", json={"text": "текст", "meeting_id": mid})
    changes = r.json()["changes"]
    assert len(changes) == 1 and changes[0]["kind"] == "remove_question"
    assert changes[0]["checked"] is False              # удаления не отмечены по умолчанию
    assert "поручениями: 1" in changes[0]["label"]

    r2 = client.post(f"/api/meetings/{mid}/agenda/apply-changes", json={"ops": [changes[0]]})
    assert r2.status_code == 200, r2.text
    m = r2.json()["meeting"]
    assert m["agenda"] == [] and m["orders"] == []     # поручение ушло вместе с вопросом


def test_apply_changes_remove_question_stale_when_colleague_added_material(client, monkeypatch):
    """Ревьюер RC-2026-09-28-14, К1: Искра предложила «убрать вопрос», подпись без материалов;
    пока человек разбирал список, коллега положил к этому вопросу материал (например, «Бриф»
    прислал саммари доклада) — формулировка вопроса не изменилась, но каскад молча унёс бы
    и материал. Это тоже устаревшая операция: 409, а не тихая потеря чужой правки."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О ремонте школ"})

    doc = {"doc_type": "повестка", "confidence": "высокая", "summary": "", "data": {"agenda": []}}
    _mock_iskra(monkeypatch, doc)
    r = client.post("/api/analyze", json={"text": "текст", "meeting_id": mid})
    changes = r.json()["changes"]
    assert len(changes) == 1 and changes[0]["kind"] == "remove_question"
    assert "материалами" not in changes[0]["label"]   # на момент предложения материалов не было

    # тем временем коллега положил материал к вопросу 1
    num = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][0]["num"]
    client.post(f"/api/meetings/{mid}/materials", json={"question_num": num, "title": "Саммари доклада"})

    r2 = client.post(f"/api/meetings/{mid}/agenda/apply-changes", json={"ops": [changes[0]]})
    assert r2.status_code == 409, r2.text
    detail = r2.json()["detail"]
    assert changes[0]["id"] in [s["id"] for s in detail["stale"]]
    # без force материал и вопрос целы
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert len(m["agenda"]) == 1 and len(m["materials"]) == 1


def test_apply_changes_remove_question_stale_when_colleague_added_order_or_scenario(client, monkeypatch):
    """То же самое для поручения и блока сценария, добавленных за время разбора."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О ремонте школ"})
    doc = {"doc_type": "повестка", "confidence": "высокая", "summary": "", "data": {"agenda": []}}
    _mock_iskra(monkeypatch, doc)
    num = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][0]["num"]

    # поручение
    r = client.post("/api/analyze", json={"text": "текст", "meeting_id": mid})
    op = r.json()["changes"][0]
    client.post(f"/api/meetings/{mid}/orders", json={"text": "Подготовить смету", "agenda_num": num})
    r2 = client.post(f"/api/meetings/{mid}/agenda/apply-changes", json={"ops": [op]})
    assert r2.status_code == 409, r2.text

    # блок сценария
    r = client.post("/api/analyze", json={"text": "текст", "meeting_id": mid})
    op = r.json()["changes"][0]
    client.post(f"/api/meetings/{mid}/scenario", json={"section": "question", "agenda_num": num, "kind": "reference"})
    r3 = client.post(f"/api/meetings/{mid}/agenda/apply-changes", json={"ops": [op]})
    assert r3.status_code == 409, r3.text
    # вопрос по-прежнему на месте — ничего не потеряно
    assert len(client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"]) == 1


def test_apply_changes_stale_conflict_then_force(client, monkeypatch):
    """Пока человек отмечал галочки, вопрос успели убрать в другом окне — 409 с раскладкой
    «что применимо / что устарело», а с `force` вносится только применимое."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "Вопрос А"})
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "2", "title": "Вопрос Б"})
    doc = {"doc_type": "повестка", "confidence": "высокая", "summary": "",
           "data": {"agenda": [{"title": "Вопрос А"}, {"title": "Вопрос В"}]}}
    _mock_iskra(monkeypatch, doc)
    r = client.post("/api/analyze", json={"text": "текст", "meeting_id": mid})
    changes = r.json()["changes"]
    remove_op = next(op for op in changes if op["kind"] == "remove_question")
    add_op = next(op for op in changes if op["kind"] == "add_question")
    remove_op["checked"] = True   # человек отметил галочку сам

    # тем временем «Вопрос Б» уже убрали в другом окне обычной кнопкой
    key_b = next(a["id"] for a in client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"]
                if a["title"] == "Вопрос Б")
    client.delete(f"/api/meetings/{mid}/agenda/{key_b}")

    r2 = client.post(f"/api/meetings/{mid}/agenda/apply-changes", json={"ops": [remove_op, add_op]})
    assert r2.status_code == 409
    detail = r2.json()["detail"]
    assert remove_op["id"] in [s["id"] for s in detail["stale"]]
    assert add_op["id"] in detail["applicable"]
    # без force ничего не внесено
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][0]["title"] == "Вопрос А"

    r3 = client.post(f"/api/meetings/{mid}/agenda/apply-changes",
                     json={"ops": [remove_op, add_op], "force": True})
    assert r3.status_code == 200, r3.text
    m = r3.json()["meeting"]
    assert {a["title"] for a in m["agenda"]} == {"Вопрос А", "Вопрос В"}
    assert r3.json()["skipped"] and r3.json()["skipped"][0]["id"] == remove_op["id"]


def test_analyze_offers_checklist_for_nonempty_participants(client, monkeypatch):
    """То же самое для «Список участников»: правка уже заполненного поля не отмечена по
    умолчанию (перезапись существующего значения — решение человека), добавление — отмечено."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/participants", json={"name": "Сидоров Сидор Сидорович", "role": "Директор"})
    doc = {"doc_type": "участники", "confidence": "высокая", "summary": "",
           "data": {"participants": [
               {"name": "Сидоров Сидор Сидорович", "role": "Заместитель директора"},
               {"name": "Новый Человек Иванович", "role": "Эксперт"},
           ]}}
    _mock_iskra(monkeypatch, doc)
    r = client.post("/api/analyze", json={"text": "текст", "meeting_id": mid})
    changes = r.json()["changes"]
    assert changes is not None
    role_op = next(op for op in changes if op["kind"] == "participant_role")
    add_op = next(op for op in changes if op["kind"] == "add_participant")
    assert role_op["checked"] is False and add_op["checked"] is True

    r2 = client.post(f"/api/meetings/{mid}/agenda/apply-changes", json={"ops": [role_op, add_op]})
    assert r2.status_code == 200, r2.text
    by_name = {p["name"]: p for p in r2.json()["meeting"]["participants"]}
    assert by_name["Сидоров Сидор Сидорович"]["role"] == "Заместитель директора"
    assert "Новый Человек Иванович" in by_name


def test_apply_changes_participant_removal_skips_agenda_speakers(client):
    """Докладчика повестки «убрать участника» не предлагает — им ведает повестка
    (`sync_agenda_participants`), а не список изменений участников."""
    from app import agenda_changes, storage
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "Вопрос",
                "speakers": ["Иванов Иван Иванович"], "speaker_roles": [""]})
    client.post(f"/api/meetings/{mid}/participants", json={"name": "Гость Без Доклада"})
    m = storage.get(mid)
    changes = agenda_changes.diff_participants(m, {"participants": []})
    removed = {op["name"] for op in changes if op["kind"] == "remove_participant"}
    assert removed == {"Гость Без Доклада"}


# ── приём из «Приёмной» «Материалов» (шаг 2 перехода «Повестка → Совещание», решение
# пользователя 29.09) — контракт списан с agenda/app/main.py::/api/inbox, только читали ──────


def _inbox_file(name="Повестка.docx", body=b"PK\x03\x04content"):
    return {"file": (name, body, "application/octet-stream")}


def test_inbox_empty_agenda_full_mode(client, monkeypatch, audit):
    """Пустая повестка — предложение целиком, как у обычной карточки «Вставить документ»."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    result = {"doc_type": "повестка", "confidence": "высокая", "summary": "Повестка совещания",
              "data": {"title": "", "date": "", "time": "", "location": "", "chair": "",
                       "agenda": [{"num": "1", "title": "О жилье", "time": "",
                                  "speakers": ["Иванов Иван Иванович"], "speaker_roles": ["Министр"],
                                  "speaker_minutes": [10]}]},
              "schema_problems": []}
    monkeypatch.setattr(llm_extract, "extract_with_file", lambda *a, **kw: result)
    r = client.post("/api/inbox", data={"meeting": mid, "sender": "89991234567"}, files=_inbox_file())
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["ok"] is True and data["mode"] == "full" and data["doc_type"] == "повестка"
    assert "вопросов 1" in data["summary"]
    assert data["url"].startswith(f"/moderator/#meeting={mid}&proposal=")
    key = data["url"].rsplit("proposal=", 1)[1]

    proposal = client.get(f"/api/meetings/{mid}/iskra/proposal/{key}").json()
    assert proposal["result"]["doc_type"] == "повестка" and proposal["mode"] == "full"
    assert audit[-1]["action"] == "Принят документ из «Приёмной» (Материалы)"
    assert "из «Приёмной»: «Повестка.docx»" in audit[-1]["details"]
    assert "89991234567" in audit[-1]["details"]


def test_inbox_dedup_same_document_not_reanalyzed(client, monkeypatch):
    """Тот же документ к тому же совещанию второй раз — прежнее предложение, «Искра» не звалась."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    calls = []

    def fake_extract(data_bytes, filename, hint):
        calls.append(1)
        return {"doc_type": "повестка", "confidence": "высокая", "summary": "",
                "data": {"agenda": [{"title": "Вопрос", "speakers": [], "speaker_roles": [],
                                     "speaker_minutes": []}]}, "schema_problems": []}
    monkeypatch.setattr(llm_extract, "extract_with_file", fake_extract)
    r1 = client.post("/api/inbox", data={"meeting": mid}, files=_inbox_file())
    r2 = client.post("/api/inbox", data={"meeting": mid}, files=_inbox_file())
    assert r1.status_code == 200 and r2.status_code == 200
    assert len(calls) == 1, "второй вызов с тем же документом не должен звать «Искру» снова"
    assert r1.json()["url"] == r2.json()["url"]

    # другой файл (другие байты) к тому же совещанию — уже новый разбор
    r3 = client.post("/api/inbox", data={"meeting": mid}, files=_inbox_file(body=b"PK\x03\x04different content"))
    assert len(calls) == 2
    assert r3.json()["url"] != r1.json()["url"]


def test_inbox_nonempty_agenda_changes_mode(client, monkeypatch):
    """Непустая повестка — список изменений через уже готовый `agenda_changes.py` (шаг 1)."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "Старый вопрос"})
    result = {"doc_type": "повестка", "confidence": "высокая", "summary": "",
              "data": {"agenda": [{"title": "Старый вопрос"}, {"title": "Новый вопрос"}]},
              "schema_problems": []}
    monkeypatch.setattr(llm_extract, "extract_with_file", lambda *a, **kw: result)
    r = client.post("/api/inbox", data={"meeting": mid}, files=_inbox_file())
    data = r.json()
    assert data["mode"] == "changes" and "изменений" in data["summary"]
    key = data["url"].rsplit("proposal=", 1)[1]
    proposal = client.get(f"/api/meetings/{mid}/iskra/proposal/{key}").json()
    assert proposal["mode"] == "changes"
    kinds = {op["kind"] for op in proposal["changes"]}
    assert "add_question" in kinds


def test_inbox_participants_summary(client, monkeypatch):
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    result = {"doc_type": "участники", "confidence": "высокая", "summary": "",
              "data": {"participants": [{"name": "Иванов Иван Иванович", "role": "Министр"}]},
              "schema_problems": []}
    monkeypatch.setattr(llm_extract, "extract_with_file", lambda *a, **kw: result)
    r = client.post("/api/inbox", data={"meeting": mid}, files=_inbox_file(name="Список участников.docx"))
    data = r.json()
    assert data["doc_type"] == "участники" and "участников 1" in data["summary"]


def test_inbox_requires_meeting_and_file_type(client):
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    assert client.post("/api/inbox", data={}, files=_inbox_file()).status_code == 400
    assert client.post("/api/inbox", data={"meeting": "нет-такого"}, files=_inbox_file()).status_code == 404
    r = client.post("/api/inbox", data={"meeting": mid}, files=_inbox_file(name="virus.exe", body=b"MZ"))
    assert r.status_code == 400 and "не принимается" in r.json()["detail"]


def test_inbox_proposal_missing_key(client):
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    r = client.get(f"/api/meetings/{mid}/iskra/proposal/no-such-key")
    # не просто «маршрута нет» (тоже 404) — конкретно «предложение устарело или не найдено»
    assert r.status_code == 404 and "не найдено" in r.json()["detail"]


def test_names_checked_before_word_and_fixed_in_one_click(client, people):
    """«Кадошников Павел» при записанном «Кадочникове» — в документ не выпускаем, заменяем одним нажатием."""
    people.rows["кадочников павел анатольевич"] = {"name": "Кадочников Павел Анатольевич",
                                                   "role": "Министр", "organization": ""}
    login(client, "alice")
    mid = create(client, date="14.10.2026", time_start="10:00")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О жилье",
                "speakers": ["Кадошников Павел Анатольевич"], "speaker_roles": ["Министр"], "speaker_minutes": [10]})
    client.post(f"/api/meetings/{mid}/scenario/generate")
    intro = [b for b in client.get(f"/api/meetings/{mid}").json()["meeting"]["scenario"] if b["kind"] == "intro"][0]
    assert "Павла Анатольевича" in intro["text"] and "Павела" not in intro["text"]   # общее склонение

    items = client.post(f"/api/meetings/{mid}/people/check").json()["items"]
    assert items[0]["status"] == "close" and items[0]["suggestions"][0]["name"] == "Кадочников Павел Анатольевич"
    r = client.get(f"/api/meetings/{mid}/export.docx")
    assert r.status_code == 409 and "Кадошников" in r.json()["detail"]

    r = client.post(f"/api/meetings/{mid}/people/replace",
                    json={"old": "Кадошников Павел Анатольевич", "new": "Кадочников Павел Анатольевич"})
    assert r.status_code == 200
    m = r.json()["meeting"]
    assert m["agenda"][0]["speakers"] == ["Кадочников Павел Анатольевич"]
    texts = " ".join(b["text"] for b in m["scenario"])
    assert "Кадошников" not in texts and "Кадочникова Павла Анатольевича" in texts
    assert client.post(f"/api/meetings/{mid}/people/check").json()["items"][0]["status"] == "exact"
    r = client.get(f"/api/meetings/{mid}/export.docx")
    assert r.status_code == 200 and r.headers["x-names-check"] == "done"


def test_opening_meeting_does_not_wait_for_portal(client, people, monkeypatch):
    """Открытие совещания не ходит в реестр: медленный портал его не тормозит (SRE 21.09, GET 10–12 с)."""
    import time as _time
    login(client, "alice")
    mid = create(client)
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О жилье",
                "speakers": ["Иванов Иван Иванович", "Петров Пётр Петрович"], "speaker_roles": ["", ""]})
    calls = []

    def slow(method, path, body=None, **_kw):
        calls.append(path)
        _time.sleep(2)
        return None
    monkeypatch.setattr(portal, "_people_call", slow)
    portal._people_cache.clear()
    started = _time.monotonic()
    assert client.get(f"/api/meetings/{mid}").status_code == 200
    assert _time.monotonic() - started < 1.0 and calls == []


def test_registry_questions_are_parallel_and_bounded(client, people, monkeypatch):
    """Разбор списка: вопросы к реестру параллельно и в пределах бюджета — ответ не ждёт медленный портал."""
    import time as _time
    login(client, "alice")
    mid = create(client)
    doc = {"doc_type": "участники", "confidence": "высокая", "summary": "",
           "data": {"participants": [{"name": f"Иванов{i} Иван Иванович", "role": "Директор"} for i in range(12)]}}

    def slow_but_alive(method, path, body=None, **_kw):
        _time.sleep(0.5)
        return people.call(method, path, body)
    monkeypatch.setattr(portal, "_people_call", slow_but_alive)
    started = _time.monotonic()
    m = client.post(f"/api/meetings/{mid}/apply", json={"result": doc}).json()["meeting"]
    assert _time.monotonic() - started < 2.5                   # 12 × 0,5 с по очереди было бы 6 с
    assert len(m["registry_pending"]) == 12

    def hanging(method, path, body=None, **_kw):
        _time.sleep(10)
        return people.call(method, path, body)
    monkeypatch.setattr(portal, "_people_call", hanging)
    monkeypatch.setattr(people_flow, "SUGGEST_BUDGET", 0.5)
    started = _time.monotonic()
    r = client.post(f"/api/meetings/{create(client)}/apply", json={"result": doc})
    assert r.status_code == 200 and _time.monotonic() - started < 2.0
    assert r.json()["meeting"]["registry_pending"] == []      # не успели — спросим в следующий раз


def test_speakers_roles_looked_up_in_parallel(client, people, monkeypatch):
    """Докладчики без должности: должности ищутся в реестре параллельно, а не по очереди."""
    import time as _time
    for i in range(4):
        people.rows[f"иванов{i} иван иванович"] = {"name": f"Иванов{i} Иван Иванович",
                                                   "role": f"Директор {i}", "organization": ""}
    login(client, "alice")
    mid = create(client)

    def slow(method, path, body=None, **_kw):
        _time.sleep(0.8)
        return people.call(method, path, body)
    monkeypatch.setattr(portal, "_people_call", slow)
    started = _time.monotonic()
    r = client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О жилье",
                    "speakers": [f"Иванов{i} Иван Иванович" for i in range(4)], "speaker_roles": ["", "", "", ""]})
    assert r.status_code == 200 and _time.monotonic() - started < 2.0     # по очереди было бы 3,2 с
    roles = {p["name"]: p["role"] for p in client.get(f"/api/meetings/{mid}").json()["meeting"]["participants"]}
    assert roles["Иванов3 Иван Иванович"] == "Директор 3"

    # «Искра» добавила вопрос с тремя докладчиками — ФИО и должности тоже одним заходом
    from app.models import Meeting
    started = _time.monotonic()
    item = llm_extract.apply_add_question(Meeting(), {"title": "О сертификатах",
                                                      "speakers": ["Иванов0", "Иванов1", "Иванов2"]})
    assert _time.monotonic() - started < 2.0
    assert item.speakers == ["Иванов0 Иван Иванович", "Иванов1 Иван Иванович", "Иванов2 Иван Иванович"]


def test_name_check_is_cached_between_ui_and_export(client, people, monkeypatch):
    """Интерфейс сверил ФИО — export.docx через секунду второй раз портал не спрашивает."""
    login(client, "alice")
    mid = _meeting_for_word(client)
    checks = []
    real = people.call

    def counting(method, path, body=None, **_kw):
        if path == "/internal/people/check":
            checks.append(body)
        return real(method, path, body, **_kw)
    monkeypatch.setattr(portal, "_people_call", counting)
    client.post(f"/api/meetings/{mid}/people/check")
    assert client.get(f"/api/meetings/{mid}/export.docx").status_code == 200
    assert len(checks) == 1
    assert portal.PEOPLE_TIMEOUT == 1.5


def test_slow_requests_logged_without_data(client, caplog, monkeypatch):
    """Медленный запрос — в server.log метод, форма адреса и длительность, без ключей и данных."""
    import logging as _logging
    login(client, "alice")
    mid = create(client, title="Секретное совещание")
    monkeypatch.setattr(main, "SLOW_SECONDS", -1.0)
    with caplog.at_level(_logging.WARNING, logger="uvicorn.error"):
        client.get(f"/api/meetings/{mid}?q=тайна")
    line = [r.getMessage() for r in caplog.records if "медленный запрос" in r.getMessage()][-1]
    assert "GET /api/meetings/{id}" in line and "код 200" in line
    assert mid not in line and "тайна" not in line and "Секретное" not in line
    assert main.path_shape("/api/meetings/meeting-16/scenario/3fa9c1d2e4") == "/api/meetings/{id}/scenario/{id}"
    assert main.path_shape("/api/meetings/find-or-create") == "/api/meetings/find-or-create"


@pytest.fixture()
def slow_portal(monkeypatch):
    """Настоящий HTTP-«портал»: поиск в реестре отвечает за 2 с, сверка — быстро и находит опечатку."""
    import http.server
    import threading

    class Handler(http.server.BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def _json(self, code, data):
            raw = json.dumps(data, ensure_ascii=False).encode()
            self.send_response(code)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(raw)))
            self.end_headers()
            self.wfile.write(raw)

        def do_GET(self):
            import time as _time
            _time.sleep(2)                      # фоновый поиск должности — медленный
            self._json(200, {"items": []})

        def do_POST(self):
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])) or b"{}")
            if self.path == "/internal/people/check":
                self._json(200, {"items": [{"name": n, "status": "close",
                                            "suggestions": [{"name": "Петрова Анна Ивановна", "role": ""}]}
                                           for n in body["names"]]})
            else:
                self._json(200, {"summary": "", "suggest": "same"})

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    monkeypatch.setenv("PORTAL_URL", f"http://127.0.0.1:{server.server_address[1]}")
    monkeypatch.setattr(portal, "_people_call", REAL_PEOPLE_CALL)
    monkeypatch.setattr(portal, "_people_down_until", 0.0)
    yield server
    server.shutdown()


def test_slow_background_search_does_not_disable_name_check(client, slow_portal):
    """К1 ревьюера RC-2026-09-21-6: медленный фоновый поиск должности включает общую паузу,
    но обязательная сверка ФИО перед Word её не соблюдает — «close» по-прежнему даёт 409, а не 200."""
    login(client, "alice")
    mid = create(client, date="14.10.2026", time_start="10:00")
    # докладчик без должности: при добавлении ищем её в реестре — поиск не укладывается в таймаут
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О жилье",
                "speakers": ["Петрава Анна Ивановна"], "speaker_roles": [""], "speaker_minutes": [10]})
    client.post(f"/api/meetings/{mid}/scenario/generate")
    assert portal._people_down_until > 0                    # общая пауза включилась
    r = client.get(f"/api/meetings/{mid}/export.docx")
    assert r.status_code == 409 and "Петрава" in r.json()["detail"]


def test_word_without_name_check_only_with_consent(client, people, audit):
    """Реестр не ответил — документ не уходит молча: 409 с просьбой подтвердить, потом ?unchecked=1."""
    login(client, "alice")
    mid = _meeting_for_word(client)
    people.down = True
    r = client.get(f"/api/meetings/{mid}/export.docx")
    assert r.status_code == 409 and r.headers["x-names-check"] == "unavailable"
    r = client.get(f"/api/meetings/{mid}/export.docx?unchecked=1")
    assert r.status_code == 200 and r.headers["x-names-check"] == "skipped"
    # в журнале — факт, а не «человек подтвердил»: прямая ссылка тоже ставит флаг
    assert "выгружено без сверки ФИО (реестр не ответил, флаг unchecked)" in audit[-1]["details"]
    # несверенный документ помечен в нём самом: первая строка, колонтитул и имя файла
    import io
    import urllib.parse
    import docx
    d = docx.Document(io.BytesIO(r.content))
    first = [p for p in d.paragraphs if p.text.strip()][0]
    assert first.text.startswith("ФИО не сверены с реестром людей") and first.runs[0].bold
    sec = d.sections[0]
    assert any("ФИО не сверены" in p.text for p in sec.footer.paragraphs)
    assert any("ФИО не сверены" in p.text for p in sec.first_page_footer.paragraphs)
    assert "(ФИО не сверены)" in urllib.parse.unquote(r.headers["x-filename"])
    # сверенный документ пометки не несёт
    people.down = False
    r = client.get(f"/api/meetings/{mid}/export.docx")
    d = docx.Document(io.BytesIO(r.content))
    assert not any("ФИО не сверены" in p.text for p in d.paragraphs)
    assert "(ФИО не сверены)" not in urllib.parse.unquote(r.headers["x-filename"])


def test_html_export_checks_names_like_word(client, people):
    """HTML-бриф — тоже документ (../CLAUDE.md п. 5а): та же сверка, то же согласие, та же пометка."""
    people.rows["иванов иван иванович"] = {"name": "Иванов Иван Иванович", "role": "", "organization": ""}
    login(client, "alice")
    mid = _meeting_for_word(client)                 # докладчики: Иванов и Петрова (её в реестре нет)
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "3", "title": "Ещё",
                "speakers": ["Иванав Иван Иванович"], "speaker_minutes": [5]})
    client.post(f"/api/meetings/{mid}/scenario/generate")
    r = client.get(f"/api/meetings/{mid}/export")
    assert r.status_code == 409 and "Иванав" in r.json()["detail"]             # опечатка не выходит
    client.post(f"/api/meetings/{mid}/people/replace", json={"old": "Иванав Иван Иванович",
                                                             "new": "Иванов Иван Иванович"})
    r = client.get(f"/api/meetings/{mid}/export")
    assert r.status_code == 200 and r.headers["x-names-check"] == "done"
    assert "ФИО не сверены" not in r.text

    people.down = True
    people_flow._checked.clear()
    r = client.get(f"/api/meetings/{mid}/export")
    assert r.status_code == 409 and r.headers["x-names-check"] == "unavailable"
    r = client.get(f"/api/meetings/{mid}/export?unchecked=1")
    assert r.status_code == 200 and r.headers["x-names-check"] == "skipped"
    assert "ФИО не сверены с реестром людей" in r.text


def test_registry_budget_is_per_request_and_late_is_visible(client, people, monkeypatch, caplog):
    """Бюджет — на весь запрос, а не на каждый вызов; опоздавшие — строка в лог и пометка человеку."""
    import logging as _logging
    import time as _time
    monkeypatch.setattr(portal, "PEOPLE_BUDGET", 0.6)
    portal.start_request("alice")
    started = _time.monotonic()
    with caplog.at_level(_logging.WARNING, logger="uvicorn.error"):
        first = portal.people_parallel(lambda k: _time.sleep(1) or k, ["а", "б"], what="первый заход")
        second = portal.people_parallel(lambda k: _time.sleep(1) or k, ["в"], what="второй заход")
    assert first == {} and second == {} and _time.monotonic() - started < 0.9   # не 0,6 + 0,6
    assert portal.registry_late()
    assert any("реестр людей не ответил вовремя: первый заход — 2 из 2" in r.getMessage() for r in caplog.records)

    # в ответе сервиса — пометка, интерфейс скажет «проверьте должности»
    login(client, "alice")
    mid = create(client)

    def hanging(method, path, body=None, **_kw):
        _time.sleep(2)
        return people.call(method, path, body)
    monkeypatch.setattr(portal, "_people_call", hanging)
    doc = {"doc_type": "участники", "confidence": "высокая", "summary": "",
           "data": {"participants": [{"name": "Сидоров Сидор Сидорович", "role": "Директор"}]}}
    r = client.post(f"/api/meetings/{mid}/apply", json={"result": doc})
    assert r.status_code == 200 and r.headers.get("x-registry-late") == "1"


def test_fast_registry_failure_leaves_a_trace(client, people, caplog):
    """Реестр отказал мгновенно — не «нет в реестре» молча, а строка в server.log и пометка человеку."""
    import logging as _logging
    login(client, "alice")
    mid = create(client)
    people.down = True                                       # отказ сразу, без ожидания
    with caplog.at_level(_logging.WARNING, logger="uvicorn.error"):
        r = client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О жилье",
                        "speakers": ["Кадочников Павел Анатольевич"], "speaker_roles": [""]})
    assert r.status_code == 200 and r.headers.get("x-registry-late") == "1"
    lines = [x.getMessage() for x in caplog.records if "реестр людей отказал" in x.getMessage()]
    assert lines == ["реестр людей отказал: должности докладчиков — 1 из 1 без ответа"]
    assert not any("Кадочников" in x.getMessage() for x in caplog.records)   # без данных

    # реестр ответил «нет такого» — это не отказ: следа нет
    people.down = False
    portal._people_down_until = 0.0
    portal._people_cache.clear()
    caplog.clear()
    with caplog.at_level(_logging.WARNING, logger="uvicorn.error"):
        r = client.post(f"/api/meetings/{mid}/participants/fill-roles")
    assert r.status_code == 200 and "x-registry-late" not in r.headers
    assert not any("реестр людей" in x.getMessage() for x in caplog.records)


def test_fill_roles_button_and_check_cache_reset(client, people):
    """Должности, не найденные при добавлении, дозаполняются кнопкой; «Заменить» сбрасывает сверку."""
    login(client, "alice")
    mid = create(client, date="14.10.2026", time_start="10:00")
    people.down = True
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О жилье",
                "speakers": ["Иванов Иван Иванович"], "speaker_roles": [""], "speaker_minutes": [10]})
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["participants"][0]["role"] == ""
    people.down = False
    portal._people_down_until = 0.0
    people.rows["иванов иван иванович"] = {"name": "Иванов Иван Иванович", "role": "Министр", "organization": ""}
    r = client.post(f"/api/meetings/{mid}/participants/fill-roles").json()
    assert r["filled"] == 1 and r["meeting"]["participants"][0]["role"] == "Министр"

    client.post(f"/api/meetings/{mid}/scenario/generate")
    client.post(f"/api/meetings/{mid}/people/check")
    assert any(k[0] == mid for k in people_flow._checked)
    people.rows["петров иван иванович"] = {"name": "Петров Иван Иванович", "role": "", "organization": ""}
    client.post(f"/api/meetings/{mid}/people/replace", json={"old": "Иванов Иван Иванович", "new": "Петров Иван Иванович"})
    assert not any(k[0] == mid for k in people_flow._checked)


def test_registry_conflicts_read_from_portal(client, people):
    people.rows["петров пётр петрович"] = {"name": "Петров Пётр Петрович", "role": "Министр", "organization": "Минтруд"}
    result = {"doc_type": "участники",
              "data": {"participants": [{"name": "Петров Пётр Петрович", "role": "Заместитель министра"}]}}
    conflicts = llm_extract.compute_registry_conflicts(None, result)
    assert conflicts == [{"name": "Петров Пётр Петрович", "new_role": "Заместитель министра",
                          "agenda_role": "", "registry_role": "Министр", "registry_org": "Минтруд"}]


def test_old_meeting_gets_stable_ids_once(client):
    """Совещания, сохранённые до появления id, получают их при чтении — и больше не меняют."""
    login(client, "alice")
    mid = create(client)
    legacy = {
        "id": mid, "title": "Старое",
        "agenda": [{"num": "1", "title": "Вопрос"}],
        "orders": [{"text": "Поручение"}],
        "materials": [{"id": "", "question_num": "1", "proposal": "Согласовать", "intro": "Подводка"}],
        "scenario": [{"section": "question", "agenda_num": "1", "kind": "reference", "text": "Подводка"},
                     {"section": "question", "agenda_num": "1", "kind": "replica", "text": "Предложение: Согласовать"}],
    }
    conn = storage._conn()
    conn.execute("UPDATE meetings SET data = ? WHERE id = ?", (json.dumps(legacy, ensure_ascii=False), mid))
    conn.commit()
    conn.close()
    first = client.get(f"/api/meetings/{mid}").json()["meeting"]
    second = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert first["orders"][0]["id"] and first["orders"][0]["id"] == second["orders"][0]["id"]
    assert first["agenda"][0]["id"] == second["agenda"][0]["id"]
    # блоки сценария старого материала связались с ним — удаление материала уносит их
    mat_id = first["materials"][0]["id"]
    assert [b["material_id"] for b in first["scenario"]] == [mat_id, mat_id]
    assert client.delete(f"/api/meetings/{mid}/materials/{mat_id}").status_code == 200
    after = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert after["scenario"] == [] and after["materials"] == []


def test_meeting_ids_unique_and_explicit_duplicate_rejected(client):
    login(client, "alice")
    ids = {create(client) for _ in range(5)}
    assert len(ids) == 5
    create(client, id="совещание-1")
    r = client.post("/api/meetings", json={"id": "совещание-1", "title": "Другое"})
    assert r.status_code == 409
    login(client, "admin")
    assert client.get("/api/meetings/совещание-1").json()["meeting"]["title"] == "Тест"


def test_cross_site_post_rejected(client):
    login(client, "alice")
    r = client.post("/api/meetings", json={"title": "x"}, headers={"Origin": "https://evil.example"})
    assert r.status_code == 403
    r = client.post("/api/meetings", json={"title": "x"}, headers={"Origin": "http://testserver"})
    assert r.status_code == 200


def test_export_cyrillic_id_and_escaping(client):
    login(client, "alice")
    mid = create(client, id="совещание-2")
    evil = 'x"><img src=x onerror=alert(1)>'
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "Вопрос"})
    client.post(f"/api/meetings/{mid}/orders", json={"text": "Сделать", "agenda_num": evil})
    r = client.get(f"/api/meetings/{mid}/export")
    assert r.status_code == 200
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    assert "<img src=x onerror" not in r.text


def test_upload_size_limit(client, monkeypatch):
    monkeypatch.setattr(deps, "MAX_UPLOAD_BYTES", 10)
    login(client, "alice")
    r = client.post("/api/extract", files={"file": ("a.txt", b"x" * 100)})
    assert r.status_code == 413


def test_zip_bomb_rejected(client, monkeypatch):
    from app import extract
    monkeypatch.setattr(extract, "MAX_UNZIPPED_BYTES", 1000)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "a" * 100_000)
    login(client, "alice")
    r = client.post("/api/extract", files={"file": ("bomb.docx", buf.getvalue())})
    assert r.status_code == 422
    assert "распаковки" in r.json()["detail"]


# ── проверки, добавленные при приёмке ───────────────────────────────────

def test_apply_checks_schema(client):
    """В /apply приходит произвольный JSON из браузера — он должен пройти проверку схемы."""
    login(client, "alice")
    mid = create(client)
    r = client.post(f"/api/meetings/{mid}/apply", json={"result": {"agenda": [
        {"num": "1", "title": "Вопрос", "speakers": ["Иванов И.И."], "speaker_minutes": ["много"]}]}})
    assert r.status_code in (200, 422)              # либо отказ, либо очищенные схемой данные
    agenda = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"]
    for item in agenda:                            # строка «много» не должна остаться в минутах
        assert all(isinstance(x, int) for x in item["speaker_minutes"])
    r = client.post(f"/api/meetings/{mid}/materials/apply", json={"result": {"theses": "абв", "orders": "ху"}})
    assert r.status_code in (200, 422)
    if r.status_code == 200:
        mat = client.get(f"/api/meetings/{mid}").json()["meeting"]["materials"][-1]
        assert mat["theses"] != ["а", "б", "в"]    # строка не должна разваливаться по буквам


def test_renumbering_moves_orders_and_materials(client):
    """Перестановка вопроса не должна уводить поручение и материал к чужому вопросу."""
    login(client, "alice")
    mid = create(client)
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "Первый"})
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "2", "title": "Второй"})
    client.post(f"/api/meetings/{mid}/orders", json={"text": "Поручение ко второму", "agenda_num": "2"})
    client.post(f"/api/meetings/{mid}/materials", json={"question_num": "2", "title": "Материал ко второму"})
    second = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][1]
    r = client.post(f"/api/meetings/{mid}/agenda/{second['id']}/move", json={"direction": "up"})
    assert r.status_code == 200
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert [a["title"] for a in m["agenda"]] == ["Второй", "Первый"]
    assert [a["num"] for a in m["agenda"]] == ["1", "2"]
    assert m["orders"][0]["agenda_num"] == "1"        # поручение уехало вместе со «Вторым»
    assert m["materials"][0]["question_num"] == "1"


def test_material_text_is_limited(client, monkeypatch):
    from app import deps
    monkeypatch.setattr(deps, "MAX_TEXT_CHARS", 100)
    login(client, "alice")
    mid = create(client)
    r = client.post(f"/api/meetings/{mid}/materials", json={"title": "Большой", "text": "я" * 200})
    assert r.status_code == 413


def test_rashod_iskry_popadaet_v_zhurnal():
    """Расход токенов копится в корзинке запроса — журнал портала берёт его оттуда."""
    from app import llm_extract, portal
    bucket = portal.start_request()
    llm_extract._note_usage({"status": "completed", "usage": {"prompt_tokens": 900, "completion_tokens": 120}})
    llm_extract._note_usage({"result": {"usage": {"input_tokens": 100, "output_tokens": 80}}})
    llm_extract._note_usage({"status": "failed"})       # без usage — просто ничего не добавилось
    assert (bucket["tokens_in"], bucket["tokens_out"]) == (1000, 200)


def test_material_carries_link_home_and_refuses_foreign_site(client):
    """Материал из другого сервиса несёт ссылку обратно; чужой сайт в неё не пропускаем."""
    login(client, "emp")
    mid = create(client, title="Проверка ссылок")

    good = client.post(f"/api/meetings/{mid}/materials", json={
        "title": "Бриф: Удмуртия", "type": "Бриф",
        "source": {"service": "brief", "id": "2026-09-20_x", "kind": "бриф", "url": "/brief/#brief=2026-09-20_x"}})
    assert good.status_code == 200
    assert good.json()["meeting"]["materials"][0]["source"]["url"] == "/brief/#brief=2026-09-20_x"

    bad = client.post(f"/api/meetings/{mid}/materials", json={
        "title": "Чужое", "source": {"service": "brief", "id": "1", "url": "https://evil.example"}})
    assert bad.status_code == 200
    assert bad.json()["meeting"]["materials"][1]["source"]["url"] == ""


def test_all_orders_collects_across_meetings(client):
    """Реестр поручений: со всех совещаний, просроченные сверху, с ссылкой на совещание."""
    login(client, "emp")
    first = create(client, title="Совещание по сиротам", date="01.09.2026")
    second = create(client, title="Отопительный сезон", date="15.10.2026")
    client.post(f"/api/meetings/{first}/orders", json={"text": "Подготовить свод", "assignee": "Иванов И.И.",
                                                        "deadline": "01.01.2020"})
    client.post(f"/api/meetings/{second}/orders", json={"text": "Проверить котельные", "assignee": "Петров П.П.",
                                                         "deadline": "31.12.2099"})
    client.post(f"/api/meetings/{second}/orders", json={"text": "Без срока", "assignee": "Иванов И.И."})

    orders = client.get("/api/orders").json()["orders"]
    assert [o["text"] for o in orders] == ["Подготовить свод", "Проверить котельные", "Без срока"]
    assert orders[0]["overdue"] is True and orders[1]["overdue"] is False
    assert orders[0]["url"] == f"/moderator/#meeting={first}"
    assert orders[0]["meeting_title"] == "Совещание по сиротам"

    mine = client.get("/api/orders?who=иванов").json()["orders"]
    assert {o["text"] for o in mine} == {"Подготовить свод", "Без срока"}
    assert [o["text"] for o in client.get("/api/orders?overdue=true").json()["orders"]] == ["Подготовить свод"]


def test_access_is_changed_by_owner_only(client, monkeypatch, audit):
    """Доступ к совещанию: открывает владелец, чужой — нет; изменения попадают в журнал."""
    monkeypatch.setattr(portal, "people", lambda: [{"login": "emp", "name": "Аня"},
                                                    {"login": "kolya", "name": "Коля"}])
    login(client, "emp")
    mid = create(client, title="Совместная работа")

    r = client.put(f"/api/meetings/{mid}/access", json={"shared_with": ["kolya", "net-takogo"]})
    assert r.status_code == 200
    assert r.json()["shared_with"] == ["kolya"]          # неизвестный логин отброшен

    login(client, "kolya")
    assert client.put(f"/api/meetings/{mid}/access", json={"shared_with": []}).status_code == 403

    login(client, "emp")
    assert client.get(f"/api/meetings/{mid}/access").json()["owner"] == "emp"
    assert any("открыт доступ: kolya" in (e.get("details") or "") for e in audit)


def test_names_come_from_portal_registry(client, people):
    """Подсказка по людям — только из общего реестра портала."""
    people.rows["северцев игорь олегович"] = {"name": "Северцев Игорь Олегович",
                                              "role": "директор департамента", "organization": "минобр"}
    login(client, "emp")
    matches = client.get("/api/names/resolve?q=севе").json()["matches"]
    assert matches == [{"name": "Северцев Игорь Олегович", "surname": "Северцев",
                        "role": "директор департамента", "organization": "минобр",
                        "source": "реестр людей"}]


def test_silent_portal_registry_does_not_break_work(client, people):
    """Портал молчит — подсказок и окна нет, но разбор применяется, а Word собирается с предупреждением."""
    people.down = True
    login(client, "emp")
    assert client.get("/api/names/resolve?q=кто-то").json()["matches"] == []
    mid = _meeting_for_word(client)
    doc = {"doc_type": "участники", "confidence": "высокая", "summary": "",
           "data": {"participants": [{"name": "Сидоров Сидор Сидорович", "role": "Директор"}]}}
    m = client.post(f"/api/meetings/{mid}/apply", json={"result": doc}).json()["meeting"]
    assert m["registry_pending"] == []
    assert client.post(f"/api/meetings/{mid}/people/check").json()["unavailable"] is True
    # без сверки документ не уходит молча: сначала вопрос человеку, потом выгрузка с его согласия
    assert client.get(f"/api/meetings/{mid}/export.docx").status_code == 409
    r = client.get(f"/api/meetings/{mid}/export.docx?unchecked=1")
    assert r.status_code == 200 and r.headers["x-names-check"] == "skipped"

def test_people_replace_fills_role_from_registry(client):
    """Разбор записал докладчика коротко — «Кравцов» (просьба пользователя 02.10). «Заменить на …»
    из реестра приносит полное ФИО и должность: ФИО меняется во всём совещании, должность встаёт
    там, где пуста (докладчик вопроса, участник), а заполненную руками не трогает. `old == new`
    с `role` — только дозаполнить должность."""
    login(client, "alice")
    mid = create(client)
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О готовности к Дню учителя"})
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "2", "title": "Разное"})
    agenda = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"]
    k1, k2 = agenda[0]["id"], agenda[1]["id"]
    client.post(f"/api/meetings/{mid}/agenda/{k1}/speakers", json={"name": "Кравцов"})          # как пишет разбор
    client.post(f"/api/meetings/{mid}/agenda/{k2}/speakers", json={"name": "Кравцов"})
    client.post(f"/api/meetings/{mid}/agenda/{k2}/speakers", json={"name": "Сидорова Анна Петровна", "role": "Директор департамента"})
    before = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert [p["name"] for p in before["participants"]] == ["Кравцов", "Сидорова Анна Петровна"]   # докладчики → участники
    assert before["participants"][0]["role"] == ""

    r = client.post(f"/api/meetings/{mid}/people/replace",
                    json={"old": "Кравцов", "new": "Кравцов Сергей Сергеевич", "role": "Министр просвещения"})
    assert r.status_code == 200, r.text
    m = r.json()["meeting"]
    assert m["agenda"][0]["speakers"] == ["Кравцов Сергей Сергеевич"]
    assert m["agenda"][0]["speaker_roles"][0] == "Министр просвещения"
    assert m["agenda"][1]["speakers"] == ["Кравцов Сергей Сергеевич", "Сидорова Анна Петровна"]
    assert m["agenda"][1]["speaker_roles"] == ["Министр просвещения", "Директор департамента"]
    assert m["participants"][0]["name"] == "Кравцов Сергей Сергеевич"
    assert m["participants"][0]["role"] == "Министр просвещения"
    # заполненную руками должность замена из реестра не трогает — то же ФИО, менять нечего: 200, replaced 0
    r = client.post(f"/api/meetings/{mid}/people/replace",
                    json={"old": "Сидорова Анна Петровна", "new": "Сидорова Анна Петровна", "role": "Заместитель директора"})
    assert r.status_code == 200 and r.json()["replaced"] == 0
    assert client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][1]["speaker_roles"][1] == "Директор департамента"

    # одно ФИО, только должность: было пусто — встало; второй раз менять нечего — 200, replaced 0
    client.post(f"/api/meetings/{mid}/participants", json={"name": "Петров Пётр Петрович"})
    r = client.post(f"/api/meetings/{mid}/people/replace",
                    json={"old": "Петров Пётр Петрович", "new": "Петров Пётр Петрович", "role": "Директор"})
    assert r.status_code == 200 and r.json()["replaced"] == 1
    assert [p["role"] for p in r.json()["meeting"]["participants"] if p["name"] == "Петров Пётр Петрович"] == ["Директор"]
    # то же ФИО, должность уже стоит — не ошибка: 200 и replaced 0 (окно «Добавить в реестр?» без правок
    # не должно показывать красный toast — ревьюер RC-2026-10-02-1, Н2); без role — тоже 200 и replaced 0
    r = client.post(f"/api/meetings/{mid}/people/replace",
                    json={"old": "Петров Пётр Петрович", "new": "Петров Пётр Петрович", "role": "Директор"})
    assert r.status_code == 200 and r.json()["replaced"] == 0
    r = client.post(f"/api/meetings/{mid}/people/replace",
                    json={"old": "Петров Пётр Петрович", "new": "Петров Пётр Петрович"})
    assert r.status_code == 200 and r.json()["replaced"] == 0      # то же ФИО без должности — тоже не ошибка
    # список должностей дополняется до числа докладчиков, не до индекса (ревьюер, Н5): у вопроса два
    # докладчика без должностей, замена первого — roles длиной 2
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "3", "title": "Третий"})
    k3 = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][2]["id"]
    for n in ("Иванов", "Смирнов"):
        client.post(f"/api/meetings/{mid}/agenda/{k3}/speakers", json={"name": n})
    r = client.post(f"/api/meetings/{mid}/people/replace",
                    json={"old": "Иванов", "new": "Иванов Иван Иванович", "role": "Начальник отдела"})
    item = r.json()["meeting"]["agenda"][2]
    assert item["speakers"] == ["Иванов Иван Иванович", "Смирнов"]
    assert item["speaker_roles"] == ["Начальник отдела", ""]


def test_apply_agenda_keeps_known_people_out_of_registry_queue(client, people):
    """Разбор написал докладчика коротко — «Кравцов», а в реестре он есть полностью. Такого человека
    не предлагают добавить в реестр (окно очереди «Добавить в реестр?» записало бы короткую запись),
    а заменяют через окно сверки; после «Заменить на …» из очереди он тоже уходит (ревьюер
    RC-2026-10-02-1, Н1). Неизвестный реестру («Новиков») — в очереди как прежде. В очередь при
    `/apply` попадают докладчики с должностью (так было и раньше) — поэтому должности в данных есть."""
    people.rows[people.key("Кравцов Сергей Сергеевич")] = {
        "name": "Кравцов Сергей Сергеевич", "role": "Министр просвещения", "organization": "Минпросвещения"}
    login(client, "alice")
    mid = create(client)
    result = {"doc_type": "повестка", "confidence": "высокая", "summary": "со слов",
              "data": {"agenda": [{"num": "1", "title": "О готовности к Дню учителя",
                                   "speakers": ["Кравцов", "Новиков"], "speaker_roles": ["Министр", "Начальник отдела"],
                                   "speaker_minutes": [10, 5]}]}}
    r = client.post(f"/api/meetings/{mid}/apply", json={"result": result, "replace": False, "resolutions": {}})
    assert r.status_code == 200, r.text
    m = r.json()["meeting"]
    assert m["agenda"][0]["speakers"] == ["Кравцов", "Новиков"]
    assert [p["name"] for p in m["registry_pending"]] == ["Новиков"]       # «Кравцов» — не добавлять, а заменить
    # сверка показывает полную запись как close — её и предложит окно «Докладчики и реестр людей»
    items = client.post(f"/api/meetings/{mid}/people/check-agenda").json()["items"]
    by = {x["name"]: x for x in items}
    assert by["Кравцов"]["status"] == "close" and by["Кравцов"]["suggestions"][0]["name"] == "Кравцов Сергей Сергеевич"
    assert by["Новиков"]["status"] == "unknown"
    # замена ФИО закрывает и очередь по старому имени
    r = client.post(f"/api/meetings/{mid}/people/replace", json={"old": "Новиков", "new": "Новиков Николай Николаевич"})
    assert r.status_code == 200
    assert r.json()["meeting"]["registry_pending"] == []


def test_known_people_leave_registry_queue_on_changes_manual_and_import(client, people, monkeypatch):
    """Очередь «Добавить в реестр?» чистится от тех, кого реестр знает полностью, не только после `/apply`:
    после списка изменений (`apply-changes`), ручного докладчика и импорта из «Повестки» (ревьюер
    RC-2026-10-02-2, Н-3). Реестр не ответил — очередь не трогаем. Сверяется тот же набор, что и
    `check-agenda`, поэтому второй запрос к порталу берётся из кэша (Н-5)."""
    people.rows[people.key("Кравцов Сергей Сергеевич")] = {
        "name": "Кравцов Сергей Сергеевич", "role": "Министр просвещения", "organization": "Минпросвещения"}
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    # ручной докладчик коротко — в очереди его нет, неизвестный «Новиков» — есть
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "Вопрос"})
    k1 = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][0]["id"]
    client.post(f"/api/meetings/{mid}/agenda/{k1}/speakers", json={"name": "Кравцов", "role": "Министр"})
    r = client.post(f"/api/meetings/{mid}/agenda/{k1}/speakers", json={"name": "Новиков", "role": "Начальник отдела"})
    assert [p["name"] for p in r.json()["meeting"]["registry_pending"]] == ["Новиков"]
    # сверка интерфейса — тот же набор, кэш: к порталу новый запрос не идёт
    checks_before = people.checks
    items = client.post(f"/api/meetings/{mid}/people/check-agenda").json()["items"]
    assert {x["name"]: x["status"] for x in items} == {"Кравцов": "close", "Новиков": "unknown"}
    assert people.checks == checks_before          # тот же набор ФИО — ответ из кэша, портал не спрашивали

    # список изменений (непустая повестка → ops от настоящего /api/analyze с подменённой «Искрой»):
    # добавить вопрос с докладчиком «Кравцов» — в очередь он снова не попадает
    people_flow.forget_checks(mid)
    doc = {"doc_type": "повестка", "confidence": "высокая", "summary": "",
           "data": {"agenda": [
               {"title": "Вопрос", "speakers": ["Кравцов", "Новиков"], "speaker_roles": ["Министр", "Начальник отдела"],
                "speaker_minutes": [10, 5]},
               {"title": "Второй", "speakers": ["Кравцов"], "speaker_roles": ["Министр"], "speaker_minutes": [5]},
           ]}}
    _mock_iskra(monkeypatch, doc)
    changes = client.post("/api/analyze", json={"text": "текст повестки", "meeting_id": mid}).json()["changes"]
    add_op = next(op for op in changes if op["kind"] == "add_question")
    r = client.post(f"/api/meetings/{mid}/agenda/apply-changes", json={"ops": [add_op]})
    assert r.status_code == 200, r.text
    assert len(r.json()["meeting"]["agenda"]) == 2
    assert [p["name"] for p in r.json()["meeting"]["registry_pending"]] == ["Новиков"]

    # счёт запросов: применение само сверило, интерфейсная сверка следом — из кэша (Н-5 RC-2 / RC-3 кр. 8)
    checks_before = people.checks
    client.post(f"/api/meetings/{mid}/people/check-agenda")
    assert people.checks == checks_before

    # участник коротко, он же докладчик — тоже не в очередь (RC-3, Н-5)
    r = client.post(f"/api/meetings/{mid}/participants", json={"name": "Кравцов", "role": "Министр", "organization": "Минпрос"})
    assert "Кравцов" not in [p["name"] for p in r.json()["meeting"]["registry_pending"]]

    # реестр молчит — очередь не трогаем: короткий «Сидоров» (в реестре есть «Сидоров Семён Семёнович») остаётся
    people.rows[people.key("Сидоров Семён Семёнович")] = {"name": "Сидоров Семён Семёнович", "role": "Советник", "organization": ""}
    people_flow.forget_checks(mid)
    monkeypatch.setattr(portal, "people_check", lambda names: None)
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "3", "title": "Третий"})
    k3 = client.get(f"/api/meetings/{mid}").json()["meeting"]["agenda"][2]["id"]
    r = client.post(f"/api/meetings/{mid}/agenda/{k3}/speakers", json={"name": "Сидоров", "role": "Советник"})
    names = [p["name"] for p in r.json()["meeting"]["registry_pending"]]
    assert "Сидоров" in names and "Новиков" in names


def test_import_from_agenda_drops_known_people_from_queue(client, people):
    """Импорт из «Повестки»: докладчик коротко, реестр знает полностью — в очередь не ставится (Н-3)."""
    people.rows[people.key("Кравцов Сергей Сергеевич")] = {
        "name": "Кравцов Сергей Сергеевич", "role": "Министр просвещения", "organization": "Минпросвещения"}
    login(client, "alice")
    body = {"source": {"service": "agenda", "meeting_id": "777"}, "target_id": "", "title": "Совещание у Кравцова",
            "date": "2026-10-14",
            "agenda": [{"num": "I", "title": "Вопрос", "speakers": ["Кравцов", "Новиков"],
                        "speaker_roles": ["Министр", "Начальник отдела"], "speaker_minutes": [10, 5],
                        "speaker_kinds": ["main", "main"], "minutes": 0}],
            "participants": [{"name": "Кравцов", "role": "Министр", "organization": "", "is_speaker": True, "attendance": "in_person"},
                             {"name": "Новиков", "role": "Начальник отдела", "organization": "", "is_speaker": True, "attendance": "in_person"}]}
    r = client.post("/api/meetings/import-from-agenda", json=body)
    assert r.status_code == 200, r.text
    m = client.get(f"/api/meetings/{r.json()['id']}").json()["meeting"]
    assert [p["name"] for p in m["registry_pending"]] == ["Новиков"]


def test_extract_retries_strictly_when_iskra_answers_in_prose(monkeypatch):
    """«Добавь вопрос Кравцова по поводу подготовки к празднику осени» — Искра приняла просьбу за команду себе и
    ответила текстом про «Планы» (пользователь 02.10). Первый ответ прозой → вторая попытка со строгим
    напоминанием про JSON впереди текста; если Искра и тогда ответила прозой — прежняя ошибка. В системном
    промпте есть правило про просьбу со слов как повестку с одним вопросом."""
    calls = []
    answers = ["Добавил в Планы задачу-вопрос: **Вопрос Кравцова** – без срока.",
               json.dumps({"doc_type": "повестка", "confidence": "высокая", "summary": "со слов",
                           "data": {"agenda": [{"num": "1", "title": "О подготовке к празднику осени",
                                                "speakers": ["Кравцов"], "speaker_roles": [""], "speaker_minutes": []}]}},
                          ensure_ascii=False)]
    def fake(system, user, *a, **kw):
        calls.append((system, user))
        return answers[len(calls) - 1]
    monkeypatch.setattr(llm_extract, "_call_iskra_raw", fake)
    res = llm_extract.extract("Добавь вопрос Кравцова по поводу подготовки к празднику осени")
    assert res["doc_type"] == "повестка" and res["data"]["agenda"][0]["speakers"] == ["Кравцов"]
    assert len(calls) == 2
    assert not calls[0][1].startswith(llm_extract.STRICT_REMINDER)          # первая попытка — как обычно
    assert calls[1][1].startswith(llm_extract.STRICT_REMINDER)              # вторая — со строгим напоминанием
    assert calls[1][1].endswith("Добавь вопрос Кравцова по поводу подготовки к празднику осени")
    assert "короткой просьбой со слов" in llm_extract.SYSTEM_PROMPT and "только JSON" in llm_extract.SYSTEM_PROMPT

    calls.clear()
    answers[:] = ["Извините, не понял.", "Всё равно текст."]
    with pytest.raises(RuntimeError, match="невалидный JSON"):
        llm_extract.extract("что-то")
    assert len(calls) == 2


def test_word_subtitle_in_genitive():
    """07.10 (пользователь): «СПИСОК / участников Совещание» → «участников совещания …»; «ПОВЕСТКА» — так же."""
    from app.agenda_docx import event_genitive
    assert event_genitive("Совещание") == "совещания"
    assert event_genitive("") == "совещания"
    assert event_genitive("Совещание у Заместителя Министра по вопросам ремонта школ") == "совещания у Заместителя Министра по вопросам ремонта школ"
    assert event_genitive("совещания у Заместителя Председателя Правительства") == "совещания у Заместителя Председателя Правительства"
    assert event_genitive("Рабочее совещание по итогам года") == "рабочего совещания по итогам года"
    assert event_genitive("Расширенное заседание коллегии") == "расширенного заседания коллегии"
    assert event_genitive("Рабочая встреча с ректорами") == "рабочей встречи с ректорами"
    assert event_genitive("Селекторное совещание-семинар") == "селекторного совещания-семинар"
    assert event_genitive("В режиме ВКС") == "в режиме ВКС"
    # ревью RC-2026-10-07-2: г/к/х — не шипящие; аббревиатура в начале; знак препинания после вида мероприятия
    assert event_genitive("Краткая встреча с ректорами") == "краткой встречи с ректорами"
    assert event_genitive("Тематическая встреча") == "тематической встречи"
    assert event_genitive("ВКС-совещание по итогам") == "ВКС-совещание по итогам"
    assert event_genitive("Совещание, посвящённое началу года") == "совещания, посвящённое началу года"


def test_export_participants_docx_heading_genitive(client, people):
    import io
    import docx
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.put(f"/api/meetings/{mid}/header", json={"subtitle": "Совещание"})
    client.post(f"/api/meetings/{mid}/participants", json={"name": "Сидоров Сидор Сидорович", "role": "Директор"})
    r = client.get(f"/api/meetings/{mid}/export/participants.docx")
    assert r.status_code == 200, r.text
    text = _docx_text(docx.Document(io.BytesIO(r.content)))
    assert "участников совещания" in text and "участников Совещание" not in text


def test_material_can_be_moved_to_another_question(client):
    """Решение пользователя 07.10: материал прикрепился не к тому вопросу — в правке материала выбирается
    нужный вопрос; за материалом переезжают его поручения и блоки сценария."""
    login(client, "alice")
    mid = create(client, date="14.10.2026")
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "1", "title": "О ремонте школ"})
    client.post(f"/api/meetings/{mid}/agenda", json={"num": "2", "title": "О питании"})
    r = client.post(f"/api/meetings/{mid}/materials", json={
        "question_num": "1", "title": "Справка о питании", "intro": "Подводка о питании",
        "conclusion": "Нужно решить", "proposals": ["Утвердить меню"], "text": "текст"})
    assert r.status_code == 200, r.text
    mat_id = r.json()["material_id"]
    client.patch(f"/api/meetings/{mid}/materials/{mat_id}", json={"orders": [{"text": "Проверить столовые"}]})
    m = client.get(f"/api/meetings/{mid}").json()["meeting"]
    assert any(o["material_id"] == mat_id and o["agenda_num"] == "1" for o in m["orders"])

    r = client.patch(f"/api/meetings/{mid}/materials/{mat_id}", json={"question_num": "2"})
    assert r.status_code == 200, r.text
    m = r.json()["meeting"]
    mat = next(x for x in m["materials"] if x["id"] == mat_id)
    assert mat["question_num"] == "2"
    assert all(o["agenda_num"] == "2" for o in m["orders"] if o["material_id"] == mat_id)
    assert all(b["agenda_num"] == "2" for b in m["scenario"] if b.get("material_id") == mat_id)
    # несуществующий вопрос — отказ, материал на месте
    assert client.patch(f"/api/meetings/{mid}/materials/{mat_id}", json={"question_num": "9"}).status_code == 400
    assert next(x for x in client.get(f"/api/meetings/{mid}").json()["meeting"]["materials"] if x["id"] == mat_id)["question_num"] == "2"
    # снять привязку — общие материалы
    assert client.patch(f"/api/meetings/{mid}/materials/{mat_id}", json={"question_num": ""}).status_code == 200
    # ревью RC-2026-10-07-5: коллега перенёс материал, у второго открыто старое окно — его сохранение даёт 409,
    # а не возвращает материал назад (question_num входит в основу правки)
    client.patch(f"/api/meetings/{mid}/materials/{mat_id}", json={"question_num": "2"})
    r = client.patch(f"/api/meetings/{mid}/materials/{mat_id}", json={"title": "Справка (правка)", "question_num": "1",
                                                                      "base": {"question_num": ""}})
    assert r.status_code == 409
    assert next(x for x in client.get(f"/api/meetings/{mid}").json()["meeting"]["materials"] if x["id"] == mat_id)["question_num"] == "2"
